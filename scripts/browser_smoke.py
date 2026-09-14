from __future__ import annotations

import os
import re
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from playwright.sync_api import sync_playwright


def create_sample(path: Path) -> None:
    document = Document()
    document.sections[0].header.paragraphs[0].text = "{{school.name}}"
    document.add_paragraph("Agent 论文排版工具研究", style="Title")
    document.add_paragraph("摘要")
    document.add_paragraph("本文验证本地 Web、独立 Worker 和 DOCX 完整性检查流程。")
    document.add_paragraph("关键词：Agent；论文排版；Word")
    document.add_paragraph("第一章 绪论", style="Heading 1")
    document.add_paragraph(
        "论文排版应先生成计划，相关研究已给出证据[1]。"
    )
    document.add_paragraph("1.1 研究背景", style="Heading 2")
    field_paragraph = document.add_paragraph(
        "这是用于浏览器端到端测试的普通正文，当前页码是 "
    )
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    field_paragraph.add_run()._r.append(begin)
    field_paragraph.add_run()._r.append(instruction)
    field_paragraph.add_run()._r.append(end)
    document.add_paragraph("参考文献")
    document.add_paragraph("[1] 张三. 论文排版方法研究[J]. 示例期刊, 2026.")
    document.add_section(WD_SECTION.NEW_PAGE)
    document.add_paragraph("附录 A 浏试材料")
    document.save(str(path))


def create_composition_template(path: Path) -> None:
    document = Document()
    document.add_paragraph("浏览器验证大学")
    document.add_paragraph("{{paper.title}}")
    document.add_paragraph("摘要：{{paper.abstract}}")
    document.add_paragraph("关键词：{{paper.keywords}}")
    document.add_paragraph("{{document.body}}")
    document.sections[0].header.paragraphs[0].text = "浏览器验证大学学位论文"
    document.save(str(path))


def main() -> None:
    base_url = os.environ.get("PAPER_SETTING_BROWSER_URL", "http://127.0.0.1:8890")
    with tempfile.TemporaryDirectory() as temporary:
        source = Path(temporary) / "browser-smoke.docx"
        create_sample(source)
        composition_template = Path(temporary) / "browser-composition-template.docx"
        create_composition_template(composition_template)
        worker_log = Path(temporary) / "worker.log"
        with worker_log.open("w", encoding="utf-8") as log:
            worker = subprocess.Popen(
                [sys.executable, "-m", "paper_setting_worker.main"],
                env=os.environ.copy(),
                stdout=log,
                stderr=subprocess.STDOUT,
                text=True,
            )
            try:
                with sync_playwright() as playwright:
                    browser = playwright.chromium.launch(headless=True)
                    page = browser.new_page(accept_downloads=True)
                    console_errors: list[str] = []
                    page.on(
                        "console",
                        lambda message: console_errors.append(message.text)
                        if message.type == "error"
                        else None,
                    )
                    page.goto(base_url)
                    page.wait_for_load_state("networkidle")
                    page.get_by_role("button", name="分析学校 DOCX 模板").click()
                    page.get_by_label("选择模板 DOCX").set_input_files(str(source))
                    page.get_by_role("button", name="分析模板结构").click()
                    page.get_by_text("能力覆盖", exact=True).wait_for(timeout=15_000)
                    page.get_by_placeholder("填写 school.name").fill("浏览器测试大学")
                    page.get_by_label("第 2 节起始页码").fill("5")
                    page.get_by_label("第 2 节页眉关系").select_option("independent_copy")
                    page.get_by_role("button", name="预览组合计划").click()
                    page.get_by_text("合计操作", exact=True).wait_for(timeout=15_000)
                    page.get_by_text("结构操作", exact=True).wait_for(timeout=15_000)
                    with page.expect_download(timeout=15_000) as download_info:
                        page.get_by_role("button", name="生成并下载组合 DOCX").click()
                    combined_template = Path(temporary) / "template-combined.docx"
                    download_info.value.save_as(str(combined_template))
                    combined_document = Document(str(combined_template))
                    if combined_document.sections[1].header.is_linked_to_previous:
                        raise AssertionError("第 2 节页眉未复制为独立部件")
                    page_number_nodes = combined_document.sections[1]._sectPr.xpath(
                        "./w:pgNumType"
                    )
                    if not page_number_nodes or page_number_nodes[0].get(qn("w:start")) != "5":
                        raise AssertionError("第 2 节起始页码未应用")
                    header_text = combined_document.sections[0].header.paragraphs[0].text
                    if "浏览器测试大学" not in header_text:
                        raise AssertionError("组合执行未保存页眉占位符替换")
                    page.get_by_role("button", name="收起学校模板分析").click()
                    page.get_by_role(
                        "button", name="从规范文本或文件生成规则包"
                    ).click()
                    page.get_by_label("规则包名称").fill("浏览器测试规则")
                    page.get_by_label("规则包 ID（可选）").fill("browser-generated-rules")
                    page.get_by_label("粘贴格式规范").fill(
                        "页面采用 A4 纸，纵向，上下边距 2.5cm，左边距 3cm，右边距 2.5cm。\n"
                        "正文：宋体，小四，Times New Roman，1.5 倍行距，"
                        "首行缩进 2 字符，两端对齐。\n"
                        "一级标题：黑体，三号，加粗，居中。\n"
                        "注释一律采用脚注，文末不列参考文献。"
                        "每页脚注重新编号，编号格式为带圈数字。"
                    )
                    page.get_by_role("button", name="识别并生成草稿").click()
                    page.get_by_text("已识别属性", exact=True).wait_for(timeout=15_000)
                    page.get_by_text("查看提取证据", exact=False).wait_for()
                    page.get_by_role("button", name="确认并导入规则包").click()
                    page.locator('option[value="browser-generated-rules"]').wait_for(
                        state="attached", timeout=15_000
                    )
                    if page.locator(".form-grid select").input_value() != "browser-generated-rules":
                        raise AssertionError("generated rule pack was not selected")
                    upload = page.locator('.dropzone input[type="file"]')
                    if upload.count() != 1:
                        page.screenshot(
                            path="/tmp/paper-setting-browser-inspect.png", full_page=True
                        )
                        raise AssertionError(
                            {
                                "title": page.title(),
                                "url": page.url,
                                "body": page.locator("body").inner_text()[:1000],
                                "console_errors": console_errors,
                            }
                        )
                    upload.set_input_files(str(source))
                    page.get_by_role("button", name="检查论文并生成计划").click()
                    try:
                        page.get_by_role("heading", name="文档检查结果").wait_for(
                            timeout=30_000
                        )
                        page.get_by_role("heading", name="审核修改计划").wait_for(
                            timeout=30_000
                        )
                    except Exception as exc:
                        page.screenshot(
                            path="/tmp/paper-setting-browser-failed.png", full_page=True
                        )
                        log.flush()
                        raise AssertionError(
                            {
                                "body": page.locator("body").inner_text()[:2000],
                                "console_errors": console_errors,
                                "worker_log": worker_log.read_text(encoding="utf-8")[-2000:],
                            }
                        ) from exc
                    page.get_by_text("查看修改前后").first.click()
                    page.get_by_text("修改后").first.wait_for()
                    page.get_by_role("button", name="选择全部可排版项").click()
                    page.get_by_text("需确认", exact=True).first.wait_for()
                    page.get_by_text("受控重构", exact=True).first.wait_for()
                    page.get_by_label(
                        "我已复核语义角色和深层操作，"
                        "确认执行计划内的文本与 OOXML 变化"
                    ).check()
                    page.get_by_role("button", name=re.compile(r"^批准并执行")).click()
                    page.get_by_role("heading", name="排版与完整性验证完成").wait_for(
                        timeout=60_000
                    )
                    page.get_by_text("内容完整性", exact=True).wait_for()
                    links = page.locator(".artifacts a")
                    format_artifact_count = links.count()
                    if format_artifact_count < 4:
                        raise AssertionError(
                            f"expected at least 4 artifacts, got {format_artifact_count}"
                        )
                    if console_errors:
                        raise AssertionError(f"browser console errors: {console_errors}")
                    page.get_by_role("button", name="新建任务").click()
                    page.locator('.dropzone input[type="file"]').set_input_files(str(source))
                    page.locator('.template-upload input[type="file"]').set_input_files(
                        str(composition_template)
                    )
                    page.get_by_role("button", name="检查论文并生成计划").click()
                    page.get_by_role(
                        "heading", name="审核整篇正文注入计划"
                    ).wait_for(timeout=30_000)
                    page.get_by_text("智能字段映射", exact=True).wait_for()
                    if page.get_by_text("整篇正文注入", exact=True).count() != 1:
                        raise AssertionError("正文注入操作未显示")

                    # 重载后仍由 localStorage 恢复当前任务。
                    page.reload()
                    page.get_by_role(
                        "heading", name="审核整篇正文注入计划"
                    ).wait_for(timeout=30_000)
                    page.get_by_role("button", name="批准并生成模板成品").click()
                    page.get_by_role(
                        "heading", name="排版与完整性验证完成"
                    ).wait_for(timeout=60_000)
                    composed_link = page.locator('.artifacts a[href$="/template_docx"]')
                    if composed_link.count() != 1:
                        raise AssertionError("模板注入成品不存在")
                    with page.expect_download(timeout=15_000) as composition_download:
                        composed_link.click()
                    composed_path = Path(temporary) / "browser-composed.docx"
                    composition_download.value.save_as(str(composed_path))
                    composed = Document(str(composed_path))
                    composed_text = "\n".join(paragraph.text for paragraph in composed.paragraphs)
                    for expected in (
                        "Agent 论文排版工具研究",
                        "本文验证本地 Web",
                        "第一章 绪论",
                    ):
                        if expected not in composed_text:
                            raise AssertionError(f"模板成品缺少内容: {expected}")
                    page.get_by_role("button", name="新建任务").click()
                    page.get_by_text("最近任务", exact=True).wait_for()
                    if page.locator(".job-history-list article").count() < 2:
                        raise AssertionError("任务历史未保留已完成任务")
                    page.locator(".job-history-list .status.completed").nth(1).wait_for()
                    page.once("dialog", lambda dialog: dialog.accept())
                    page.locator(".job-history-delete").last.click()
                    page.locator(".job-history-list article").nth(1).wait_for(
                        state="detached", timeout=15_000
                    )
                    if console_errors:
                        raise AssertionError(f"browser console errors: {console_errors}")
                    page.screenshot(path="/tmp/paper-setting-browser-smoke.png", full_page=True)
                    print(
                        {
                            "status": "completed",
                            "format_artifact_count": format_artifact_count,
                            "composition": "completed",
                            "history_count_after_cleanup": page.locator(
                                ".job-history-list article"
                            ).count(),
                            "screenshot": "/tmp/paper-setting-browser-smoke.png",
                        }
                    )
                    browser.close()
            finally:
                worker.terminate()
                try:
                    worker.wait(timeout=10)
                except subprocess.TimeoutExpired:
                    worker.kill()
                    worker.wait(timeout=5)
            if worker.returncode not in {0, -15}:
                raise RuntimeError(worker_log.read_text(encoding="utf-8"))


if __name__ == "__main__":
    main()
