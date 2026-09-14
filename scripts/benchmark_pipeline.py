from __future__ import annotations

import time
from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document
from paper_setting_core.documents import inspect_document
from paper_setting_core.formatting import apply_plan
from paper_setting_core.planning import generate_plan
from paper_setting_core.rulepacks import default_rule_pack
from paper_setting_core.verification import validate_result


def main() -> None:
    with TemporaryDirectory() as temporary:
        root = Path(temporary)
        source = root / "benchmark.docx"
        output = root / "formatted.docx"
        document = Document()
        document.add_paragraph("大规模论文排版性能测试", style="Title")
        for chapter in range(1, 11):
            document.add_paragraph(f"第{chapter}章 测试章节", style="Heading 1")
            for section in range(1, 11):
                document.add_paragraph(f"{chapter}.{section} 测试小节", style="Heading 2")
                for paragraph in range(25):
                    document.add_paragraph(
                        f"这是第 {chapter} 章第 {section} 节的第 {paragraph} 个性能测试段落。"
                        "内容用于模拟中文学位论文正文，并验证排版计划和完整性检查性能。"
                    )
            document.add_page_break()
        document.save(source)

        started = time.perf_counter()
        inspection = inspect_document(source)
        plan = generate_plan(inspection, default_rule_pack())
        planning_seconds = time.perf_counter() - started
        approved = {
            operation.operation_id
            for operation in plan.operations
            if operation.status == "proposed"
        }
        started = time.perf_counter()
        operations = apply_plan(source, output, plan, approved)
        report = validate_result(source, output, operations)
        apply_seconds = time.perf_counter() - started
        print(
            {
                "paragraphs": len(inspection.items),
                "plan_seconds": round(planning_seconds, 3),
                "apply_validate_seconds": round(apply_seconds, 3),
                "integrity_ok": report.integrity_ok,
            }
        )
        if planning_seconds >= 15 or apply_seconds >= 30 or not report.integrity_ok:
            raise SystemExit(1)


if __name__ == "__main__":
    main()

