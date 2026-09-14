"""Academic OOXML Post-Processing Sanitizer.

Normalizes Word OOXML nodes for 100% academic thesis compliance:
1. Eliminates unintended italics and unifies text colors to pure black (RGB 000000).
2. Clears stray single/double/wavy underlines and removes proofErr squiggles.
3. Preserves original author header/footer distances.
4. Clears highlight and run/paragraph shading; removes strikethroughs.
5. Smart semantic correction of false superscripts/subscripts in unbroken sentences.
6. Enforces unified paragraph indents (firstLine 2 chars for body, hanging 2 chars for references).
7. Enforces unified paragraph spacing per semantic role.
8. Prevents orphan headings (w:keepNext, w:keepLines) and enables widow/orphan control (w:widowControl).
9. Academic Three-Line Table (三线表): top/bottom 1.5pt, header dividing line 0.75pt, centered, tblHeader, cantSplit.
10. Caption formatting: center alignment, proper before/after spacing, keepNext on table captions.
11. Dual font normalization at run level (w:rFonts): 宋体/黑体 + Times New Roman.
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

CITATION_REGEX = re.compile(r"^(\[|〔|\()?\s*\d+(?:[\s,\-–—~，、]\s*\d+)*\s*(\]|〕|\))?$")
FOOTNOTE_MARK_REGEX = re.compile(r"^[①②③④⑤⑥⑦⑧⑨⑩*†‡#※]$")
MATH_SCI_SUBSCRIPTS = {"max", "min", "avg", "ref", "eff", "in", "out", "opt", "sub", "sup", "tot", "obs", "pred"}

HEADING_ROLES = {
    "heading_1",
    "heading_2",
    "heading_3",
    "heading_4",
    "paper_title",
    "abstract_heading",
    "bibliography_heading",
    "toc_heading",
    "appendix_heading",
    "acknowledgments_heading",
}


def is_legitimate_vertical_align(text: str) -> bool:
    """Determine if a superscript or subscript run is academically legitimate (citation/math)."""
    s = text.strip()
    if not s:
        return False
    if CITATION_REGEX.match(s):
        return True
    if FOOTNOTE_MARK_REGEX.match(s):
        return True
    if text != s:
        return False
    if any("\u4e00" <= ch <= "\u9fa5" for ch in s):
        return False
    if re.search(r"[。，、：；！？“”\"\'—（）\(\)]", s):
        return False
    if re.match(r"^[\+\-±]?\d+(?:\.\d+)?$", s) and len(s) <= 4:
        return True
    if len(s) == 1 and s.isalpha():
        return True
    if s.lower() in MATH_SCI_SUBSCRIPTS:
        return True
    return False


def sanitize_formatted_docx(
    output_path: Path,
    original_distances: list[tuple[Any, Any]] | None = None,
    inspection: Any = None,
    standard_id: str = "gb-t-7713-1",
) -> None:
    """Execute deep academic OOXML sanitization and normalization."""
    doc = Document(str(output_path))

    # 1. Restore original header and footer distances to prevent unnecessary distortion
    if original_distances:
        for idx, sec in enumerate(doc.sections):
            if idx < len(original_distances):
                h_dist, f_dist = original_distances[idx]
                if h_dist is not None:
                    sec.header_distance = h_dist
                if f_dist is not None:
                    sec.footer_distance = f_dist

    # 2. Extract paragraph semantic roles from inspection if available
    body_roles = []
    if inspection is not None and hasattr(inspection, "items"):
        body_roles = [item.semantic_role for item in inspection.items if getattr(item, "story", "") == "body"]

    # 3. Process every paragraph in the document body
    for p_idx, p in enumerate(doc.paragraphs):
        role = body_roles[p_idx] if p_idx < len(body_roles) else "body"
        ppr = p._p.get_or_add_pPr()

        # Clear paragraph-level shading
        for shd in ppr.findall(qn("w:shd")):
            ppr.remove(shd)

        # 3.1 Orphan heading prevention (w:keepNext, w:keepLines) & Widow control
        if role in HEADING_ROLES or role in ("table_caption", "table_caption_en"):
            if ppr.find(qn("w:keepNext")) is None:
                ppr.append(OxmlElement("w:keepNext"))
            if ppr.find(qn("w:keepLines")) is None:
                ppr.append(OxmlElement("w:keepLines"))
        elif role in ("body", "abstract_body"):
            if ppr.find(qn("w:widowControl")) is None:
                ppr.append(OxmlElement("w:widowControl"))

        # 3.2 Alignment
        jc = ppr.find(qn("w:jc"))
        if jc is None:
            jc = OxmlElement("w:jc")
            ppr.append(jc)
        if role in (
            "paper_title",
            "abstract_heading",
            "bibliography_heading",
            "toc_heading",
            "appendix_heading",
            "acknowledgments_heading",
            "figure_caption",
            "figure_caption_en",
            "table_caption",
            "table_caption_en",
        ):
            jc.set(qn("w:val"), "center")
        elif role == "heading_1":
            jc.set(qn("w:val"), "left" if standard_id in ("cass-humanities", "ieee-style") else "center")
        elif role in ("heading_2", "heading_3", "heading_4"):
            jc.set(qn("w:val"), "left")
        elif role in ("body", "abstract_body", "reference_entry"):
            jc.set(qn("w:val"), "both")

        # 3.3 Unified paragraph indents
        ind = ppr.find(qn("w:ind"))
        if ind is None:
            ind = OxmlElement("w:ind")
            ppr.append(ind)
        # Clear erratic left/right indentation
        ind.attrib.pop(qn("w:left"), None)
        ind.attrib.pop(qn("w:right"), None)
        ind.attrib.pop(qn("w:leftChars"), None)
        ind.attrib.pop(qn("w:rightChars"), None)

        if role in ("body", "abstract_body"):
            ind.attrib.pop(qn("w:hanging"), None)
            ind.attrib.pop(qn("w:hangingChars"), None)
            ind.set(qn("w:firstLineChars"), "200")
            ind.set(qn("w:firstLine"), "480")
        elif role == "reference_entry":
            ind.attrib.pop(qn("w:firstLine"), None)
            ind.attrib.pop(qn("w:firstLineChars"), None)
            ind.set(qn("w:left"), "420")
            ind.set(qn("w:hanging"), "420")
        else:
            # headings, titles, captions, keywords
            ind.attrib.pop(qn("w:firstLine"), None)
            ind.attrib.pop(qn("w:firstLineChars"), None)
            ind.attrib.pop(qn("w:hanging"), None)
            ind.attrib.pop(qn("w:hangingChars"), None)

        # 3.4 Unified paragraph spacing
        spacing = ppr.find(qn("w:spacing"))
        if spacing is None:
            spacing = OxmlElement("w:spacing")
            ppr.append(spacing)
        spacing.attrib.pop(qn("w:beforeLines"), None)
        spacing.attrib.pop(qn("w:afterLines"), None)

        if role in ("body", "abstract_body"):
            spacing.set(qn("w:before"), "0")
            spacing.set(qn("w:after"), "0")
            line_val = "360"
            if standard_id == "cass-humanities":
                line_val = "324"  # 1.35x
            elif standard_id == "ieee-style":
                line_val = "276"  # 1.15x
            elif standard_id == "ucas-thesis":
                line_val = "336"  # 1.4x
            spacing.set(qn("w:line"), line_val)
            spacing.set(qn("w:lineRule"), "auto")
        elif role == "heading_1":
            if standard_id == "ucas-thesis":
                spacing.set(qn("w:before"), "360")  # 18pt
                spacing.set(qn("w:after"), "240")   # 12pt
            else:
                spacing.set(qn("w:before"), "240")  # 12pt
                spacing.set(qn("w:after"), "120")   # 6pt
            spacing.set(qn("w:line"), "360")
            spacing.set(qn("w:lineRule"), "auto")
        elif role == "heading_2":
            spacing.set(qn("w:before"), "180")
            spacing.set(qn("w:after"), "60")
            spacing.set(qn("w:line"), "360")
            spacing.set(qn("w:lineRule"), "auto")
        elif role == "heading_3":
            spacing.set(qn("w:before"), "120")
            spacing.set(qn("w:after"), "0")
            spacing.set(qn("w:line"), "360")
            spacing.set(qn("w:lineRule"), "auto")
        elif role in ("paper_title", "abstract_heading", "bibliography_heading"):
            spacing.set(qn("w:before"), "240")
            spacing.set(qn("w:after"), "120")
        elif role in ("table_caption", "table_caption_en"):
            spacing.set(qn("w:before"), "120")  # 6pt before table caption
            spacing.set(qn("w:after"), "60")    # 3pt after (close to table)
            spacing.set(qn("w:line"), "300")
            spacing.set(qn("w:lineRule"), "auto")
        elif role in ("figure_caption", "figure_caption_en"):
            spacing.set(qn("w:before"), "60")   # 3pt before (close to figure)
            spacing.set(qn("w:after"), "120")   # 6pt after
            spacing.set(qn("w:line"), "300")
            spacing.set(qn("w:lineRule"), "auto")
        elif role == "keywords":
            spacing.set(qn("w:before"), "60")
            spacing.set(qn("w:after"), "60")
            spacing.set(qn("w:line"), "360")
            spacing.set(qn("w:lineRule"), "auto")
        elif role == "reference_entry":
            spacing.set(qn("w:before"), "0")
            spacing.set(qn("w:after"), "0")
            spacing.set(qn("w:line"), "240")
            spacing.set(qn("w:lineRule"), "auto")

        # 3.5 Process each run in the paragraph
        for r in p.runs:
            rpr = r._r.get_or_add_rPr()

            # Clear highlights and run shading
            for h in rpr.findall(qn("w:highlight")):
                rpr.remove(h)
            for s in rpr.findall(qn("w:shd")):
                rpr.remove(s)

            # Clear all strikethrough
            for st in rpr.findall(qn("w:strike")):
                rpr.remove(st)
            for dst in rpr.findall(qn("w:dstrike")):
                rpr.remove(dst)

            # Clear Italics: force val="0"
            for t in rpr.findall(qn("w:i")):
                rpr.remove(t)
            for t in rpr.findall(qn("w:iCs")):
                rpr.remove(t)
            i_elem = OxmlElement("w:i")
            i_elem.set(qn("w:val"), "0")
            rpr.append(i_elem)
            ics_elem = OxmlElement("w:iCs")
            ics_elem.set(qn("w:val"), "0")
            rpr.append(ics_elem)

            # Unify Color to Black: enforce 000000
            for t in rpr.findall(qn("w:color")):
                rpr.remove(t)
            c_elem = OxmlElement("w:color")
            c_elem.set(qn("w:val"), "000000")
            rpr.append(c_elem)

            # Clear Underlines
            for t in rpr.findall(qn("w:u")):
                rpr.remove(t)
            u_elem = OxmlElement("w:u")
            u_elem.set(qn("w:val"), "none")
            rpr.append(u_elem)

            # Intelligent semantic check for vertAlign (superscript/subscript)
            va = rpr.find(qn("w:vertAlign"))
            if va is not None:
                if is_legitimate_vertical_align(r.text):
                    # Citations in academic writing must always be superscript, not subscript
                    if CITATION_REGEX.match(r.text.strip()):
                        va.set(qn("w:val"), "superscript")
                else:
                    rpr.remove(va)
                    target_sz = "21" if role == "reference_entry" else "24"
                    for sz in rpr.findall(qn("w:sz")):
                        sz.set(qn("w:val"), target_sz)
                    for szcs in rpr.findall(qn("w:szCs")):
                        szcs.set(qn("w:val"), target_sz)

            # Dual font normalization: East Asian + Western Latin (rFonts)
            rfonts = rpr.find(qn("w:rFonts"))
            if rfonts is None:
                rfonts = OxmlElement("w:rFonts")
                rpr.insert(0, rfonts)

            if role in HEADING_ROLES:
                rfonts.set(qn("w:eastAsia"), "黑体")
                rfonts.set(qn("w:ascii"), "Times New Roman")
                rfonts.set(qn("w:hAnsi"), "Times New Roman")
                rfonts.set(qn("w:cs"), "Times New Roman")
                if role == "heading_1":
                    sz_val = "36" if standard_id == "ucas-thesis" else "32"
                    for sz in rpr.findall(qn("w:sz")):
                        sz.set(qn("w:val"), sz_val)
                    for szcs in rpr.findall(qn("w:szCs")):
                        szcs.set(qn("w:val"), sz_val)
            elif role in ("table_caption", "table_caption_en", "figure_caption", "figure_caption_en"):
                rfonts.set(qn("w:eastAsia"), "宋体")
                rfonts.set(qn("w:ascii"), "Times New Roman")
                rfonts.set(qn("w:hAnsi"), "Times New Roman")
                rfonts.set(qn("w:cs"), "Times New Roman")
                for sz in rpr.findall(qn("w:sz")):
                    sz.set(qn("w:val"), "21")
                for szcs in rpr.findall(qn("w:szCs")):
                    szcs.set(qn("w:val"), "21")
            else:
                target_east_asia = "仿宋" if standard_id == "cass-humanities" else "宋体"
                rfonts.set(qn("w:eastAsia"), target_east_asia)
                rfonts.set(qn("w:ascii"), "Times New Roman")
                rfonts.set(qn("w:hAnsi"), "Times New Roman")
                rfonts.set(qn("w:cs"), "Times New Roman")

            # Inject w:noProof to suppress Word spell/grammar squiggles
            for np in rpr.findall(qn("w:noProof")):
                rpr.remove(np)
            rpr.append(OxmlElement("w:noProof"))

    # 4. Standard Academic Three-Line Table (三线表)
    for table in doc.tables:
        tbl_pr = table._tbl.tblPr
        tbl_jc = tbl_pr.find(qn("w:jc"))
        if tbl_jc is None:
            tbl_jc = OxmlElement("w:jc")
            tbl_pr.append(tbl_jc)
        tbl_jc.set(qn("w:val"), "center")

        # Rebuild three-line table borders
        tbl_borders = tbl_pr.find(qn("w:tblBorders"))
        if tbl_borders is None:
            tbl_borders = OxmlElement("w:tblBorders")
            tbl_pr.append(tbl_borders)
        else:
            tbl_borders.clear()

        top_border = OxmlElement("w:top")
        top_border.set(qn("w:val"), "single")
        top_border.set(qn("w:sz"), "12")  # 1.5 pt
        top_border.set(qn("w:space"), "0")
        top_border.set(qn("w:color"), "000000")
        tbl_borders.append(top_border)

        bottom_border = OxmlElement("w:bottom")
        bottom_border.set(qn("w:val"), "single")
        bottom_border.set(qn("w:sz"), "12")  # 1.5 pt
        bottom_border.set(qn("w:space"), "0")
        bottom_border.set(qn("w:color"), "000000")
        tbl_borders.append(bottom_border)

        for b_name in ("left", "right", "insideH", "insideV"):
            b_elem = OxmlElement(f"w:{b_name}")
            b_elem.set(qn("w:val"), "none")
            tbl_borders.append(b_elem)

        for r_idx, row in enumerate(table.rows):
            tr_pr = row._tr.get_or_add_trPr()
            # CantSplit: prevent row breaking mid-cell across pages
            if tr_pr.find(qn("w:cantSplit")) is None:
                tr_pr.append(OxmlElement("w:cantSplit"))

            # Header row repeat (tblHeader) and dividing line (0.75 pt)
            if r_idx == 0:
                if tr_pr.find(qn("w:tblHeader")) is None:
                    tr_pr.append(OxmlElement("w:tblHeader"))
                for cell in row.cells:
                    tc_pr = cell._tc.get_or_add_tcPr()
                    tc_borders = tc_pr.find(qn("w:tcBorders"))
                    if tc_borders is None:
                        tc_borders = OxmlElement("w:tcBorders")
                        tc_pr.append(tc_borders)
                    b_elem = tc_borders.find(qn("w:bottom"))
                    if b_elem is None:
                        b_elem = OxmlElement("w:bottom")
                        tc_borders.append(b_elem)
                    b_elem.set(qn("w:val"), "single")
                    b_elem.set(qn("w:sz"), "6")  # 0.75 pt header line
                    b_elem.set(qn("w:space"), "0")
                    b_elem.set(qn("w:color"), "000000")

            for cell in row.cells:
                tc_pr = cell._tc.get_or_add_tcPr()
                # Vertical center
                if tc_pr.find(qn("w:vAlign")) is None:
                    v_align = OxmlElement("w:vAlign")
                    v_align.set(qn("w:val"), "center")
                    tc_pr.append(v_align)

                for cp in cell.paragraphs:
                    cp_pr = cp._p.get_or_add_pPr()
                    # Remove cell paragraph indents
                    c_ind = cp_pr.find(qn("w:ind"))
                    if c_ind is not None:
                        c_ind.attrib.pop(qn("w:firstLine"), None)
                        c_ind.attrib.pop(qn("w:firstLineChars"), None)
                        c_ind.attrib.pop(qn("w:left"), None)
                        c_ind.attrib.pop(qn("w:right"), None)
                    # Single spacing
                    c_sp = cp_pr.find(qn("w:spacing"))
                    if c_sp is None:
                        c_sp = OxmlElement("w:spacing")
                        cp_pr.append(c_sp)
                    c_sp.set(qn("w:before"), "0")
                    c_sp.set(qn("w:after"), "0")
                    c_sp.set(qn("w:line"), "240")
                    c_sp.set(qn("w:lineRule"), "auto")

                    for r in cp.runs:
                        rpr = r._r.get_or_add_rPr()
                        for h in rpr.findall(qn("w:highlight")):
                            rpr.remove(h)
                        for s in rpr.findall(qn("w:shd")):
                            rpr.remove(s)
                        for st in rpr.findall(qn("w:strike")):
                            rpr.remove(st)
                        for dst in rpr.findall(qn("w:dstrike")):
                            rpr.remove(dst)
                        for u in rpr.findall(qn("w:u")):
                            rpr.remove(u)
                        u_elem = OxmlElement("w:u")
                        u_elem.set(qn("w:val"), "none")
                        rpr.append(u_elem)
                        for t in rpr.findall(qn("w:color")):
                            rpr.remove(t)
                        c_elem = OxmlElement("w:color")
                        c_elem.set(qn("w:val"), "000000")
                        rpr.append(c_elem)
                        for t in rpr.findall(qn("w:i")):
                            rpr.remove(t)
                        for t in rpr.findall(qn("w:iCs")):
                            rpr.remove(t)
                        # Font: 10.5pt (五号)
                        for sz in rpr.findall(qn("w:sz")):
                            sz.set(qn("w:val"), "21")
                        for szcs in rpr.findall(qn("w:szCs")):
                            szcs.set(qn("w:val"), "21")
                        # Header row bold
                        if r_idx == 0:
                            if rpr.find(qn("w:b")) is None:
                                rpr.append(OxmlElement("w:b"))
                        # Fonts: 宋体 / Times New Roman
                        rfonts = rpr.find(qn("w:rFonts"))
                        if rfonts is None:
                            rfonts = OxmlElement("w:rFonts")
                            rpr.insert(0, rfonts)
                        rfonts.set(qn("w:eastAsia"), "宋体")
                        rfonts.set(qn("w:ascii"), "Times New Roman")
                        rfonts.set(qn("w:hAnsi"), "Times New Roman")
                        rfonts.set(qn("w:cs"), "Times New Roman")
                        if rpr.find(qn("w:noProof")) is None:
                            rpr.append(OxmlElement("w:noProof"))

    # 5. Remove all proofErr markers (Word red/green wavy squiggles)
    for err in doc._element.xpath(".//w:proofErr"):
        err.getparent().remove(err)

    # 6. Word settings level: hide spelling and grammar errors document-wide
    try:
        settings_elm = doc.settings.element
        if settings_elm.find(qn("w:hideSpellingErrors")) is None:
            settings_elm.append(OxmlElement("w:hideSpellingErrors"))
        if settings_elm.find(qn("w:hideGrammaticalErrors")) is None:
            settings_elm.append(OxmlElement("w:hideGrammaticalErrors"))
        if settings_elm.find(qn("w:doNotCheckSpelling")) is None:
            settings_elm.append(OxmlElement("w:doNotCheckSpelling"))
        proof_state = settings_elm.find(qn("w:proofState"))
        if proof_state is None:
            proof_state = OxmlElement("w:proofState")
            settings_elm.append(proof_state)
        proof_state.set(qn("w:spelling"), "clean")
        proof_state.set(qn("w:grammar"), "clean")
    except Exception:
        pass

    doc.save(str(output_path))
