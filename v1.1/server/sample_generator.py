"""Generator for realistic unformatted thesis DOCX samples."""

from __future__ import annotations

from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


def generate_sample_thesis(output_path: Path) -> Path:
    """Generate a realistic thesis draft DOCX with typical raw formatting flaws."""
    doc = Document()

    # 1. Title with improper font and alignment
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run_title = p_title.add_run("基于深度图神经网络的学术论文智能排版与语义一致性推断系统研究")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(18)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    # Author
    p_author = doc.add_paragraph("张三李四   指导教师：王教授")
    p_author.paragraph_format.space_after = Pt(20)

    # 2. Abstract
    p_abs_h = doc.add_paragraph("摘  要")
    p_abs_h.paragraph_format.space_before = Pt(10)
    p_abs_h.paragraph_format.space_after = Pt(6)

    p_abs_body = doc.add_paragraph(
        "随着学术科研成果的爆发式增长，高校学位论文与学术期刊排版在格式规范性、图表交叉引用及学术隐私保护等方面面临着前所未有的挑战。"
        "传统的人工 Word 排版操作繁琐且极易出错，而直接将未发表的学术论文上传至商业大模型云端又存在极大的查重污染与论文泄露风险。"
        "针对上述痛点，本文提出了一种纯本地离线单机运行的学术论文智能排版与图结构重构框架。通过建立确定性 OpenXML AST 语法树解析管线，"
        "实现秒级自动对齐国标 GB/T 7713.1 规范，并在严格保护学术隐私的前提下提供专家级深度结构攻坚能力。"
    )
    # Notice: NO first line indent, small line spacing
    p_abs_body.paragraph_format.first_line_indent = Pt(0)
    p_abs_body.paragraph_format.line_spacing = 1.0

    p_kw = doc.add_paragraph("关键词：学术论文排版；本地离线；GB/T 7713.1；OOXML 语义树；学术隐私")
    p_kw.paragraph_format.space_after = Pt(16)

    # 3. Heading 1
    h1_1 = doc.add_paragraph("1  绪论")
    h1_1.runs[0].font.size = Pt(15)

    p_body_1 = doc.add_paragraph(
        "学术论文是衡量科研产出与高校人才培养质量的重要载体[1]。然而，长期以来，研究生在撰写毕业论文时，"
        "往往需要花费大量时间处理复杂的版芯边距、标题多级层级、宋体与 Times New Roman 混排等格式要求[2]。"
        "若排版不符合规范，盲审时常遭到格式扣分甚至延期答辩。"
    )
    p_body_1.paragraph_format.first_line_indent = Pt(0)

    # 4. Heading 2
    h2_1 = doc.add_paragraph("1.1 国内外研究现状")
    h2_1.runs[0].font.size = Pt(13)

    p_body_2 = doc.add_paragraph(
        "在现有文献中，学术文献结构化处理主要包括 LaTeX 编译与基于 Office 的 VBA 宏模板[3]。"
        "然而，绝大多数高校强制要求提交标准 DOCX 终稿，而现有的云端排版工具大多需要用户上传全文，这在涉密项目与高水平论文投稿中被严格禁止。"
        "如表 1 所示，主流排版方式在隐私安全与易用性上存在显著差距。"
    )

    # Table
    table = doc.add_table(rows=3, cols=3)
    table.style = "Table Grid"
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "排版方案"
    hdr_cells[1].text = "隐私保护等级"
    hdr_cells[2].text = "排版耗时与体验"
    
    row1_cells = table.rows[1].cells
    row1_cells[0].text = "纯手工调整"
    row1_cells[1].text = "高 (单机)"
    row1_cells[2].text = "低效 (平均 6~12 小时)"

    row2_cells = table.rows[2].cells
    row2_cells[0].text = "Paper Setting 本地引擎"
    row2_cells[1].text = "绝对安全 (100% 离线)"
    row2_cells[2].text = "极速秒级完成 (<3 秒)"

    p_tbl_cap = doc.add_paragraph("表 1 主流论文排版模式综合特性对比")
    p_tbl_cap.paragraph_format.space_before = Pt(4)
    p_tbl_cap.paragraph_format.space_after = Pt(10)

    # 5. Heading 1 Chapter 2
    h1_2 = doc.add_paragraph("2  双视图系统架构与排版算法")
    
    p_body_3 = doc.add_paragraph(
        "为了兼顾 90% 的普通格式对齐诉求与 10% 的疑难结构攻坚，本系统设计了独特的双视图有限状态机体系。"
        "在默认视图下，纯确定性规则引擎以零 Token 消耗、零网络请求执行 4 步流排版；"
        "当用户切换至详细模式时，系统支持正文引注转换为页下真脚注，并将孤立编号升级为官方 REF/SEQ 域代码。"
    )

    # 6. References
    p_ref_h = doc.add_paragraph("参考文献")
    p_ref_h.paragraph_format.space_before = Pt(14)
    p_ref_h.paragraph_format.space_after = Pt(6)

    doc.add_paragraph("[1] 陈某某, 李某. 学位论文格式规范化对科研评价的影响研究[J]. 中国高等教育研究, 2021, 38(4): 45-52.")
    doc.add_paragraph("[2] 王某. 基于 Office OpenXML 的科技文档自动化处理技术[D]. 北京: 清华大学, 2022.")
    doc.add_paragraph("[3] Smith J, Doe A. Deep structure parsing of scientific manuscripts[J]. Journal of Computer Science, 2023, 19(2): 112-125.")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    return output_path
