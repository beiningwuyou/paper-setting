"""Test 100% offline minimal 4-step wizard compilation."""

from __future__ import annotations

from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from server.sample_generator import generate_sample_thesis
from server.engine import format_document_minimal, inspect_thesis


def test_sample_generation_and_inspection(tmp_path: Path):
    source_docx = tmp_path / "raw_sample.docx"
    generate_sample_thesis(source_docx)

    assert source_docx.exists()
    info = inspect_thesis(source_docx)
    assert info["total_paragraphs"] > 5
    assert info["headings_count"] >= 1
    assert info["tables_count"] >= 1


def test_minimal_flow_formatting_execution(tmp_path: Path):
    source_docx = tmp_path / "raw_sample.docx"
    output_docx = tmp_path / "formatted.docx"
    generate_sample_thesis(source_docx)

    result = format_document_minimal(
        source_path=source_docx,
        output_path=output_docx,
        standard_id="gb-t-7713-1",
    )

    assert result["status"] == "success"
    assert result["mode"] == "minimalist"
    assert result["operations_applied"] > 0
    assert result["compliance_score"] == 100
    assert len(result["highlights"]) >= 5
    assert len(result["diff_previews"]) > 0
    assert output_docx.exists()
    assert "xray_data" in result
    assert len(result["xray_data"]) > 0
    first_xray = result["xray_data"][0]
    assert "text" in first_xray
    assert "before" in first_xray and "after" in first_xray
    assert "font" in first_xray["before"] and "font" in first_xray["after"]
    assert "size_pt" in first_xray["before"] and "size_pt" in first_xray["after"]
    assert "align" in first_xray["before"] and "align" in first_xray["after"]
    assert "indent_pt" in first_xray["before"] and "indent_pt" in first_xray["after"]

    # Validate output DOCX is readable
    formatted_doc = Document(str(output_docx))
    assert len(formatted_doc.paragraphs) > 5


def test_six_issues_fixes(tmp_path: Path):
    """Regression test ensuring italics, non-black colors, underlines, and squiggles are removed, and header/footer distance is preserved."""
    from docx.shared import RGBColor, Pt, Mm
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    src_docx = tmp_path / "test_six_issues_src.docx"
    out_docx = tmp_path / "test_six_issues_out.docx"

    doc = Document()
    sec = doc.sections[0]
    sec.header_distance = Mm(6.0)
    sec.footer_distance = Mm(14.0)

    p1 = doc.add_paragraph()
    r1 = p1.add_run("这是带有斜体和红色字体的段落")
    r1.font.italic = True
    r1.font.color.rgb = RGBColor(255, 0, 0)

    p2 = doc.add_paragraph()
    r2 = p2.add_run("这是带有下划线和绿色字体的段落")
    r2.font.underline = True
    r2.font.color.rgb = RGBColor(0, 180, 0)

    # Add proofErr
    p2._p.append(OxmlElement("w:proofErr"))

    doc.save(str(src_docx))

    result = format_document_minimal(src_docx, out_docx, "gb-t-7713-1")
    assert result["status"] == "success"

    out_doc = Document(str(out_docx))

    # Check 1: Italics eliminated
    for rpr in out_doc._element.xpath(".//w:rPr"):
        i_elem = rpr.find(qn("w:i"))
        if i_elem is not None:
            assert i_elem.get(qn("w:val")) in ("0", "false", "off")
        ics_elem = rpr.find(qn("w:iCs"))
        if ics_elem is not None:
            assert ics_elem.get(qn("w:val")) in ("0", "false", "off")

    # Check 2: Colors unified to black
    for rpr in out_doc._element.xpath(".//w:rPr"):
        c = rpr.find(qn("w:color"))
        if c is not None:
            assert c.get(qn("w:val")) in ("000000", "auto")

    # Check 3: Underlines eliminated
    for rpr in out_doc._element.xpath(".//w:rPr"):
        u = rpr.find(qn("w:u"))
        if u is not None:
            assert u.get(qn("w:val")) in (None, "none")

    # Check 4: Proofing errors eliminated
    assert len(out_doc._element.xpath(".//w:proofErr")) == 0

    # Check 5: Header/footer distance preserved
    assert abs(out_doc.sections[0].header_distance.mm - 6.0) < 0.1
    assert abs(out_doc.sections[0].footer_distance.mm - 14.0) < 0.1


def test_seven_issues_chaotic_fixes(tmp_path: Path):
    """Test comprehensive resolution of all 7 issues on chaotic formatting."""
    src_docx = tmp_path / "chaotic_test.docx"
    out_docx = tmp_path / "chaotic_test_fixed.docx"

    doc = Document()
    p_title = doc.add_paragraph("测试论文标题")

    # Paragraph with highlights, strikes, and erroneous subscript
    p_body1 = doc.add_paragraph()
    r_hl = p_body1.add_run("这是带有高亮色块的文本")
    hl_elem = OxmlElement("w:highlight")
    hl_elem.set(qn("w:val"), "red")
    r_hl._r.get_or_add_rPr().append(hl_elem)

    r_strike = p_body1.add_run("这是带有删除线的文本")
    r_strike._r.get_or_add_rPr().append(OxmlElement("w:strike"))

    r_err_sub = p_body1.add_run("未断句长句中的错误下标文字")
    va_elem = OxmlElement("w:vertAlign")
    va_elem.set(qn("w:val"), "subscript")
    r_err_sub._r.get_or_add_rPr().append(va_elem)

    # Legitimate citation marker
    r_cit = p_body1.add_run("〔1〕")
    cit_va = OxmlElement("w:vertAlign")
    cit_va.set(qn("w:val"), "subscript")  # accidentally subscripted in draft
    r_cit._r.get_or_add_rPr().append(cit_va)

    # Set erratic indent and spacing
    p_body1_ind = p_body1._p.get_or_add_pPr().get_or_add_ind()
    p_body1_ind.set(qn("w:left"), "1080")
    p_body1_ind.set(qn("w:right"), "360")
    p_body1_sp = p_body1._p.get_or_add_pPr().get_or_add_spacing()
    p_body1_sp.set(qn("w:before"), "280")
    p_body1_sp.set(qn("w:after"), "180")

    # Add bibliography heading and entry
    p_ref_head = doc.add_paragraph("参考文献")
    p_ref1 = doc.add_paragraph("王飞跃,缪青海. 人工智能驱动的科学研究新范式：从AI4S到智能科学 [J]. 中国科学院院刊, 2023, 38 (04): 536-540.")
    p_ref1_ind = p_ref1._p.get_or_add_pPr().get_or_add_ind()
    p_ref1_ind.set(qn("w:left"), "0")
    p_ref1_ind.set(qn("w:firstLine"), "480")  # wrongly set to firstLine instead of hanging

    doc.save(str(src_docx))

    result = format_document_minimal(src_docx, out_docx, "gb-t-7713-1")
    assert result["status"] == "success"

    out_doc = Document(str(out_docx))

    # Issue 1: Highlights and shadings completely eliminated
    assert len(out_doc._element.xpath(".//w:rPr/w:highlight")) == 0
    assert len(out_doc._element.xpath(".//w:rPr/w:shd")) == 0
    assert len(out_doc._element.xpath(".//w:pPr/w:shd")) == 0

    # Issue 2: Underlines eliminated
    for u in out_doc._element.xpath(".//w:rPr/w:u"):
        assert u.attrib.get(qn("w:val")) in (None, "none")

    # Issue 3: Strikethroughs completely eliminated
    assert len(out_doc._element.xpath(".//w:rPr/w:strike")) == 0
    assert len(out_doc._element.xpath(".//w:rPr/w:dstrike")) == 0

    # Issue 4: Erroneous Chinese words subscript/superscript reset to baseline
    for p in out_doc.paragraphs:
        for r in p.runs:
            va = r._r.xpath(".//w:vertAlign")
            if va:
                # Citations like 〔1〕 are kept and normalized to superscript
                assert r.text == "〔1〕"
                assert va[0].attrib.get(qn("w:val")) == "superscript"

    # Issue 5 & 6: Body paragraph indent and spacing unified
    out_p_body = out_doc.paragraphs[1]
    ppr_body = out_p_body._p.find(qn("w:pPr"))
    assert ppr_body is not None
    ind = ppr_body.find(qn("w:ind"))
    assert ind is not None
    assert ind.attrib.get(qn("w:firstLineChars")) == "200"
    assert qn("w:left") not in ind.attrib
    assert qn("w:right") not in ind.attrib

    sp = ppr_body.find(qn("w:spacing"))
    assert sp is not None
    assert sp.attrib.get(qn("w:before")) == "0"
    assert sp.attrib.get(qn("w:after")) == "0"
    assert sp.attrib.get(qn("w:line")) == "360"

    # Issue 7: Reference entry has hanging indent and single spacing
    out_p_ref = out_doc.paragraphs[3]
    ppr_ref = out_p_ref._p.find(qn("w:pPr"))
    assert ppr_ref is not None
    ref_ind = ppr_ref.find(qn("w:ind"))
    assert ref_ind is not None
    assert ref_ind.attrib.get(qn("w:hanging")) == "420"
    assert ref_ind.attrib.get(qn("w:left")) == "420"
    assert qn("w:firstLine") not in ref_ind.attrib


def test_academic_advanced_rules_and_presets(tmp_path: Path):
    """Verify newly audited and implemented academic typesetting rules:
    1. Orphan heading prevention (w:keepNext, w:keepLines) & widowControl.
    2. Academic three-line table (w:tblBorders, top/bottom 1.5pt, header line 0.75pt, centered, tblHeader, cantSplit).
    3. Caption formatting: center alignment, table caption keepNext, proper spacing before/after.
    4. Dual font normalization (w:rFonts).
    5. Preset customization (ucas-thesis: 18pt heading 1, 1.4x line spacing, 28mm top margin).
    """
    src_docx = tmp_path / "academic_rules_src.docx"
    out_docx_gbt = tmp_path / "academic_rules_gbt.docx"
    out_docx_ucas = tmp_path / "academic_rules_ucas.docx"

    doc = Document()
    p_title = doc.add_paragraph("基于深度学习的学术排版优化研究")
    p_h1 = doc.add_paragraph("第1章 绪论")
    p_body = doc.add_paragraph("这是正文测试段落，用于验证段落孤行控制以及中西文字体规范设置。")
    p_tbl_cap = doc.add_paragraph("表 1-1 常见论文格式规范参数对比")

    # Add a table
    tbl = doc.add_table(rows=2, cols=2)
    tbl.rows[0].cells[0].text = "参数名称"
    tbl.rows[0].cells[1].text = "标准值"
    tbl.rows[1].cells[0].text = "行间距"
    tbl.rows[1].cells[1].text = "1.5倍"

    p_fig_cap = doc.add_paragraph("图 1-1 深度神经网络架构示意图")

    doc.save(str(src_docx))

    # 1. Format with default GB/T 7713.1
    res_gbt = format_document_minimal(src_docx, out_docx_gbt, "gb-t-7713-1")
    assert res_gbt["status"] == "success"

    doc_gbt = Document(str(out_docx_gbt))

    # Check 1: Headings have w:keepNext and w:keepLines
    h1_p = None
    for p in doc_gbt.paragraphs:
        if "第1章" in p.text:
            h1_p = p
            break
    assert h1_p is not None
    ppr_h1 = h1_p._p.find(qn("w:pPr"))
    assert ppr_h1.find(qn("w:keepNext")) is not None
    assert ppr_h1.find(qn("w:keepLines")) is not None
    assert ppr_h1.find(qn("w:jc")).attrib.get(qn("w:val")) == "center"

    # Check 2: Table caption has w:keepNext, centered, spacing before 6pt / after 3pt
    tbl_cap_p = None
    for p in doc_gbt.paragraphs:
        if "表 1-1" in p.text:
            tbl_cap_p = p
            break
    assert tbl_cap_p is not None
    ppr_tbl = tbl_cap_p._p.find(qn("w:pPr"))
    assert ppr_tbl.find(qn("w:keepNext")) is not None
    assert ppr_tbl.find(qn("w:keepLines")) is not None
    assert ppr_tbl.find(qn("w:jc")).attrib.get(qn("w:val")) == "center"
    tbl_cap_sp = ppr_tbl.find(qn("w:spacing"))
    assert tbl_cap_sp.attrib.get(qn("w:before")) == "120"  # 6pt
    assert tbl_cap_sp.attrib.get(qn("w:after")) == "60"    # 3pt

    # Check 3: Figure caption centered, spacing before 3pt / after 6pt
    fig_cap_p = None
    for p in doc_gbt.paragraphs:
        if "图 1-1" in p.text:
            fig_cap_p = p
            break
    assert fig_cap_p is not None
    ppr_fig = fig_cap_p._p.find(qn("w:pPr"))
    assert ppr_fig.find(qn("w:jc")).attrib.get(qn("w:val")) == "center"
    fig_cap_sp = ppr_fig.find(qn("w:spacing"))
    assert fig_cap_sp.attrib.get(qn("w:before")) == "60"   # 3pt
    assert fig_cap_sp.attrib.get(qn("w:after")) == "120"  # 6pt

    # Check 4: Body has widowControl and dual font
    body_p = None
    for p in doc_gbt.paragraphs:
        if "这是正文测试段落" in p.text:
            body_p = p
            break
    assert body_p is not None
    ppr_body = body_p._p.find(qn("w:pPr"))
    assert ppr_body.find(qn("w:widowControl")) is not None
    rfonts = body_p.runs[0]._r.find(qn("w:rPr")).find(qn("w:rFonts"))
    assert rfonts is not None
    assert rfonts.attrib.get(qn("w:eastAsia")) == "宋体"
    assert rfonts.attrib.get(qn("w:ascii")) == "Times New Roman"

    # Check 5: Three-line table
    assert len(doc_gbt.tables) == 1
    t = doc_gbt.tables[0]
    t_pr = t._tbl.tblPr
    assert t_pr.find(qn("w:jc")).attrib.get(qn("w:val")) == "center"
    t_borders = t_pr.find(qn("w:tblBorders"))
    assert t_borders is not None
    assert t_borders.find(qn("w:top")).attrib.get(qn("w:sz")) == "12"    # 1.5pt
    assert t_borders.find(qn("w:bottom")).attrib.get(qn("w:sz")) == "12" # 1.5pt
    assert t_borders.find(qn("w:left")).attrib.get(qn("w:val")) == "none"
    assert t_borders.find(qn("w:right")).attrib.get(qn("w:val")) == "none"
    # Row 0 has tblHeader and cantSplit
    assert t.rows[0]._tr.find(qn("w:trPr")).find(qn("w:tblHeader")) is not None
    assert t.rows[0]._tr.find(qn("w:trPr")).find(qn("w:cantSplit")) is not None
    # Row 0 cell bottom border is 0.75pt (sz=6)
    tc_b = t.rows[0].cells[0]._tc.find(qn("w:tcPr")).find(qn("w:tcBorders")).find(qn("w:bottom"))
    assert tc_b is not None
    assert tc_b.attrib.get(qn("w:sz")) == "6"

    # 2. Format with UCAS standard (Chinese Academy of Sciences thesis)
    res_ucas = format_document_minimal(src_docx, out_docx_ucas, "ucas-thesis")
    assert res_ucas["status"] == "success"

    doc_ucas = Document(str(out_docx_ucas))
    # Check margins: top 2.8cm (28mm)
    assert abs(doc_ucas.sections[0].top_margin.mm - 28.0) < 0.2

    # Check heading 1: 18pt (sz=36), space before 18pt (360), after 12pt (240)
    for p in doc_ucas.paragraphs:
        if "第1章" in p.text:
            sp = p._p.find(qn("w:pPr")).find(qn("w:spacing"))
            assert sp.attrib.get(qn("w:before")) == "360"
            assert sp.attrib.get(qn("w:after")) == "240"
            for r in p.runs:
                sz = r._r.find(qn("w:rPr")).find(qn("w:sz"))
                if sz is not None:
                    assert sz.attrib.get(qn("w:val")) == "36"

    # Check body: line spacing 1.4x (line=336)
    for p in doc_ucas.paragraphs:
        if "这是正文测试段落" in p.text:
            sp = p._p.find(qn("w:pPr")).find(qn("w:spacing"))
            assert sp.attrib.get(qn("w:line")) == "336"



