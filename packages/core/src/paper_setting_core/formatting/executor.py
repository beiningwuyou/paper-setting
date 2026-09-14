from __future__ import annotations

import os
from collections.abc import Iterator
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from docx.text.paragraph import Paragraph

from paper_setting_core.advanced_word.cross_references import apply_cross_references
from paper_setting_core.advanced_word.transforms import (
    apply_automatic_numbering,
    convert_notes_to_footnotes,
    normalize_citations,
    rebuild_toc,
    reformat_formulas,
)
from paper_setting_core.documents.inspection import _iter_body, inspect_document
from paper_setting_core.errors import CrossReferenceInvalidError, JobConflictError
from paper_setting_core.planning.models import PatchOperation, PatchPlan

ALIGNMENTS = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}


def _iter_all_paragraphs(document: Any) -> Iterator[Paragraph]:
    for _, paragraph, _ in _iter_body(document):
        yield paragraph


def _ensure_style(document: Any, operation: PatchOperation, original_style: Any = None) -> Any:
    style_name = operation.after["style_name"]
    try:
        style = document.styles[style_name]
    except KeyError:
        style = document.styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = (
            original_style if original_style is not None else document.styles["Normal"]
        )
    if original_style is not None:
        # Shared styles may also be used by rejected/compliant paragraphs.
        return style
    paragraph_values = operation.after.get("paragraph", {})
    character_values = operation.after.get("character", {})
    _apply_paragraph_format(style.paragraph_format, paragraph_values)
    _apply_font(style.font, style.element, character_values)
    return style


def _apply_paragraph_format(paragraph_format: Any, values: dict[str, Any]) -> None:
    if "alignment" in values:
        paragraph_format.alignment = ALIGNMENTS[values["alignment"]]
    if "first_line_indent_pt" in values:
        paragraph_format.first_line_indent = Pt(values["first_line_indent_pt"])
    if "left_indent_pt" in values:
        paragraph_format.left_indent = Pt(values["left_indent_pt"])
    if "right_indent_pt" in values:
        paragraph_format.right_indent = Pt(values["right_indent_pt"])
    if "space_before_pt" in values:
        paragraph_format.space_before = Pt(values["space_before_pt"])
    if "space_after_pt" in values:
        paragraph_format.space_after = Pt(values["space_after_pt"])
    spacing = paragraph_format._element.get_or_add_pPr().get_or_add_spacing()
    if "space_before_lines" in values:
        spacing.attrib.pop(qn("w:before"), None)
        spacing.set(qn("w:beforeLines"), str(round(values["space_before_lines"] * 100)))
    if "space_after_lines" in values:
        spacing.attrib.pop(qn("w:after"), None)
        spacing.set(qn("w:afterLines"), str(round(values["space_after_lines"] * 100)))
    if "line_spacing" in values:
        mode = values.get("line_spacing_mode", "multiple")
        if mode == "multiple":
            paragraph_format.line_spacing = values["line_spacing"]
            paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        elif mode == "exact":
            paragraph_format.line_spacing = Pt(values["line_spacing"])
            paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        elif mode == "at_least":
            paragraph_format.line_spacing = Pt(values["line_spacing"])
            paragraph_format.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
        else:
            raise ValueError(f"unsupported line spacing mode: {mode}")
    if "keep_with_next" in values:
        paragraph_format.keep_with_next = values["keep_with_next"]
    if "keep_together" in values:
        paragraph_format.keep_together = values["keep_together"]
    if "page_break_before" in values:
        paragraph_format.page_break_before = values["page_break_before"]


def _apply_font(font: Any, style_element: Any, values: dict[str, Any]) -> None:
    latin = values.get("latin_font")
    east_asia = values.get("east_asia_font")
    if latin:
        font.name = latin
    if values.get("size_pt") is not None:
        font.size = Pt(values["size_pt"])
    if values.get("bold") is not None:
        font.bold = values["bold"]
    if values.get("italic") is not None:
        font.italic = values["italic"]
    if values.get("underline") is not None:
        font.underline = values["underline"]
    if values.get("color"):
        font.color.rgb = RGBColor.from_string(values["color"].upper())
    rpr = style_element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    if east_asia:
        rfonts.set(qn("w:eastAsia"), east_asia)
    if latin:
        rfonts.set(qn("w:ascii"), latin)
        rfonts.set(qn("w:hAnsi"), latin)


def _run_is_protected(run: Any) -> bool:
    xml = run._r.xml
    return any(
        marker in xml
        for marker in (
            "<w:fldChar",
            "<w:instrText",
            "<w:drawing",
            "<w:object",
            "<m:oMath",
            "<w:pict",
        )
    )


def _apply_run_format(run: Any, values: dict[str, Any]) -> None:
    if _run_is_protected(run):
        return
    latin = values.get("latin_font")
    east_asia = values.get("east_asia_font")
    if latin:
        run.font.name = latin
    if values.get("size_pt") is not None:
        run.font.size = Pt(values["size_pt"])
    if values.get("color"):
        run.font.color.rgb = RGBColor.from_string(values["color"].upper())
    if values.get("bold") is not None:
        run.font.bold = values["bold"]
    if values.get("italic") is not None:
        run.font.italic = values["italic"]
    if values.get("underline") is not None:
        run.font.underline = values["underline"]
    rfonts = run._r.get_or_add_rPr().get_or_add_rFonts()
    if east_asia:
        rfonts.set(qn("w:eastAsia"), east_asia)
    if latin:
        rfonts.set(qn("w:ascii"), latin)
        rfonts.set(qn("w:hAnsi"), latin)


def _fit_plain_tables_to_width(document: Any, maximum_width_twips: int) -> int:
    fitted = 0
    for table in document.tables:
        table_xml = table._tbl
        if table_xml.xpath(".//w:tc//w:tbl | .//w:gridSpan | .//w:vMerge"):
            continue
        columns = table_xml.xpath("./w:tblGrid/w:gridCol")
        widths = [int(column.get(qn("w:w"), "0")) for column in columns]
        if not widths or any(width <= 0 for width in widths):
            continue
        if maximum_width_twips <= len(widths):
            continue
        current_width = sum(widths)
        if current_width <= maximum_width_twips:
            continue
        rows = [row.xpath("./w:tc") for row in table_xml.xpath("./w:tr")]
        if any(len(cells) != len(widths) for cells in rows):
            continue

        scale = maximum_width_twips / current_width
        scaled_widths = [max(1, round(width * scale)) for width in widths]
        scaled_widths[-1] += maximum_width_twips - sum(scaled_widths)
        for column, width in zip(columns, scaled_widths, strict=True):
            column.set(qn("w:w"), str(width))
        for cells in rows:
            for cell, width in zip(cells, scaled_widths, strict=True):
                cell_width = cell.find("./w:tcPr/w:tcW", namespaces=cell.nsmap)
                if cell_width is not None:
                    cell_width.set(qn("w:type"), "dxa")
                    cell_width.set(qn("w:w"), str(width))
        table_width = table_xml.tblPr.find(qn("w:tblW"))
        if table_width is None:
            table_width = OxmlElement("w:tblW")
            table_xml.tblPr.insert(0, table_width)
        table_width.set(qn("w:type"), "dxa")
        table_width.set(qn("w:w"), str(maximum_width_twips))
        table.autofit = False
        fitted += 1
    return fitted


def _apply_page_format(document: Any, values: dict[str, Any]) -> dict[str, int]:
    if set(values) != {
        "width_mm",
        "height_mm",
        "orientation",
        "margin_top_mm",
        "margin_bottom_mm",
        "margin_left_mm",
        "margin_right_mm",
        "header_distance_mm",
        "footer_distance_mm",
    }:
        attributes = {
            "width_mm": "page_width",
            "height_mm": "page_height",
            "margin_top_mm": "top_margin",
            "margin_bottom_mm": "bottom_margin",
            "margin_left_mm": "left_margin",
            "margin_right_mm": "right_margin",
            "header_distance_mm": "header_distance",
            "footer_distance_mm": "footer_distance",
        }
        for section in document.sections:
            for key, attribute in attributes.items():
                if key in values:
                    setattr(section, attribute, Mm(values[key]))
            if "orientation" in values:
                landscape = values["orientation"] == "landscape"
                section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
                width, height = section.page_width, section.page_height
                if width is not None and height is not None and (width > height) != landscape:
                    section.page_width, section.page_height = height, width
        return {"sections_updated": len(document.sections), "tables_fitted": 0}
    width = values["width_mm"]
    height = values["height_mm"]
    if values.get("orientation") == "landscape":
        width, height = height, width
    for section in document.sections:
        if values.get("orientation") == "landscape":
            section.orientation = WD_ORIENT.LANDSCAPE
        else:
            section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Mm(width)
        section.page_height = Mm(height)
        section.top_margin = Mm(values["margin_top_mm"])
        section.bottom_margin = Mm(values["margin_bottom_mm"])
        section.left_margin = Mm(values["margin_left_mm"])
        section.right_margin = Mm(values["margin_right_mm"])
        section.header_distance = Mm(values["header_distance_mm"])
        section.footer_distance = Mm(values["footer_distance_mm"])
    content_width_mm = width - values["margin_left_mm"] - values["margin_right_mm"]
    tables_fitted = _fit_plain_tables_to_width(document, int(Mm(content_width_mm).twips))
    return {"sections_updated": len(document.sections), "tables_fitted": tables_fitted}


def apply_plan(
    source: Path,
    output: Path,
    plan: PatchPlan,
    approved_operation_ids: set[str],
    *,
    confirmed_manual_operation_ids: set[str] | None = None,
) -> list[PatchOperation]:
    confirmed_manual_ids = confirmed_manual_operation_ids or set()
    inspection = inspect_document(source)
    if inspection.source_sha256 != plan.source_sha256:
        raise JobConflictError("源文件已变化，不能执行旧计划", "STALE_PLAN")
    document = Document(str(source))
    paragraphs = list(_iter_all_paragraphs(document))
    target_items = {item.stable_id: index for index, item in enumerate(inspection.items)}
    executed: list[PatchOperation] = []

    for operation in plan.operations:
        copied = operation.model_copy(deep=True)
        if operation.operation_id not in approved_operation_ids:
            copied.status = "skipped"
            copied.reasons = ["未获批准，因此未执行"]
            executed.append(copied)
            continue
        if (
            operation.status == "manual_review"
            and operation.operation_id not in confirmed_manual_ids
        ):
            copied.status = "skipped"
            copied.reasons.append("需要在工作台明确确认后才能执行")
            executed.append(copied)
            continue
        try:
            if operation.operation_type == "apply_page_format":
                copied.result = _apply_page_format(document, operation.after)
            elif operation.operation_type == "reformat_formulas":
                copied.result = reformat_formulas(document, operation.after)
            elif operation.operation_type == "normalize_citations":
                copied.result = normalize_citations(
                    document, inspection, paragraphs, operation.after
                )
            elif operation.operation_type == "rebuild_toc":
                copied.result = rebuild_toc(document, inspection, paragraphs, operation.after)
            elif operation.operation_type == "apply_automatic_numbering":
                copied.result = apply_automatic_numbering(
                    document, inspection, paragraphs, operation.after
                )
            elif operation.operation_type == "create_cross_references":
                copied.result = apply_cross_references(
                    document, inspection, paragraphs, operation.after
                )
            elif operation.operation_type == "convert_notes_to_footnotes":
                copied.result = convert_notes_to_footnotes(document, operation.after)
            else:
                item_index = target_items.get(operation.target_id)
                if item_index is None or item_index >= len(paragraphs):
                    raise JobConflictError("修改目标已变化", "STALE_PLAN")
                item = inspection.items[item_index]
                if item.exact_text_hash != operation.precondition_hash:
                    raise JobConflictError("修改目标文本已变化", "STALE_PLAN")
                paragraph = paragraphs[item_index]
                original_style = (
                    paragraph.style if operation.after.get("preserve_original_style") else None
                )
                if not operation.after.get("preserve_style_identity"):
                    style = _ensure_style(document, operation, original_style)
                    paragraph.style = style
                _apply_paragraph_format(paragraph.paragraph_format, operation.after["paragraph"])
                for run in paragraph.runs:
                    _apply_run_format(run, operation.after["character"])
                if operation.execution_scope == "preserve_protected_content":
                    copied.reasons.append("已保留受保护对象的内部 XML")
            copied.status = "applied"
        except (CrossReferenceInvalidError, JobConflictError):
            raise
        except Exception as exc:  # operation failure remains auditable
            copied.status = "failed"
            copied.reasons.append(type(exc).__name__)
        executed.append(copied)

    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = output.with_suffix(".tmp.docx")
    document.save(str(temporary))
    Document(str(temporary))
    os.replace(temporary, output)
    return executed
