from __future__ import annotations

import hashlib
import json
import re
import zipfile
from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Literal, cast

from docx import Document
from lxml import etree

from paper_setting_core.documents import sha256_file, validate_docx_package
from paper_setting_core.errors import (
    IntegrityCheckFailedError,
    StaleTemplateFillPlanError,
    TemplateFillBlockedError,
    TemplateFillInvalidError,
)
from paper_setting_core.templates.inspection import NS, PLACEHOLDER, W, _parse, _story_name
from paper_setting_core.templates.models import (
    TemplateFillBlockedTarget,
    TemplateFillOperation,
    TemplateFillPreview,
    TemplateFillResult,
)

XML = "http://www.w3.org/XML/1998/namespace"
M = "http://schemas.openxmlformats.org/officeDocument/2006/math"
EDITABLE_PART = re.compile(r"word/(?:document|header\d+|footer\d+)\.xml")
PROTECTED_ANCESTORS = {
    f"{{{M}}}oMath",
    f"{{{M}}}oMathPara",
    f"{{{W}}}del",
    f"{{{W}}}ins",
    f"{{{W}}}moveFrom",
    f"{{{W}}}moveTo",
    f"{{{W}}}fldSimple",
}
PROTECTED_CONTENT = (
    ".//w:fldChar | .//w:instrText | .//m:oMath | .//m:oMathPara | "
    ".//w:drawing | .//w:object | .//w:del | .//w:ins | .//w:moveFrom | .//w:moveTo"
)
PROTECTION_COUNTERS = {
    "fields": ".//w:fldChar | .//w:instrText | .//w:fldSimple",
    "formulas": ".//m:oMath | .//m:oMathPara",
    "drawings": ".//w:drawing | .//w:object",
    "revisions": ".//w:del | .//w:ins | .//w:moveFrom | .//w:moveTo",
}


def _xpath(root: etree._Element, expression: str) -> list[Any]:
    namespaces = {**NS, "m": M}
    return cast(list[Any], root.xpath(expression, namespaces=namespaces))


@dataclass
class _PlaceholderTarget:
    key: str
    token: str
    part_name: str
    story: str
    paragraph: etree._Element
    text_nodes: list[etree._Element]
    start: int
    end: int
    blocked_reason: str | None


@dataclass
class _ControlTarget:
    key: str
    part_name: str
    story: str
    control: etree._Element
    blocked_reason: str | None


def parse_template_values(payload: str) -> dict[str, str]:
    try:
        raw = json.loads(payload)
    except json.JSONDecodeError as exc:
        raise TemplateFillInvalidError("填写值必须是有效的 JSON 对象") from exc
    if not isinstance(raw, dict):
        raise TemplateFillInvalidError("填写值必须是由字段名和值组成的 JSON 对象")
    if len(raw) > 200:
        raise TemplateFillInvalidError("一次最多填写 200 个模板字段")
    values: dict[str, str] = {}
    field_errors: list[dict[str, str]] = []
    total_length = 0
    for key, value in raw.items():
        if (
            not isinstance(key, str)
            or not key
            or len(key) > 128
            or any(ord(character) < 32 for character in key)
        ):
            field_errors.append(
                {
                    "field": str(key),
                    "message": "字段名必须是 1–128 个可见字符",
                    "code": "invalid_key",
                }
            )
            continue
        if not isinstance(value, str):
            field_errors.append(
                {"field": key, "message": "字段值必须是字符串", "code": "invalid_type"}
            )
            continue
        if len(value) > 20_000:
            field_errors.append(
                {"field": key, "message": "单个字段值不能超过 20000 个字符", "code": "too_long"}
            )
            continue
        if any(ord(character) < 32 for character in value):
            field_errors.append(
                {
                    "field": key,
                    "message": "当前仅支持不含换行或控制字符的单行字段值",
                    "code": "control_character",
                }
            )
            continue
        total_length += len(value)
        values[key] = value
    if total_length > 1_000_000:
        field_errors.append(
            {
                "field": "values",
                "message": "字段值总长度不能超过 1000000 个字符",
                "code": "too_large",
            }
        )
    if field_errors:
        raise TemplateFillInvalidError("模板填写值校验失败", field_errors)
    return values


def _story_parts(archive: zipfile.ZipFile) -> dict[str, etree._Element]:
    return {
        name: _parse(archive.read(name))
        for name in sorted(archive.namelist())
        if EDITABLE_PART.fullmatch(name)
    }


def _nearest(node: etree._Element, tag: str) -> etree._Element | None:
    return next((ancestor for ancestor in node.iterancestors() if ancestor.tag == tag), None)


def _text_nodes(paragraph: etree._Element) -> list[etree._Element]:
    return [
        node
        for node in _xpath(paragraph, ".//w:t")
        if _nearest(node, f"{{{W}}}p") is paragraph
        and _nearest(node, f"{{{W}}}sdt") is None
    ]


def _placeholder_block_reason(nodes: list[etree._Element]) -> str | None:
    for node in nodes:
        for ancestor in node.iterancestors():
            if ancestor.tag == f"{{{W}}}p":
                break
            if ancestor.tag in PROTECTED_ANCESTORS:
                return "占位符位于字段、公式或修订区域内"
    return None


def _placeholder_targets(parts: dict[str, etree._Element]) -> list[_PlaceholderTarget]:
    targets: list[_PlaceholderTarget] = []
    for part_name, root in parts.items():
        for paragraph in _xpath(root, ".//w:p"):
            nodes = _text_nodes(paragraph)
            text = "".join(str(node.text or "") for node in nodes)
            if not text:
                continue
            spans: list[tuple[int, int, etree._Element]] = []
            cursor = 0
            for node in nodes:
                length = len(str(node.text or ""))
                spans.append((cursor, cursor + length, node))
                cursor += length
            for match in PLACEHOLDER.finditer(text):
                key = match.group("brace") or match.group("bracket") or match.group("angle")
                involved = [
                    node
                    for start, end, node in spans
                    if start < match.end() and end > match.start()
                ]
                targets.append(
                    _PlaceholderTarget(
                        key=key,
                        token=match.group("token"),
                        part_name=part_name,
                        story=_story_name(part_name, paragraph),
                        paragraph=paragraph,
                        text_nodes=nodes,
                        start=match.start(),
                        end=match.end(),
                        blocked_reason=_placeholder_block_reason(involved),
                    )
                )
    return targets


def _control_key(control: etree._Element, values: dict[str, str] | None = None) -> str | None:
    tags = _xpath(control, "./w:sdtPr/w:tag")
    aliases = _xpath(control, "./w:sdtPr/w:alias")
    tag = str(tags[0].get(f"{{{W}}}val")) if tags and tags[0].get(f"{{{W}}}val") else None
    alias = (
        str(aliases[0].get(f"{{{W}}}val"))
        if aliases and aliases[0].get(f"{{{W}}}val")
        else None
    )
    if values is not None:
        if tag in values:
            return tag
        if alias in values:
            return alias
    return tag or alias


def _control_block_reason(control: etree._Element) -> str | None:
    if _xpath(control, "./w:sdtPr/w:dataBinding"):
        return "内容控件使用数据绑定，直接写入可能被 Word 覆盖"
    if _xpath(control, PROTECTED_CONTENT):
        return "内容控件包含字段、公式、图形或修订对象"
    if _xpath(control, "./w:sdtContent//w:sdt"):
        return "内容控件包含嵌套内容控件"
    return None


def _control_targets(
    parts: dict[str, etree._Element], values: dict[str, str] | None = None
) -> list[_ControlTarget]:
    targets: list[_ControlTarget] = []
    for part_name, root in parts.items():
        for control in _xpath(root, ".//w:sdt"):
            key = _control_key(control, values)
            if not key:
                continue
            targets.append(
                _ControlTarget(
                    key=key,
                    part_name=part_name,
                    story=(
                        "text_box"
                        if any(
                            ancestor.tag == f"{{{W}}}txbxContent"
                            for ancestor in control.iterancestors()
                        )
                        else _story_name(part_name)
                    ),
                    control=control,
                    blocked_reason=_control_block_reason(control),
                )
            )
    return targets


def _operation_id(
    mechanism: str, key: str, part_name: str, story: str, occurrences: int
) -> str:
    payload = f"{mechanism}|{key}|{part_name}|{story}|{occurrences}".encode()
    return hashlib.sha256(payload).hexdigest()[:24]


def _safe_output_filename(source_filename: str) -> str:
    stem = Path(source_filename).stem[:100] or "template"
    return f"{stem}-filled.docx"


def build_template_fill_preview(
    path: Path,
    values: dict[str, str],
    *,
    source_filename: str | None = None,
    expected_source_sha256: str | None = None,
    max_upload_bytes: int = 50 * 1024 * 1024,
    max_uncompressed_bytes: int = 500 * 1024 * 1024,
    max_entries: int = 20_000,
) -> TemplateFillPreview:
    validate_docx_package(
        path,
        max_upload_bytes=max_upload_bytes,
        max_uncompressed_bytes=max_uncompressed_bytes,
        max_entries=max_entries,
    )
    source_sha256 = sha256_file(path)
    if expected_source_sha256 is not None and expected_source_sha256 != source_sha256:
        raise StaleTemplateFillPlanError("模板文件已变化，请重新分析并生成填写预览")
    with zipfile.ZipFile(path) as archive:
        parts = _story_parts(archive)
        placeholder_targets = _placeholder_targets(parts)
        control_targets = _control_targets(parts, values)

    grouped: Counter[tuple[str, str, str, str]] = Counter()
    blocked: list[TemplateFillBlockedTarget] = []
    detected_keys: set[str] = set()
    for placeholder_target in placeholder_targets:
        detected_keys.add(placeholder_target.key)
        if placeholder_target.key not in values:
            continue
        if placeholder_target.blocked_reason:
            blocked.append(
                TemplateFillBlockedTarget(
                    key=placeholder_target.key,
                    mechanism="placeholder",
                    story=placeholder_target.story,
                    part_name=placeholder_target.part_name,
                    reason=placeholder_target.blocked_reason,
                )
            )
            continue
        grouped[
            (
                "placeholder",
                placeholder_target.key,
                placeholder_target.story,
                placeholder_target.part_name,
            )
        ] += 1
    for control_target in control_targets:
        detected_keys.add(control_target.key)
        if control_target.key not in values:
            continue
        if control_target.blocked_reason:
            blocked.append(
                TemplateFillBlockedTarget(
                    key=control_target.key,
                    mechanism="content_control",
                    story=control_target.story,
                    part_name=control_target.part_name,
                    reason=control_target.blocked_reason,
                )
            )
            continue
        grouped[
            (
                "content_control",
                control_target.key,
                control_target.story,
                control_target.part_name,
            )
        ] += 1

    operations = [
        TemplateFillOperation(
            operation_id=_operation_id(mechanism, key, part_name, story, occurrences),
            key=key,
            mechanism=cast(Literal["placeholder", "content_control"], mechanism),
            story=story,
            part_name=part_name,
            occurrences=occurrences,
            value_preview=values[key][:80],
        )
        for (mechanism, key, story, part_name), occurrences in sorted(grouped.items())
    ]
    version_payload = {
        "source_sha256": source_sha256,
        "operations": [item.operation_id for item in operations],
        "values": {key: hashlib.sha256(values[key].encode()).hexdigest() for key in sorted(values)},
        "blocked": [item.model_dump() for item in blocked],
    }
    plan_version = hashlib.sha256(
        json.dumps(version_payload, ensure_ascii=False, sort_keys=True).encode()
    ).hexdigest()
    filename = source_filename or path.name
    return TemplateFillPreview(
        source_filename=filename,
        source_sha256=source_sha256,
        plan_version=plan_version,
        operations=operations,
        missing_keys=sorted(detected_keys - set(values)),
        unused_keys=sorted(set(values) - detected_keys),
        blocked_targets=blocked,
        replacement_count=sum(item.occurrences for item in operations),
        can_generate=bool(operations) and not blocked,
        output_filename=_safe_output_filename(filename),
    )


def _preserve_space(node: etree._Element) -> None:
    text = str(node.text or "")
    attribute = f"{{{XML}}}space"
    if text.startswith(" ") or text.endswith(" "):
        node.set(attribute, "preserve")
    elif attribute in node.attrib:
        del node.attrib[attribute]


def _apply_placeholder_targets(
    targets: list[_PlaceholderTarget], values: dict[str, str]
) -> int:
    grouped: defaultdict[int, list[_PlaceholderTarget]] = defaultdict(list)
    for target in targets:
        if target.key in values and not target.blocked_reason:
            grouped[id(target.paragraph)].append(target)
    applied = 0
    for paragraph_targets in grouped.values():
        for target in sorted(paragraph_targets, key=lambda item: item.start, reverse=True):
            nodes = target.text_nodes
            spans: list[tuple[int, int, etree._Element]] = []
            cursor = 0
            for node in nodes:
                length = len(str(node.text or ""))
                spans.append((cursor, cursor + length, node))
                cursor += length
            start_entry = next(item for item in spans if item[0] <= target.start < item[1])
            end_entry = next(item for item in spans if item[0] < target.end <= item[1])
            start_index = spans.index(start_entry)
            end_index = spans.index(end_entry)
            start_node = start_entry[2]
            end_node = end_entry[2]
            local_start = target.start - start_entry[0]
            local_end = target.end - end_entry[0]
            start_text = str(start_node.text or "")
            end_text = str(end_node.text or "")
            if start_node is end_node:
                start_node.text = (
                    start_text[:local_start] + values[target.key] + start_text[local_end:]
                )
                _preserve_space(start_node)
            else:
                start_node.text = start_text[:local_start] + values[target.key]
                _preserve_space(start_node)
                for node in nodes[start_index + 1 : end_index]:
                    node.text = ""
                    _preserve_space(node)
                end_node.text = end_text[local_end:]
                _preserve_space(end_node)
            applied += 1
    return applied


def _direct_control_text_nodes(control: etree._Element) -> list[etree._Element]:
    return [
        node
        for node in _xpath(control, "./w:sdtContent//w:t")
        if _nearest(node, f"{{{W}}}sdt") is control
    ]


def _replace_control_content(control: etree._Element, value: str) -> None:
    content_nodes = _xpath(control, "./w:sdtContent")
    if not content_nodes:
        raise TemplateFillInvalidError("内容控件缺少 sdtContent，无法安全填写")
    content = content_nodes[0]
    text_nodes = _direct_control_text_nodes(control)
    if not text_nodes:
        paragraphs = _xpath(content, ".//w:p")
        paragraph = paragraphs[0] if paragraphs else etree.SubElement(content, f"{{{W}}}p")
        run = etree.SubElement(paragraph, f"{{{W}}}r")
        text_node = etree.SubElement(run, f"{{{W}}}t")
        text_nodes = [text_node]
    text_nodes[0].text = value
    _preserve_space(text_nodes[0])
    for node in text_nodes[1:]:
        node.text = ""
        _preserve_space(node)
    for marker in _xpath(control, "./w:sdtPr/w:showingPlcHdr"):
        marker.getparent().remove(marker)


def _apply_control_targets(targets: list[_ControlTarget], values: dict[str, str]) -> int:
    applied = 0
    for target in targets:
        if target.key not in values or target.blocked_reason:
            continue
        _replace_control_content(target.control, values[target.key])
        applied += 1
    return applied


def _serialize(root: etree._Element) -> bytes:
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def _story_parts_from_archive(parts: dict[str, bytes]) -> dict[str, etree._Element]:
    return {
        name: _parse(payload)
        for name, payload in sorted(parts.items())
        if EDITABLE_PART.fullmatch(name)
    }


def apply_fill_to_parts(parts: dict[str, bytes], values: dict[str, str]) -> tuple[set[str], int]:
    story_parts = _story_parts_from_archive(parts)
    placeholders = _placeholder_targets(story_parts)
    controls = _control_targets(story_parts, values)
    placeholder_count = _apply_placeholder_targets(placeholders, values)
    control_count = _apply_control_targets(controls, values)
    applied_count = placeholder_count + control_count
    if applied_count == 0:
        return set(), 0
    changed_part_names = {
        target.part_name
        for target in placeholders
        if target.key in values and not target.blocked_reason
    }
    changed_part_names.update(
        target.part_name
        for target in controls
        if target.key in values and not target.blocked_reason
    )
    for name in changed_part_names:
        parts[name] = _serialize(story_parts[name])
    return changed_part_names, applied_count


def _protection_inventory(root: etree._Element) -> dict[str, int]:
    return {name: len(_xpath(root, expression)) for name, expression in PROTECTION_COUNTERS.items()}


def _verify_output(
    source: Path, output: Path, changed_parts: list[str], source_sha256: str
) -> dict[str, bool]:
    checks = {
        "source_unchanged": sha256_file(source) == source_sha256,
        "part_inventory_unchanged": False,
        "unchanged_parts_exact": False,
        "protected_objects_unchanged": False,
        "reopens_with_word_model": False,
    }
    with zipfile.ZipFile(source) as before, zipfile.ZipFile(output) as after:
        before_names = before.namelist()
        after_names = after.namelist()
        checks["part_inventory_unchanged"] = before_names == after_names
        changed = set(changed_parts)
        checks["unchanged_parts_exact"] = all(
            before.read(name) == after.read(name) for name in before_names if name not in changed
        )
        checks["protected_objects_unchanged"] = all(
            _protection_inventory(_parse(before.read(name)))
            == _protection_inventory(_parse(after.read(name)))
            for name in changed_parts
        )
    try:
        Document(str(output))
        checks["reopens_with_word_model"] = True
    except Exception:
        checks["reopens_with_word_model"] = False
    if not all(checks.values()):
        failures = ", ".join(name for name, passed in checks.items() if not passed)
        raise IntegrityCheckFailedError(f"模板填写结果未通过完整性检查: {failures}")
    return checks


def fill_template(
    source: Path,
    output: Path,
    values: dict[str, str],
    *,
    source_filename: str | None = None,
    expected_source_sha256: str,
    expected_plan_version: str,
    max_upload_bytes: int = 50 * 1024 * 1024,
    max_uncompressed_bytes: int = 500 * 1024 * 1024,
    max_entries: int = 20_000,
) -> TemplateFillResult:
    preview = build_template_fill_preview(
        source,
        values,
        source_filename=source_filename,
        expected_source_sha256=expected_source_sha256,
        max_upload_bytes=max_upload_bytes,
        max_uncompressed_bytes=max_uncompressed_bytes,
        max_entries=max_entries,
    )
    if preview.plan_version != expected_plan_version:
        raise StaleTemplateFillPlanError()
    if preview.blocked_targets:
        raise TemplateFillBlockedError(
            "部分填写目标位于受保护结构中，未生成文件",
            [f"{item.key}: {item.reason}" for item in preview.blocked_targets],
        )
    if not preview.operations:
        raise TemplateFillInvalidError("没有可执行的模板填写操作")

    with zipfile.ZipFile(source) as archive:
        parts = {info.filename: archive.read(info.filename) for info in archive.infolist()}
        changed_part_names, applied_count = apply_fill_to_parts(parts, values)
        if applied_count != preview.replacement_count:
            raise StaleTemplateFillPlanError("模板目标数量已变化，请重新生成填写预览")
        patched = {name: parts[name] for name in changed_part_names}
        output.parent.mkdir(parents=True, exist_ok=True)
        temporary = output.with_suffix(".tmp.docx")
        with zipfile.ZipFile(temporary, "w") as target_archive:
            for info in archive.infolist():
                data = patched.get(info.filename, archive.read(info.filename))
                target_archive.writestr(info, data)
    validate_docx_package(
        temporary,
        max_upload_bytes=max(max_upload_bytes, temporary.stat().st_size),
        max_uncompressed_bytes=max_uncompressed_bytes,
        max_entries=max_entries,
    )
    checks = _verify_output(source, temporary, sorted(patched), preview.source_sha256)
    temporary.replace(output)
    return TemplateFillResult(
        preview=preview,
        output_sha256=sha256_file(output),
        output_filename=preview.output_filename,
        changed_parts=sorted(patched),
        integrity_checks=checks,
    )
