from __future__ import annotations

import hashlib
import json
import posixpath
import re
import zipfile
from collections import Counter
from copy import deepcopy
from pathlib import Path
from typing import Any, cast

from docx import Document
from lxml import etree

from paper_setting_core.documents import inspect_document, sha256_file, validate_docx_package
from paper_setting_core.errors import (
    IntegrityCheckFailedError,
    StaleTemplateCombinedPlanError,
    TemplateCombinedBlockedError,
)
from paper_setting_core.templates.filling import apply_fill_to_parts, build_template_fill_preview
from paper_setting_core.templates.inspection import PLACEHOLDER, PR, R, W, _parse
from paper_setting_core.templates.models import (
    ManuscriptBodyInjectionPreview,
    ManuscriptCompositionPreview,
    ManuscriptCompositionResult,
    ManuscriptFieldMapping,
    TemplateSectionStructureConfig,
)
from paper_setting_core.templates.structure import (
    CONTENT_TYPES_PART,
    DOCUMENT_PART,
    DOCUMENT_RELS_PART,
    _structure_matches,
    apply_structure_to_parts,
    build_section_structure_preview,
)

M = "http://schemas.openxmlformats.org/officeDocument/2006/math"
CT = "http://schemas.openxmlformats.org/package/2006/content-types"
REL_IMAGE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"
REL_HYPERLINK = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink"
REL_NUMBERING = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/numbering"
NS = {"w": W, "r": R, "pr": PR, "m": M, "ct": CT}
BODY_ANCHORS = {
    "document.body",
    "manuscript.body",
    "paper.body",
    "body",
    "正文",
    "论文正文",
}
STYLE_PART = "word/styles.xml"
NUMBERING_PART = "word/numbering.xml"

FIELD_ALIASES: dict[str, set[str]] = {
    "title": {
        "title",
        "paper.title",
        "thesis.title",
        "document.title",
        "论文题目",
        "题目",
    },
    "abstract": {"abstract", "paper.abstract", "thesis.abstract", "摘要"},
    "keywords": {
        "keywords",
        "paper.keywords",
        "thesis.keywords",
        "关键词",
        "关键字",
    },
    "author": {
        "author",
        "paper.author",
        "student.name",
        "student_name",
        "作者",
        "姓名",
        "学生姓名",
    },
    "student_id": {
        "student.id",
        "student_id",
        "student.number",
        "student_number",
        "学号",
        "学生编号",
    },
    "school": {"school", "university", "institution", "paper.school", "学校", "院校"},
    "college": {"college", "department", "school.department", "学院", "院系"},
    "major": {"major", "student.major", "paper.major", "专业"},
    "advisor": {
        "advisor",
        "supervisor",
        "teacher",
        "paper.advisor",
        "导师",
        "指导教师",
        "指导老师",
    },
}

LABEL_PATTERNS: dict[str, re.Pattern[str]] = {
    "author": re.compile(r"^(?:作者|学生姓名|姓名)\s*[：:]\s*(.+)$"),
    "student_id": re.compile(r"^(?:学号|学生编号)\s*[：:]\s*(.+)$"),
    "school": re.compile(r"^(?:学校|院校)\s*[：:]\s*(.+)$"),
    "college": re.compile(r"^(?:学院|院系|系别)\s*[：:]\s*(.+)$"),
    "major": re.compile(r"^(?:专业|专业名称)\s*[：:]\s*(.+)$"),
    "advisor": re.compile(r"^(?:导师|指导教师|指导老师)\s*[：:]\s*(.+)$"),
}


def _xpath(root: etree._Element, expression: str) -> list[Any]:
    return cast(list[Any], root.xpath(expression, namespaces=NS))


def _serialize(root: etree._Element) -> bytes:
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def _paragraph_text(paragraph: etree._Element) -> str:
    return "".join(str(node.text or "") for node in paragraph.iter(f"{{{W}}}t"))


def _block_text(block: etree._Element) -> str:
    return re.sub(
        r"\s+", " ", "".join(str(node.text or "") for node in block.iter(f"{{{W}}}t"))
    ).strip()


def _read_parts(path: Path) -> tuple[list[zipfile.ZipInfo], dict[str, bytes]]:
    with zipfile.ZipFile(path) as archive:
        infos = archive.infolist()
        return infos, {info.filename: archive.read(info.filename) for info in infos}


def _full_text_by_path(path: Path) -> dict[str, str]:
    document = Document(str(path))
    values: dict[str, str] = {}
    paragraph_index = 0
    table_index = 0
    for child in document.element.body.iterchildren():
        if child.tag == f"{{{W}}}p":
            values[f"body/p[{paragraph_index}]"] = _paragraph_text(child)
            paragraph_index += 1
        elif child.tag == f"{{{W}}}tbl":
            from docx.table import Table

            table = Table(child, document)
            for row_index, row in enumerate(table.rows):
                for cell_index, cell in enumerate(row.cells):
                    for item_index, paragraph in enumerate(cell.paragraphs):
                        values[
                            f"table[{table_index}]/r[{row_index}]/c[{cell_index}]/p[{item_index}]"
                        ] = paragraph.text
            table_index += 1
    return values


def _template_field_keys(template: Path) -> list[str]:
    from paper_setting_core.templates.inspection import inspect_template

    inspection = inspect_template(template)
    return sorted(
        {
            *[item.key for item in inspection.placeholders],
            *[
                value
                for item in inspection.content_controls
                for value in (item.tag, item.alias)
                if value
            ],
        }
    )


def _canonical_field(key: str) -> str | None:
    normalized = key.strip().lower().replace("-", "_")
    for canonical, aliases in FIELD_ALIASES.items():
        if normalized in {item.replace("-", "_") for item in aliases}:
            return canonical
    return None


def _semantic_values(manuscript: Path) -> dict[str, tuple[str, str, float, str]]:
    inspection = inspect_document(manuscript)
    full_text = _full_text_by_path(manuscript)
    role_values: dict[str, list[str]] = {}
    for item in inspection.items:
        if item.story not in {"body", "table"}:
            continue
        text = full_text.get(item.path, item.text_preview).strip()
        if text:
            role_values.setdefault(item.semantic_role, []).append(text)

    values: dict[str, tuple[str, str, float, str]] = {}
    titles = role_values.get("paper_title", [])
    if titles:
        values["title"] = (titles[0], "semantic_role", 0.96, "识别为论文标题")
    abstracts = role_values.get("abstract_body", [])
    if abstracts:
        values["abstract"] = (
            " ".join(abstracts),
            "semantic_role",
            0.92,
            f"合并 {len(abstracts)} 段摘要正文",
        )
    keywords = role_values.get("keywords", [])
    if keywords:
        value = re.sub(r"^(?:关键词|关键字|Keywords?)\s*[：:]\s*", "", keywords[0], flags=re.I)
        values["keywords"] = (value, "semantic_role", 0.94, "识别为关键词段落")

    early_texts = [
        full_text.get(item.path, item.text_preview).strip()
        for item in inspection.items
        if item.story in {"body", "table"} and item.order < 40
    ]
    for canonical, pattern in LABEL_PATTERNS.items():
        for text in early_texts:
            match = pattern.match(text)
            if match and match.group(1).strip():
                values[canonical] = (
                    match.group(1).strip(),
                    "label",
                    0.9,
                    f"由标签“{text[:40]}”提取",
                )
                break
    return values


def infer_template_values(
    manuscript: Path,
    template: Path,
    explicit_values: dict[str, str],
) -> tuple[dict[str, str], list[ManuscriptFieldMapping]]:
    semantic = _semantic_values(manuscript)
    values = dict(explicit_values)
    mappings: list[ManuscriptFieldMapping] = []
    for key in _template_field_keys(template):
        if key in BODY_ANCHORS:
            continue
        if key in explicit_values:
            value = explicit_values[key]
            mappings.append(
                ManuscriptFieldMapping(
                    key=key,
                    value_preview=value[:160],
                    source="rule_pack",
                    confidence=1,
                    evidence="规则包显式填写值",
                )
            )
            continue
        canonical = _canonical_field(key)
        candidate = semantic.get(canonical or "")
        if candidate is None:
            mappings.append(
                ManuscriptFieldMapping(
                    key=key,
                    value_preview="",
                    source="unmapped",
                    confidence=0,
                    evidence="未找到足够可靠的来源，不自动猜测",
                )
            )
            continue
        value, source, confidence, evidence = candidate
        values[key] = value
        mappings.append(
            ManuscriptFieldMapping(
                key=key,
                value_preview=value[:160],
                source=cast(Any, source),
                confidence=confidence,
                evidence=evidence,
            )
        )
    return values, mappings


def _source_body_blocks(
    manuscript_parts: dict[str, bytes],
    manuscript: Path,
    mapped_front_matter: bool,
) -> tuple[list[etree._Element], int]:
    root = _parse(manuscript_parts[DOCUMENT_PART])
    body = _xpath(root, ".//w:body")[0]
    blocks = [child for child in body if child.tag != f"{{{W}}}sectPr"]
    start_index = 0
    if mapped_front_matter:
        inspection = inspect_document(manuscript)
        first_heading = next(
            (
                item
                for item in inspection.items
                if item.story == "body" and item.semantic_role == "heading_1"
            ),
            None,
        )
        if first_heading is not None:
            paragraph_target = int(first_heading.path.removeprefix("body/p[").removesuffix("]"))
            paragraph_seen = 0
            for index, block in enumerate(blocks):
                if block.tag == f"{{{W}}}p":
                    if paragraph_seen == paragraph_target:
                        start_index = index
                        break
                    paragraph_seen += 1
    return blocks[start_index:], start_index


def _relationship_map(parts: dict[str, bytes]) -> dict[str, etree._Element]:
    if DOCUMENT_RELS_PART not in parts:
        return {}
    root = _parse(parts[DOCUMENT_RELS_PART])
    return {
        str(node.get("Id")): node for node in _xpath(root, "./pr:Relationship") if node.get("Id")
    }


def _body_preview(
    manuscript: Path,
    template: Path,
    mappings: list[ManuscriptFieldMapping],
) -> ManuscriptBodyInjectionPreview:
    _, source_parts = _read_parts(manuscript)
    _, template_parts = _read_parts(template)
    source_blocks, excluded = _source_body_blocks(
        source_parts,
        manuscript,
        any(item.source != "unmapped" for item in mappings),
    )
    stripped_sections = sum(len(_xpath(block, ".//w:sectPr")) for block in source_blocks)
    analysis_blocks = [deepcopy(block) for block in source_blocks]
    for block in analysis_blocks:
        for section in _xpath(block, ".//w:sectPr"):
            parent = section.getparent()
            if parent is not None:
                parent.remove(section)
    template_root = _parse(template_parts[DOCUMENT_PART])
    anchors: list[tuple[etree._Element, str]] = []
    for paragraph in _xpath(template_root, ".//w:body/w:p"):
        text = _paragraph_text(paragraph).strip()
        match = PLACEHOLDER.fullmatch(text)
        if match:
            key = match.group("brace") or match.group("bracket") or match.group("angle")
            if key in BODY_ANCHORS:
                anchors.append((paragraph, key))

    blockers: list[str] = []
    if not anchors:
        blockers.append("模板正文缺少独立段落占位符 {{document.body}}")
    elif len(anchors) > 1:
        blockers.append("模板中存在多个正文占位符，请只保留一个")
    if not source_blocks:
        blockers.append("原稿没有可注入的正文块")

    source_rels = _relationship_map(source_parts)
    relationship_ids = {
        str(value)
        for block in analysis_blocks
        for node in block.iter()
        for name, value in node.attrib.items()
        if name in {f"{{{R}}}id", f"{{{R}}}embed", f"{{{R}}}link"}
    }
    image_relationships = 0
    for relationship_id in sorted(relationship_ids):
        relationship = source_rels.get(relationship_id)
        relationship_type = str(relationship.get("Type") or "") if relationship is not None else ""
        if relationship_type == REL_IMAGE:
            image_relationships += 1
        elif relationship_type != REL_HYPERLINK:
            blockers.append(
                f"正文包含暂不支持迁移的关系对象：{relationship_type or relationship_id}"
            )
    unsupported_nodes = {
        "脚注": ".//w:footnoteReference",
        "尾注": ".//w:endnoteReference",
        "批注": ".//w:commentReference | .//w:commentRangeStart | .//w:commentRangeEnd",
        "嵌入对象": ".//w:object | .//w:altChunk",
    }
    for label, expression in unsupported_nodes.items():
        if any(_xpath(block, expression) for block in analysis_blocks):
            blockers.append(f"正文包含{label}，当前版本为避免引用 ID 损坏而停止注入")
    operation_payload = {
        "source": sha256_file(manuscript),
        "template": sha256_file(template),
        "anchor": anchors[0][1] if len(anchors) == 1 else "document.body",
        "blocks": [_block_text(block) for block in source_blocks],
    }
    operation_id = hashlib.sha256(
        json.dumps(operation_payload, ensure_ascii=False, sort_keys=True).encode()
    ).hexdigest()[:24]
    return ManuscriptBodyInjectionPreview(
        operation_id=operation_id,
        anchor_key=anchors[0][1] if len(anchors) == 1 else "document.body",
        source_blocks=len(source_blocks),
        paragraph_blocks=sum(block.tag == f"{{{W}}}p" for block in source_blocks),
        table_blocks=sum(block.tag == f"{{{W}}}tbl" for block in source_blocks),
        image_relationships=image_relationships,
        formula_count=sum(
            len(_xpath(block, ".//m:oMath | .//m:oMathPara")) for block in analysis_blocks
        ),
        stripped_section_properties=stripped_sections,
        excluded_front_matter_blocks=excluded,
        blockers=blockers,
    )


def build_manuscript_composition_preview(
    manuscript: Path,
    template: Path,
    explicit_values: dict[str, str],
    structure_configuration: TemplateSectionStructureConfig,
    *,
    source_filename: str | None = None,
    template_filename: str | None = None,
    expected_source_sha256: str | None = None,
    expected_template_sha256: str | None = None,
    max_upload_bytes: int = 50 * 1024 * 1024,
    max_uncompressed_bytes: int = 500 * 1024 * 1024,
    max_entries: int = 20_000,
) -> ManuscriptCompositionPreview:
    for path in (manuscript, template):
        validate_docx_package(
            path,
            max_upload_bytes=max_upload_bytes,
            max_uncompressed_bytes=max_uncompressed_bytes,
            max_entries=max_entries,
        )
    source_sha256 = sha256_file(manuscript)
    template_sha256 = sha256_file(template)
    if expected_source_sha256 is not None and source_sha256 != expected_source_sha256:
        raise StaleTemplateCombinedPlanError("原稿已变化，请重新生成模板注入计划")
    if expected_template_sha256 is not None and template_sha256 != expected_template_sha256:
        raise StaleTemplateCombinedPlanError("模板已变化，请重新生成模板注入计划")

    values, mappings = infer_template_values(manuscript, template, explicit_values)
    fill = build_template_fill_preview(
        template,
        values,
        source_filename=template_filename,
        expected_source_sha256=template_sha256,
        max_upload_bytes=max_upload_bytes,
        max_uncompressed_bytes=max_uncompressed_bytes,
        max_entries=max_entries,
    )
    structure = build_section_structure_preview(
        template,
        structure_configuration,
        source_filename=template_filename,
        expected_source_sha256=template_sha256,
        max_upload_bytes=max_upload_bytes,
        max_uncompressed_bytes=max_uncompressed_bytes,
        max_entries=max_entries,
    )
    body = _body_preview(manuscript, template, mappings)
    blockers = [*body.blockers, *structure.blockers]
    blockers.extend(f"{item.key}: {item.reason}" for item in fill.blocked_targets)
    version_payload = {
        "source_sha256": source_sha256,
        "template_sha256": template_sha256,
        "mappings": [item.model_dump() for item in mappings],
        "fill_plan": fill.plan_version,
        "body_operation": body.operation_id,
        "structure_plan": structure.plan_version,
    }
    plan_version = hashlib.sha256(
        json.dumps(version_payload, ensure_ascii=False, sort_keys=True).encode()
    ).hexdigest()
    filename = source_filename or manuscript.name
    return ManuscriptCompositionPreview(
        source_filename=filename,
        source_sha256=source_sha256,
        template_filename=template_filename or template.name,
        template_sha256=template_sha256,
        plan_version=plan_version,
        mappings=mappings,
        fill=fill,
        body_injection=body,
        structure=structure,
        can_generate=bool(body.source_blocks) and not blockers,
        blockers=blockers,
        output_filename=f"{Path(filename).stem[:100] or 'manuscript'}-composed.docx",
    )


def _next_relationship_id(root: etree._Element) -> str:
    used = {str(node.get("Id")) for node in _xpath(root, "./pr:Relationship")}
    index = 1
    while f"rId{index}" in used:
        index += 1
    return f"rId{index}"


def _ensure_content_type(
    parts: dict[str, bytes], source_parts: dict[str, bytes], suffix: str
) -> None:
    target_root = _parse(parts[CONTENT_TYPES_PART])
    source_root = _parse(source_parts[CONTENT_TYPES_PART])
    extension = suffix.lstrip(".").lower()
    existing = {
        str(node.get("Extension") or "").lower() for node in _xpath(target_root, "./ct:Default")
    }
    if extension not in existing:
        source_default = next(
            (
                node
                for node in _xpath(source_root, "./ct:Default")
                if str(node.get("Extension") or "").lower() == extension
            ),
            None,
        )
        if source_default is not None:
            target_root.append(deepcopy(source_default))
            parts[CONTENT_TYPES_PART] = _serialize(target_root)


def _copy_styles(
    parts: dict[str, bytes], source_parts: dict[str, bytes], blocks: list[etree._Element]
) -> bool:
    if STYLE_PART not in parts or STYLE_PART not in source_parts:
        return False
    target = _parse(parts[STYLE_PART])
    source = _parse(source_parts[STYLE_PART])
    target_ids = {
        str(node.get(f"{{{W}}}styleId"))
        for node in _xpath(target, "./w:style")
        if node.get(f"{{{W}}}styleId")
    }
    source_by_id = {
        str(node.get(f"{{{W}}}styleId")): node
        for node in _xpath(source, "./w:style")
        if node.get(f"{{{W}}}styleId")
    }
    required = {
        str(node.get(f"{{{W}}}val"))
        for block in blocks
        for node in _xpath(block, ".//w:pStyle | .//w:rStyle | .//w:tblStyle")
        if node.get(f"{{{W}}}val")
    }
    queue = list(required)
    changed = False
    while queue:
        style_id = queue.pop()
        if style_id in target_ids or style_id not in source_by_id:
            continue
        style = deepcopy(source_by_id[style_id])
        target.append(style)
        target_ids.add(style_id)
        changed = True
        queue.extend(
            str(node.get(f"{{{W}}}val"))
            for node in _xpath(style, "./w:basedOn | ./w:next | ./w:link")
            if node.get(f"{{{W}}}val")
        )
    if changed:
        parts[STYLE_PART] = _serialize(target)
    return changed


def _copy_numbering(
    parts: dict[str, bytes], source_parts: dict[str, bytes], blocks: list[etree._Element]
) -> bool:
    used = {
        int(str(node.get(f"{{{W}}}val")))
        for block in blocks
        for node in _xpath(block, ".//w:numPr/w:numId")
        if str(node.get(f"{{{W}}}val") or "").isdigit()
    }
    if not used or NUMBERING_PART not in source_parts or NUMBERING_PART not in parts:
        return False
    source = _parse(source_parts[NUMBERING_PART])
    target = _parse(parts[NUMBERING_PART])
    source_nums = {
        int(str(node.get(f"{{{W}}}numId"))): node
        for node in _xpath(source, "./w:num")
        if str(node.get(f"{{{W}}}numId") or "").isdigit()
    }
    source_abstract = {
        int(str(node.get(f"{{{W}}}abstractNumId"))): node
        for node in _xpath(source, "./w:abstractNum")
        if str(node.get(f"{{{W}}}abstractNumId") or "").isdigit()
    }
    next_num = (
        max(
            [
                int(str(node.get(f"{{{W}}}numId")))
                for node in _xpath(target, "./w:num")
                if str(node.get(f"{{{W}}}numId") or "").isdigit()
            ]
            or [0]
        )
        + 1
    )
    next_abstract = (
        max(
            [
                int(str(node.get(f"{{{W}}}abstractNumId")))
                for node in _xpath(target, "./w:abstractNum")
                if str(node.get(f"{{{W}}}abstractNumId") or "").isdigit()
            ]
            or [0]
        )
        + 1
    )
    remap: dict[int, int] = {}
    for old_num in sorted(used):
        source_num = source_nums.get(old_num)
        if source_num is None:
            continue
        abstract_nodes = _xpath(source_num, "./w:abstractNumId")
        if not abstract_nodes:
            continue
        old_abstract = int(str(abstract_nodes[0].get(f"{{{W}}}val")))
        source_definition = source_abstract.get(old_abstract)
        if source_definition is None:
            continue
        abstract_copy = deepcopy(source_definition)
        abstract_copy.set(f"{{{W}}}abstractNumId", str(next_abstract))
        num_copy = deepcopy(source_num)
        num_copy.set(f"{{{W}}}numId", str(next_num))
        _xpath(num_copy, "./w:abstractNumId")[0].set(f"{{{W}}}val", str(next_abstract))
        target.append(abstract_copy)
        target.append(num_copy)
        remap[old_num] = next_num
        next_num += 1
        next_abstract += 1
    for block in blocks:
        for node in _xpath(block, ".//w:numPr/w:numId"):
            old_value = str(node.get(f"{{{W}}}val") or "")
            if old_value.isdigit() and int(old_value) in remap:
                node.set(f"{{{W}}}val", str(remap[int(old_value)]))
    if remap:
        parts[NUMBERING_PART] = _serialize(target)
    return bool(remap)


def _copy_relationships(
    parts: dict[str, bytes], source_parts: dict[str, bytes], blocks: list[etree._Element]
) -> tuple[int, set[str]]:
    source_map = _relationship_map(source_parts)
    target_root = _parse(parts[DOCUMENT_RELS_PART])
    copied_images = 0
    added_parts: set[str] = set()
    relationship_remap: dict[str, str] = {}
    for block in blocks:
        for node in block.iter():
            for attribute in (f"{{{R}}}id", f"{{{R}}}embed", f"{{{R}}}link"):
                old_id = node.get(attribute)
                if not old_id:
                    continue
                if old_id not in relationship_remap:
                    relationship = source_map[old_id]
                    relationship_type = str(relationship.get("Type") or "")
                    new_id = _next_relationship_id(target_root)
                    relationship_copy = deepcopy(relationship)
                    relationship_copy.set("Id", new_id)
                    if relationship_type == REL_IMAGE:
                        target = str(relationship.get("Target") or "")
                        source_name = posixpath.normpath(posixpath.join("word", target))
                        payload = source_parts[source_name]
                        suffix = Path(source_name).suffix
                        digest = hashlib.sha256(payload).hexdigest()[:16]
                        target_name = f"word/media/injected-{digest}{suffix}"
                        counter = 1
                        while target_name in parts and parts[target_name] != payload:
                            target_name = f"word/media/injected-{digest}-{counter}{suffix}"
                            counter += 1
                        if target_name not in parts:
                            parts[target_name] = payload
                            added_parts.add(target_name)
                            _ensure_content_type(parts, source_parts, suffix)
                        relationship_copy.set("Target", target_name.removeprefix("word/"))
                        copied_images += 1
                    target_root.append(relationship_copy)
                    relationship_remap[old_id] = new_id
                node.set(attribute, relationship_remap[old_id])
    parts[DOCUMENT_RELS_PART] = _serialize(target_root)
    return copied_images, added_parts


def _inject_body(
    parts: dict[str, bytes],
    source_parts: dict[str, bytes],
    manuscript: Path,
    mappings: list[ManuscriptFieldMapping],
) -> tuple[set[str], set[str], list[str], int]:
    source_blocks, _ = _source_body_blocks(
        source_parts,
        manuscript,
        any(item.source != "unmapped" for item in mappings),
    )
    blocks = [deepcopy(block) for block in source_blocks]
    for block in blocks:
        for section in _xpath(block, ".//w:sectPr"):
            parent = section.getparent()
            if parent is not None:
                parent.remove(section)
    root = _parse(parts[DOCUMENT_PART])
    anchors = []
    for paragraph in _xpath(root, ".//w:body/w:p"):
        match = PLACEHOLDER.fullmatch(_paragraph_text(paragraph).strip())
        key = (
            match.group("brace") or match.group("bracket") or match.group("angle")
            if match
            else None
        )
        if key in BODY_ANCHORS:
            anchors.append(paragraph)
    if len(anchors) != 1:
        raise TemplateCombinedBlockedError("正文占位符无效", ["模板必须包含一个正文占位符"])

    changed = {DOCUMENT_PART, DOCUMENT_RELS_PART}
    if _copy_styles(parts, source_parts, blocks):
        changed.add(STYLE_PART)
    if _copy_numbering(parts, source_parts, blocks):
        changed.add(NUMBERING_PART)
    copied_images, added = _copy_relationships(parts, source_parts, blocks)
    if added:
        changed.add(CONTENT_TYPES_PART)
    anchor = anchors[0]
    parent = anchor.getparent()
    index = parent.index(anchor)
    parent.remove(anchor)
    for offset, block in enumerate(blocks):
        parent.insert(index + offset, block)
    parts[DOCUMENT_PART] = _serialize(root)
    return (
        changed,
        added,
        [_block_text(block) for block in source_blocks if _block_text(block)],
        copied_images,
    )


def apply_manuscript_composition(
    manuscript: Path,
    template: Path,
    output: Path,
    explicit_values: dict[str, str],
    structure_configuration: TemplateSectionStructureConfig,
    *,
    source_filename: str | None = None,
    template_filename: str | None = None,
    expected_source_sha256: str,
    expected_template_sha256: str,
    expected_plan_version: str,
    max_upload_bytes: int = 50 * 1024 * 1024,
    max_uncompressed_bytes: int = 500 * 1024 * 1024,
    max_entries: int = 20_000,
) -> ManuscriptCompositionResult:
    preview = build_manuscript_composition_preview(
        manuscript,
        template,
        explicit_values,
        structure_configuration,
        source_filename=source_filename,
        template_filename=template_filename,
        expected_source_sha256=expected_source_sha256,
        expected_template_sha256=expected_template_sha256,
        max_upload_bytes=max_upload_bytes,
        max_uncompressed_bytes=max_uncompressed_bytes,
        max_entries=max_entries,
    )
    if preview.plan_version != expected_plan_version:
        raise StaleTemplateCombinedPlanError("模板注入计划已变化，请重新审核")
    if preview.blockers:
        raise TemplateCombinedBlockedError("模板注入计划包含不可执行项", preview.blockers)

    template_infos, parts = _read_parts(template)
    _, source_parts = _read_parts(manuscript)
    values, _ = infer_template_values(manuscript, template, explicit_values)
    changed, added, injected_text, copied_images = _inject_body(
        parts, source_parts, manuscript, preview.mappings
    )
    fill_changed, applied_count = apply_fill_to_parts(parts, values)
    if applied_count != preview.fill.replacement_count:
        raise StaleTemplateCombinedPlanError("智能字段目标数量已变化，请重新审核")
    changed.update(fill_changed)
    structure_changed, structure_added = apply_structure_to_parts(parts, structure_configuration)
    changed.update(structure_changed)
    added.update(structure_added)

    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = output.with_suffix(".tmp.docx")
    original_names = {info.filename for info in template_infos}
    with zipfile.ZipFile(temporary, "w") as archive:
        for info in template_infos:
            archive.writestr(info, parts[info.filename])
        for name in sorted(set(parts) - original_names):
            archive.writestr(name, parts[name], compress_type=zipfile.ZIP_DEFLATED)
    validate_docx_package(
        temporary,
        max_upload_bytes=max(max_upload_bytes, temporary.stat().st_size),
        max_uncompressed_bytes=max_uncompressed_bytes,
        max_entries=max_entries,
    )
    output_text = []
    with zipfile.ZipFile(temporary) as archive:
        output_root = _parse(archive.read(DOCUMENT_PART))
        output_text = [
            _block_text(block)
            for block in _xpath(output_root, ".//w:body/*[self::w:p or self::w:tbl]")
            if _block_text(block)
        ]
    source_counter = Counter(injected_text)
    output_counter = Counter(output_text)
    body_preserved = all(output_counter[text] >= count for text, count in source_counter.items())
    checks = {
        "source_unchanged": sha256_file(manuscript) == preview.source_sha256,
        "template_unchanged": sha256_file(template) == preview.template_sha256,
        "body_text_preserved": body_preserved,
        "field_targets_applied": applied_count == preview.fill.replacement_count,
        "images_copied": copied_images == preview.body_injection.image_relationships,
        "requested_structure_applied": _structure_matches(temporary, structure_configuration),
        "reopens_with_word_model": False,
    }
    try:
        Document(str(temporary))
        checks["reopens_with_word_model"] = True
    except Exception:
        pass
    if not all(checks.values()):
        failures = ", ".join(name for name, passed in checks.items() if not passed)
        temporary.unlink(missing_ok=True)
        raise IntegrityCheckFailedError(f"整篇正文注入未通过完整性检查: {failures}")
    temporary.replace(output)
    return ManuscriptCompositionResult(
        preview=preview,
        output_sha256=sha256_file(output),
        output_filename=preview.output_filename,
        changed_parts=sorted(changed),
        added_parts=sorted(added),
        integrity_checks=checks,
    )
