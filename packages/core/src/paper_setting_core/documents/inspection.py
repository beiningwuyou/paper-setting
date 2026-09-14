from __future__ import annotations

import hashlib
import re
from collections import Counter
from collections.abc import Iterator
from pathlib import Path

from docx import Document
from docx.document import Document as DocumentObject
from docx.table import Table
from docx.text.paragraph import Paragraph

from paper_setting_core.advanced_word.inspection import snapshot_advanced_word
from paper_setting_core.classification.service import KEYWORDS, classify_paragraph, division_role
from paper_setting_core.documents.format_snapshot import (
    FormatContext,
    build_format_context,
    snapshot_paragraph_format,
    snapshot_sections,
)
from paper_setting_core.documents.models import (
    DocumentInspection,
    DocumentItem,
    InspectionSummary,
    RiskFlags,
)


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def normalized_text(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def text_hash(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def _stable_id(story: str, path: str, text: str, previous_hash: str) -> str:
    material = f"{story}|{path}|{text_hash(text)}|{previous_hash}"
    return hashlib.sha256(material.encode("utf-8")).hexdigest()[:24]


def paragraph_risks(paragraph: Paragraph, *, nested_table: bool = False) -> RiskFlags:
    xml = paragraph._p.xml
    revision_nodes = paragraph._p.xpath(
        ".//w:ins | .//w:del | .//w:moveFrom | .//w:moveTo | "
        ".//w:moveFromRangeStart | .//w:moveFromRangeEnd | "
        ".//w:moveToRangeStart | .//w:moveToRangeEnd"
    )
    return RiskFlags(
        formula="<m:oMath" in xml or "<m:oMathPara" in xml,
        text_box="txbxContent" in xml,
        drawing="<w:drawing" in xml or "<w:pict" in xml,
        revision=bool(revision_nodes),
        field="<w:fldChar" in xml or "<w:instrText" in xml or "<w:fldSimple" in xml,
        external_link='TargetMode="External"' in xml,
        embedded_object="<w:object" in xml or "<o:OLEObject" in xml,
        nested_table=nested_table,
    )


def _iter_body(document: DocumentObject) -> Iterator[tuple[str, Paragraph, bool]]:
    body = document.element.body
    paragraph_index = 0
    table_index = 0
    for child in body.iterchildren():
        if child.tag.endswith("}p"):
            yield f"body/p[{paragraph_index}]", Paragraph(child, document), False
            paragraph_index += 1
        elif child.tag.endswith("}tbl"):
            table = Table(child, document)
            for row_index, row in enumerate(table.rows):
                for cell_index, cell in enumerate(row.cells):
                    nested = bool(
                        cell._tc.xpath("./w:tbl//w:tbl | ./w:tcPr/w:gridSpan | ./w:tcPr/w:vMerge")
                    )
                    for item_index, paragraph in enumerate(cell.paragraphs):
                        path = (
                            f"table[{table_index}]/r[{row_index}]/c[{cell_index}]/p[{item_index}]"
                        )
                        yield path, paragraph, nested
            table_index += 1


def _append_item(
    items: list[DocumentItem],
    paragraph: Paragraph,
    *,
    story: str,
    path: str,
    order: int,
    previous_hash: str,
    in_abstract: bool,
    in_references: bool,
    format_context: FormatContext,
    nested_table: bool = False,
) -> DocumentItem:
    text = paragraph.text
    style = paragraph.style
    style_name = style.name if style is not None else None
    style_id = style.style_id if style is not None else None
    role, confidence = classify_paragraph(
        text,
        style_name,
        order=order,
        in_abstract=in_abstract,
        in_references=in_references,
        story=story,
    )
    risks = paragraph_risks(paragraph, nested_table=nested_table)
    item = DocumentItem(
        stable_id=_stable_id(story, path, text, previous_hash),
        story=story,
        path=path,
        order=order,
        text_preview=normalized_text(text)[:80],
        exact_text_hash=text_hash(text),
        normalized_text_hash=text_hash(normalized_text(text)),
        style_id=style_id,
        style_name=style_name,
        semantic_role=role,
        confidence=confidence,
        risks=risks,
        effective_format=snapshot_paragraph_format(paragraph, format_context),
    )
    items.append(item)
    return item


def inspect_document(path: Path, *, source_filename: str | None = None) -> DocumentInspection:
    document = Document(str(path))
    format_context = build_format_context(document)
    items: list[DocumentItem] = []
    previous_hash = "root"
    in_abstract = False
    in_references = False
    backmatter_role = None

    body_entries = list(_iter_body(document))
    # A recognizable thesis cover may also contain declarations/signatures/spine
    # instructions. Never treat these as ordinary body text before the first abstract.
    first_abstract = next(
        (
            i
            for i, (_, p, _) in enumerate(body_entries)
            if division_role(p.text.strip()) == "abstract_heading"
        ),
        None,
    )
    has_cover = first_abstract is not None and any(
        p.text.strip() in {"博士/硕士学位论文", "博士学位论文", "硕士学位论文"}
        for _, p, _ in body_entries[:first_abstract]
    )
    for order, (item_path, paragraph, nested) in enumerate(body_entries):
        story = "table" if item_path.startswith("table") else "body"
        stripped = paragraph.text.strip()
        style_name = paragraph.style.name if paragraph.style is not None else ""
        preliminary_role, _ = classify_paragraph(stripped, style_name, order=order, story=story)
        division = division_role(stripped) if story == "body" else None
        if (
            preliminary_role.startswith("toc_entry_")
            or preliminary_role == "figure_table_list_entry"
        ):
            division = None
        if division == "abstract_heading":
            in_abstract = True
            in_references = False
            backmatter_role = None
        elif story == "body" and (
            KEYWORDS.match(stripped) or division or preliminary_role.startswith("heading_")
        ):
            in_abstract = False
        if division == "bibliography_heading":
            in_references = True
            in_abstract = False
            backmatter_role = None
        elif division and division != "abstract_heading":
            in_references = False
            backmatter_role = {
                "appendix_heading": "appendix_body",
                "acknowledgments_heading": "acknowledgments_body",
                "cv_heading": "cv_body",
            }.get(division)
        elif story == "body" and preliminary_role.startswith("heading_"):
            in_references = False
        item = _append_item(
            items,
            paragraph,
            story=story,
            path=item_path,
            order=order,
            previous_hash=previous_hash,
            in_abstract=in_abstract and division != "abstract_heading",
            in_references=in_references,
            format_context=format_context,
            nested_table=nested,
        )
        if (
            has_cover
            and first_abstract is not None
            and order < first_abstract
            and story == "body"
            and item.semantic_role
            not in {
                "empty",
                "degree_label",
            }
        ):
            item.semantic_role = "cover_frontmatter"
        if backmatter_role and item.semantic_role == "body" and story == "body":
            item.semantic_role = backmatter_role
        # Notes belong to a nearby caption, not arbitrary prose beginning with 注.
        previous_role = next(
            (
                entry.semantic_role
                for entry in reversed(items[:-1])
                if entry.semantic_role != "empty" and entry.story == "body"
            ),
            None,
        )
        if (
            story == "body"
            and re.match(r"^注[：:]", stripped)
            and previous_role
            in {
                "figure_caption",
                "figure_caption_en",
                "table_caption",
                "table_caption_en",
            }
        ):
            item.semantic_role = "caption_note"
            item.confidence = 0.96
        if item.risks.formula and item.semantic_role in {"body", "abstract_body", "appendix_body"}:
            item.semantic_role = "formula_paragraph"
        if backmatter_role == "acknowledgments_body" and re.fullmatch(
            r"\s*\d{0,4}\s*年\s*\d{0,2}\s*月\s*(?:\d{0,2}\s*日)?\s*", stripped
        ):
            item.semantic_role = "date_line"
        previous_hash = item.exact_text_hash

    order = len(items)
    seen_headers: set[int] = set()
    seen_footers: set[int] = set()
    for section_index, section in enumerate(document.sections):
        for story, container, seen in (
            ("header", section.header, seen_headers),
            ("footer", section.footer, seen_footers),
        ):
            element_id = id(container._element)
            if element_id in seen:
                continue
            seen.add(element_id)
            for paragraph_index, paragraph in enumerate(container.paragraphs):
                item = _append_item(
                    items,
                    paragraph,
                    story=story,
                    path=f"section[{section_index}]/{story}/p[{paragraph_index}]",
                    order=order,
                    previous_hash=previous_hash,
                    in_abstract=False,
                    in_references=False,
                    format_context=format_context,
                )
                item.section_index = section_index
                previous_hash = item.exact_text_hash
                order += 1

    role_counts = Counter(item.semantic_role for item in items if item.semantic_role != "empty")
    risk_counts: Counter[str] = Counter()
    for item in items:
        for risk, enabled in item.risks.model_dump().items():
            if enabled:
                risk_counts[risk] += 1
    summary = InspectionSummary(
        paragraphs=sum(item.story == "body" for item in items),
        table_paragraphs=sum(item.story == "table" for item in items),
        headers=sum(item.story == "header" for item in items),
        footers=sum(item.story == "footer" for item in items),
        protected_items=sum(item.risks.protected for item in items),
        role_counts=dict(role_counts),
        risk_counts=dict(risk_counts),
    )
    advanced = snapshot_advanced_word(
        document,
        [item for item in items if item.story in {"body", "table"}],
        [paragraph for _, paragraph, _ in body_entries],
    )
    warnings = []
    if summary.protected_items:
        warnings.append(
            f"检测到 {summary.protected_items} 个受保护对象；复核确认后可排版安全层，"
            "对象内部 XML 保持不变"
        )
    warnings.extend(advanced.cross_reference_issues)
    return DocumentInspection(
        source_sha256=sha256_file(path),
        source_filename=source_filename or path.name,
        items=items,
        sections=snapshot_sections(document),
        summary=summary,
        advanced=advanced,
        warnings=warnings,
    )
