from pathlib import Path

import pytest
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt
from paper_setting_core.documents import inspect_document, validate_docx_package
from paper_setting_core.formatting import apply_plan
from paper_setting_core.planning import generate_plan
from paper_setting_core.rulepacks import default_rule_pack
from paper_setting_core.rulepacks.models import RulePack
from paper_setting_core.verification import snapshot_package, validate_result


def test_inspect_plan_apply_preserves_content(sample_docx: Path, tmp_path: Path) -> None:
    validate_docx_package(sample_docx)
    inspection = inspect_document(sample_docx)
    assert inspection.summary.paragraphs >= 10
    assert inspection.summary.role_counts["heading_1"] == 1
    assert inspection.summary.role_counts["reference_entry"] == 1

    plan = generate_plan(inspection, default_rule_pack())
    approved = {
        operation.operation_id for operation in plan.operations if operation.status == "proposed"
    }
    output = tmp_path / "formatted.docx"
    operations = apply_plan(sample_docx, output, plan, approved)
    report = validate_result(
        sample_docx,
        output,
        operations,
        already_compliant=plan.summary.get("compliant", 0),
    )

    assert report.integrity_ok, report.differences
    assert report.operation_counts["applied"] > 0
    assert report.total_rule_targets == len(operations) + report.already_compliant
    assert Document(output).paragraphs[0].text == "Agent 论文排版工具研究"
    assert snapshot_package(sample_docx).media_hashes == snapshot_package(output).media_hashes
    result_plan = generate_plan(inspect_document(output), default_rule_pack())
    assert result_plan.operations == []
    assert result_plan.summary["compliant"] > 0


def test_field_paragraph_is_protected(tmp_path: Path) -> None:
    source = tmp_path / "field.docx"
    document = Document()
    paragraph = document.add_paragraph("目录页码 ")
    begin = OxmlElement("w:fldChar")
    begin.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fldCharType", "begin")
    instruction = OxmlElement("w:instrText")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fldCharType", "end")
    paragraph.add_run()._r.append(begin)
    paragraph.add_run()._r.append(instruction)
    paragraph.add_run()._r.append(end)
    document.save(source)

    inspection = inspect_document(source)
    protected = next(item for item in inspection.items if item.risks.field)
    assert protected.risks.revision is False
    plan = generate_plan(inspection, default_rule_pack())
    operation = next(item for item in plan.operations if item.target_id == protected.stable_id)
    assert operation.status == "manual_review"
    assert operation.risk == "high"
    assert operation.execution_scope == "preserve_protected_content"

    skipped_output = tmp_path / "field-unconfirmed.docx"
    skipped = apply_plan(source, skipped_output, plan, {operation.operation_id})
    skipped_operation = next(
        item for item in skipped if item.operation_id == operation.operation_id
    )
    assert skipped_operation.status == "skipped"

    output = tmp_path / "field-confirmed.docx"
    executed = apply_plan(
        source,
        output,
        plan,
        {operation.operation_id},
        confirmed_manual_operation_ids={operation.operation_id},
    )
    applied = next(item for item in executed if item.operation_id == operation.operation_id)
    report = validate_result(source, output, executed)
    assert applied.status == "applied"
    assert "已保留受保护对象的内部 XML" in applied.reasons
    assert report.integrity_ok, report.differences
    assert report.checks["field_instructions_unchanged"] is True


def test_nested_and_merged_tables_are_protected(tmp_path: Path) -> None:
    source = tmp_path / "complex-table.docx"
    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "合并单元格"
    table.cell(1, 0).add_table(rows=1, cols=1).cell(0, 0).text = "嵌套表格"
    document.save(source)

    inspection = inspect_document(source)
    risky_items = [item for item in inspection.items if item.risks.nested_table]
    assert risky_items
    risky_ids = {item.stable_id for item in risky_items}
    plan = generate_plan(inspection, default_rule_pack())
    protected_operations = [item for item in plan.operations if item.target_id in risky_ids]
    assert protected_operations
    assert all(item.status == "manual_review" for item in protected_operations)
    page_operation = next(
        operation
        for operation in plan.operations
        if operation.operation_type == "apply_page_format"
    )
    original_widths = [
        int(column.get(qn("w:w"), "0"))
        for column in table._tbl.xpath("./w:tblGrid/w:gridCol")
    ]
    output = tmp_path / "complex-table-formatted.docx"
    executed = apply_plan(source, output, plan, {page_operation.operation_id})
    applied = next(
        operation
        for operation in executed
        if operation.operation_id == page_operation.operation_id
    )
    assert applied.result["tables_fitted"] == 0
    result = Document(output)
    result_widths = [
        int(column.get(qn("w:w"), "0"))
        for column in result.tables[0]._tbl.xpath("./w:tblGrid/w:gridCol")
    ]
    assert result_widths == original_widths


def test_page_format_fits_plain_wide_table_to_content_width(tmp_path: Path) -> None:
    source = tmp_path / "wide-table.docx"
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11)
    section.page_height = Inches(8.5)
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)
    table = document.add_table(rows=2, cols=3)
    for index, cell in enumerate(table.rows[0].cells):
        cell.text = f"列 {index + 1}"
    document.save(source)

    rule_pack = default_rule_pack()
    plan = generate_plan(inspect_document(source), rule_pack)
    page_operation = next(
        operation
        for operation in plan.operations
        if operation.operation_type == "apply_page_format"
    )
    output = tmp_path / "wide-table-formatted.docx"
    executed = apply_plan(source, output, plan, {page_operation.operation_id})

    applied = next(
        operation
        for operation in executed
        if operation.operation_id == page_operation.operation_id
    )
    assert applied.result["tables_fitted"] == 1
    result = Document(output)
    widths = [
        int(column.get(qn("w:w"), "0"))
        for column in result.tables[0]._tbl.xpath("./w:tblGrid/w:gridCol")
    ]
    expected_width = int(
        Mm(
            rule_pack.page.width_mm
            - rule_pack.page.margin_left_mm
            - rule_pack.page.margin_right_mm
        ).twips
    )
    assert sum(widths) == expected_width
    report = validate_result(source, output, executed)
    assert report.integrity_ok, report.differences


def test_plan_omits_already_compliant_formatting(tmp_path: Path) -> None:
    source = tmp_path / "compliant.docx"
    document = Document()
    rule_pack = default_rule_pack()
    page = rule_pack.page
    section = document.sections[0]
    section.page_width = Mm(page.width_mm)
    section.page_height = Mm(page.height_mm)
    section.top_margin = Mm(page.margin_top_mm)
    section.bottom_margin = Mm(page.margin_bottom_mm)
    section.left_margin = Mm(page.margin_left_mm)
    section.right_margin = Mm(page.margin_right_mm)
    section.header_distance = Mm(page.header_distance_mm)
    section.footer_distance = Mm(page.footer_distance_mm)

    paragraph = document.add_paragraph("这是一段已经符合规则要求的普通正文内容。")
    paragraph.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.first_line_indent = Pt(24)
    paragraph.paragraph_format.line_spacing = 1.5
    run = paragraph.runs[0]
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    run._r.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")
    document.save(source)

    inspection = inspect_document(source)
    body = next(item for item in inspection.items if item.semantic_role == "body")
    assert body.effective_format.paragraph["alignment"] == "justify"
    assert body.effective_format.character["east_asia_font"] == "宋体"

    plan = generate_plan(inspection, rule_pack)
    assert plan.operations == []
    assert plan.summary == {"proposed": 0, "manual_review": 0, "compliant": 2}


def test_plan_reports_only_changed_effective_fields(tmp_path: Path) -> None:
    source = tmp_path / "one-mismatch.docx"
    document = Document()
    rule_pack = default_rule_pack()
    page = rule_pack.page
    section = document.sections[0]
    section.page_width = Mm(page.width_mm)
    section.page_height = Mm(page.height_mm)
    section.top_margin = Mm(page.margin_top_mm)
    section.bottom_margin = Mm(page.margin_bottom_mm)
    section.left_margin = Mm(page.margin_left_mm)
    section.right_margin = Mm(page.margin_right_mm)
    section.header_distance = Mm(page.header_distance_mm)
    section.footer_distance = Mm(page.footer_distance_mm)

    paragraph = document.add_paragraph("这是一段只有字号不符合规则要求的普通正文内容。")
    paragraph.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.first_line_indent = Pt(24)
    paragraph.paragraph_format.line_spacing = 1.5
    run = paragraph.runs[0]
    run.font.name = "Times New Roman"
    run.font.size = Pt(11)
    run._r.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")
    document.save(source)

    plan = generate_plan(inspect_document(source), rule_pack)
    assert len(plan.operations) == 1
    operation = plan.operations[0]
    assert operation.semantic_role == "body"
    assert operation.changed_fields == ["character.size_pt"]
    assert operation.before["character"]["size_pt"] == 11.0


def test_inspection_resolves_doc_defaults_and_theme_fonts(tmp_path: Path) -> None:
    source = tmp_path / "defaults.docx"
    document = Document()
    document.add_paragraph("这是一段使用 Word 文档默认格式的普通正文。")
    document.save(source)

    inspection = inspect_document(source)
    body = next(item for item in inspection.items if item.semantic_role == "body")

    assert body.effective_format.paragraph["space_after_pt"] == 10.0
    assert body.effective_format.paragraph["line_spacing"] == 1.15
    assert body.effective_format.paragraph["line_spacing_mode"] == "multiple"
    assert body.effective_format.character["east_asia_font"] == "宋体"
    assert body.effective_format.character["latin_font"] == "Cambria"
    assert body.effective_format.character["size_pt"] == 11.0


@pytest.mark.parametrize(("mode", "value"), [("exact", 20.0), ("at_least", 18.0)])
def test_apply_and_verify_point_based_line_spacing(
    tmp_path: Path,
    mode: str,
    value: float,
) -> None:
    source = tmp_path / f"{mode}.docx"
    output = tmp_path / f"{mode}-formatted.docx"
    document = Document()
    document.add_paragraph("这是一段需要设置特定行距模式的普通正文内容。")
    document.save(source)

    payload = default_rule_pack().model_dump(mode="json")
    body_rule = next(rule for rule in payload["roles"] if rule["role"] == "body")
    body_rule["paragraph"]["line_spacing"] = value
    body_rule["paragraph"]["line_spacing_mode"] = mode
    rule_pack = RulePack.model_validate(payload)

    plan = generate_plan(inspect_document(source), rule_pack)
    approved = {
        operation.operation_id for operation in plan.operations if operation.status == "proposed"
    }
    apply_plan(source, output, plan, approved)

    result = inspect_document(output)
    body = next(item for item in result.items if item.semantic_role == "body")
    assert body.effective_format.paragraph["line_spacing_mode"] == mode
    assert body.effective_format.paragraph["line_spacing"] == value
    assert generate_plan(result, rule_pack).operations == []
