from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient
from paper_setting_api.main import create_app
from paper_setting_runtime.config import Settings


def _settings(tmp_path: Path, database_name: str) -> Settings:
    return Settings(
        env="test",
        database_url=f"sqlite:///{tmp_path / database_name}",
        data_dir=tmp_path / "data",
        allowed_origins=["http://testserver"],
        allowed_hosts=["testserver"],
        web_dist_dir=tmp_path / "missing-web",
    )


def _create_job(client: TestClient, document: Path, template: Path | None = None) -> dict:
    with document.open("rb") as source_handle:
        files: dict[str, tuple[str, object, str]] = {
            "document": (
                document.name,
                source_handle,
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        }
        if template is None:
            response = client.post(
                "/api/v1/jobs",
                files=files,
                data={"rule_pack_id": "zh-thesis-default", "render_preview": "false"},
            )
        else:
            with template.open("rb") as template_handle:
                files["template_document"] = (
                    template.name,
                    template_handle,
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
                response = client.post(
                    "/api/v1/jobs",
                    files=files,
                    data={"rule_pack_id": "zh-thesis-default", "render_preview": "false"},
                )
    assert response.status_code == 202, response.text
    return response.json()


def test_job_history_paginates_and_terminal_job_cleanup_is_safe(
    sample_docx: Path, tmp_path: Path
) -> None:
    app = create_app(_settings(tmp_path, "history.db"))
    with TestClient(app, base_url="http://testserver") as client:
        jobs = [_create_job(client, sample_docx) for _ in range(3)]
        active_delete = client.delete(f"/api/v1/jobs/{jobs[-1]['id']}")
        assert active_delete.status_code == 409
        assert active_delete.json()["code"] == "JOB_DELETE_CONFLICT"
        for job in jobs:
            assert client.post(f"/api/v1/jobs/{job['id']}/cancel").status_code == 200

        first = client.get("/api/v1/jobs", params={"limit": 2})
        assert first.status_code == 200
        assert len(first.json()["items"]) == 2
        assert first.json()["next_cursor"]
        second = client.get(
            "/api/v1/jobs",
            params={"limit": 2, "cursor": first.json()["next_cursor"]},
        )
        assert second.status_code == 200
        assert len(second.json()["items"]) == 1
        assert {
            item["id"] for item in [*first.json()["items"], *second.json()["items"]]
        } == {item["id"] for item in jobs}

        deleted_id = jobs[0]["id"]
        job_directory = app.state.settings.jobs_dir / deleted_id
        assert job_directory.exists()
        assert client.delete(f"/api/v1/jobs/{deleted_id}").status_code == 204
        assert not job_directory.exists()
        assert client.get(f"/api/v1/jobs/{deleted_id}").status_code == 404


def test_job_pipeline_composes_manuscript_into_uploaded_template(tmp_path: Path) -> None:
    manuscript = tmp_path / "paper.docx"
    document = Document()
    document.add_paragraph("论文排版台验证", style="Title")
    document.add_paragraph("摘要")
    document.add_paragraph("这是用于验证智能映射的摘要。")
    document.add_paragraph("关键词：模板；正文")
    document.add_paragraph("第一章 绪论", style="Heading 1")
    document.add_paragraph("这是应当注入模板的正文。")
    document.save(manuscript)

    template = tmp_path / "template.docx"
    destination = Document()
    destination.add_paragraph("{{paper.title}}")
    destination.add_paragraph("摘要：{{paper.abstract}}")
    destination.add_paragraph("关键词：{{paper.keywords}}")
    destination.add_paragraph("{{document.body}}")
    destination.save(template)

    app = create_app(_settings(tmp_path, "composition.db"))
    with TestClient(app, base_url="http://testserver") as client:
        created = _create_job(client, manuscript, template)
        assert created["mode"] == "template"
        assert created["template_filename"] == "template.docx"
        assert app.state.job_service.process_next("composition-worker")

        plan_response = client.get(f"/api/v1/jobs/{created['id']}/template-plan")
        assert plan_response.status_code == 200
        plan = plan_response.json()
        assert plan["plan_kind"] == "manuscript_composition"
        assert plan["can_generate"] is True
        operation_ids = [
            *[item["operation_id"] for item in plan["fill"]["operations"]],
            plan["body_injection"]["operation_id"],
            *[item["operation_id"] for item in plan["structure"]["operations"]],
        ]
        approved = client.post(
            f"/api/v1/jobs/{created['id']}/approve",
            json={
                "plan_version": plan["plan_version"],
                "approved_operation_ids": operation_ids,
            },
        )
        assert approved.status_code == 200, approved.text
        assert app.state.job_service.process_next("composition-worker")
        completed = client.get(f"/api/v1/jobs/{created['id']}").json()
        assert completed["status"] == "completed", completed
        output = client.get(f"/api/v1/jobs/{created['id']}/artifacts/formatted_docx")
        assert output.status_code == 200
        output_path = tmp_path / "result.docx"
        output_path.write_bytes(output.content)
        text = "\n".join(paragraph.text for paragraph in Document(output_path).paragraphs)
        assert "论文排版台验证" in text
        assert "这是用于验证智能映射的摘要" in text
        assert "这是应当注入模板的正文" in text
