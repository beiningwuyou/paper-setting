from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from fastapi.testclient import TestClient
from lxml import etree
from paper_setting_api.main import create_app
from paper_setting_core.advanced_word.citations import normalize_numeric_citations
from paper_setting_core.documents import inspect_document
from paper_setting_core.errors import CrossReferenceInvalidError
from paper_setting_core.formatting import apply_plan
from paper_setting_core.planning import generate_plan
from paper_setting_core.rulepacks import deep_rule_pack
from paper_setting_core.rulepacks.models import AdvancedWordFormat, NotesFormat
from paper_setting_core.verification import validate_result
from paper_setting_runtime.config import Settings


def _append_formula(paragraph: object, text: str) -> None:
    math = OxmlElement("m:oMath")
    run = OxmlElement("m:r")
    math_text = OxmlElement("m:t")
    math_text.text = text
    run.append(math_text)
    math.append(run)
    paragraph._p.append(math)  # type: ignore[attr-defined]


def _create_deep_document(path: Path) -> None:
    document = Document()
    document.add_paragraph("深层 Word 自动化研究", style="Title")
    document.add_page_break()
    document.add_paragraph("第一章 绪论", style="Heading 1")
    document.add_paragraph("研究结论已由多项工作证实[3，1, 2, 2]。")
    document.add_paragraph("1.1 研究背景", style="Heading 2")
    formula_paragraph = document.add_paragraph("公式如下：")
    _append_formula(formula_paragraph, "x+1")
    document.add_paragraph("图 1-1 系统处理流程")
    document.add_paragraph("表 1-1 核心功能清单")
    document.add_paragraph("参考文献")
    document.add_paragraph("[1] 张三. 论文排版方法研究[J]. 示例期刊, 2026.")
    document.save(path)


def _create_cross_reference_document(path: Path) -> None:
    document = Document()
    document.add_page_break()
    document.add_paragraph(
        "如[[xref:fig-system]]所示，详见第[[xref-page:fig-system]]页；"
        "数据见[[xref:tab-results]]。"
    )
    figure = document.add_paragraph()
    figure.add_run("[[xref-target:")
    figure.add_run("fig-system]]系统架构")
    document.add_paragraph("[[xref-target:tab-results]]实验结果")
    document.save(path)


def _append_ref_field(paragraph: object, number: int) -> None:
    for field_type, instruction, visible in (
        ("begin", None, None),
        (None, f" REF _Ref{number} \\r \\h ", None),
        ("separate", None, None),
        (None, None, f"〔{number}〕"),
        ("end", None, None),
    ):
        run = OxmlElement("w:r")
        if field_type:
            field = OxmlElement("w:fldChar")
            field.set(qn("w:fldCharType"), field_type)
            run.append(field)
        elif instruction:
            node = OxmlElement("w:instrText")
            node.text = instruction
            run.append(node)
        else:
            node = OxmlElement("w:t")
            node.text = visible
            run.append(node)
        paragraph._p.append(run)  # type: ignore[attr-defined]


def test_inline_reference_fields_become_real_footnotes_and_bibliography_is_removed(
    tmp_path: Path,
) -> None:
    source = tmp_path / "notes-source.docx"
    document = Document()
    paragraph = document.add_paragraph("这一结论已被论证")
    _append_ref_field(paragraph, 1)
    document.add_paragraph("。")
    document.add_paragraph("参考文献", style="Heading 1")
    document.add_paragraph("[1] 张三. 哲学分析方法研究[M]. 北京: 示例出版社, 2026.")
    document.save(source)

    rule_pack = deep_rule_pack().model_copy(
        update={
            "advanced": AdvancedWordFormat(
                notes=NotesFormat(
                    enabled=True,
                    delete_bibliography=True,
                    numbering_restart="each_page",
                    number_format="decimal_enclosed_circle",
                )
            )
        }
    )
    inspection = inspect_document(source)
    assert inspection.advanced.inline_note_candidate_count == 1
    assert inspection.advanced.bibliography_entry_count == 1
    plan = generate_plan(inspection, rule_pack)
    operation = next(
        operation
        for operation in plan.operations
        if operation.operation_type == "convert_notes_to_footnotes"
    )
    output = tmp_path / "notes-result.docx"
    executed = apply_plan(
        source,
        output,
        plan,
        {operation.operation_id},
        confirmed_manual_operation_ids={operation.operation_id},
    )

    executed_operation = next(
        item for item in executed if item.operation_type == "convert_notes_to_footnotes"
    )
    assert executed_operation.status == "applied"
    assert executed_operation.result["converted_inline_citations"] == 1
    assert executed_operation.result["removed_bibliography_entries"] == 1
    assert executed_operation.result["added_parts"] == ["word/footnotes.xml"]
    report = validate_result(source, output, executed)
    assert report.integrity_ok, report.differences

    result = Document(output)
    assert all(paragraph.text.strip() != "参考文献" for paragraph in result.paragraphs)
    result_inspection = inspect_document(output)
    assert result_inspection.advanced.footnote_count == 1
    assert result_inspection.advanced.inline_note_candidate_count == 0
    assert result_inspection.advanced.bibliography_entry_count == 0
    assert result_inspection.advanced.footnote_numbering_restarts == ["each_page"]
    assert result_inspection.advanced.footnote_number_formats == [
        "decimal_enclosed_circle"
    ]
    second_plan = generate_plan(result_inspection, rule_pack)
    assert not any(
        item.operation_type == "convert_notes_to_footnotes"
        for item in second_plan.operations
    )
    with ZipFile(output) as archive:
        footnotes = etree.fromstring(archive.read("word/footnotes.xml"))
        document_xml = etree.fromstring(archive.read("word/document.xml"))
    namespaces = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    section_note_properties = document_xml.xpath(
        ".//w:sectPr/w:footnotePr",
        namespaces=namespaces,
    )
    assert len(section_note_properties) == 1
    assert section_note_properties[0].xpath(
        "./w:numRestart[@w:val='eachPage']",
        namespaces=namespaces,
    )
    assert section_note_properties[0].xpath(
        "./w:numFmt[@w:val='decimalEnclosedCircle']",
        namespaces=namespaces,
    )
    assert section_note_properties[0].xpath(
        "./w:numStart[@w:val='1']",
        namespaces=namespaces,
    )
    note_paragraphs = footnotes.xpath(
        "./w:footnote[not(@w:type)]/w:p",
        namespaces=namespaces,
    )
    assert len(note_paragraphs) == 1
    assert note_paragraphs[0].xpath(
        "./w:pPr/w:jc[@w:val='left']",
        namespaces=namespaces,
    )
    assert note_paragraphs[0].xpath(
        "./w:pPr/w:wordWrap[@w:val='off']",
        namespaces=namespaces,
    )


def test_numeric_citations_are_sorted_deduplicated_and_collapsed() -> None:
    assert normalize_numeric_citations("见[3，1, 2, 2, 5]") == "见[1–3, 5]"
    assert normalize_numeric_citations("保留[1-3]") == "保留[1-3]"


def test_explicit_cross_reference_markers_create_fields_and_bookmarks(
    tmp_path: Path,
) -> None:
    source = tmp_path / "cross-reference-source.docx"
    _create_cross_reference_document(source)
    rule_pack = deep_rule_pack()
    inspection = inspect_document(source)
    assert inspection.advanced.cross_reference_target_count == 2
    assert inspection.advanced.cross_reference_marker_count == 2
    assert inspection.advanced.cross_reference_page_marker_count == 1
    assert inspection.advanced.cross_reference_issues == []

    plan = generate_plan(inspection, rule_pack)
    operation = next(
        item for item in plan.operations if item.operation_type == "create_cross_references"
    )
    output = tmp_path / "cross-reference-result.docx"
    executed = apply_plan(
        source,
        output,
        plan,
        {operation.operation_id},
        confirmed_manual_operation_ids={operation.operation_id},
    )
    result_operation = next(
        item for item in executed if item.operation_type == "create_cross_references"
    )
    assert result_operation.status == "applied"
    assert result_operation.result["targets"] == 2
    assert result_operation.result["references"] == 2
    assert result_operation.result["page_references"] == 1
    report = validate_result(source, output, executed)
    assert report.integrity_ok, report.differences

    with ZipFile(output) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
    namespaces = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    instructions = [
        " ".join(str(value).split())
        for value in root.xpath(".//w:instrText/text()", namespaces=namespaces)
    ]
    assert sum(value.startswith("SEQ Figure") for value in instructions) == 1
    assert sum(value.startswith("SEQ Table") for value in instructions) == 1
    assert sum(value.startswith("REF _PSX_") for value in instructions) == 2
    assert sum(value.startswith("PAGEREF _PSX_") for value in instructions) == 1
    assert len(root.xpath(".//w:bookmarkStart", namespaces=namespaces)) == 2

    result = Document(output)
    visible = "".join(paragraph.text for paragraph in result.paragraphs)
    assert "[[xref" not in visible
    assert "图 1" in visible
    assert "表 1" in visible
    result_inspection = inspect_document(output)
    assert result_inspection.advanced.cross_reference_target_count == 0
    assert result_inspection.advanced.cross_reference_marker_count == 0
    assert result_inspection.advanced.cross_reference_existing_field_count == 5
    second_plan = generate_plan(result_inspection, rule_pack)
    assert not any(
        item.operation_type in {"apply_automatic_numbering", "create_cross_references"}
        for item in second_plan.operations
    )


def test_unresolved_cross_reference_blocks_plan(tmp_path: Path) -> None:
    source = tmp_path / "unresolved-cross-reference.docx"
    document = Document()
    document.add_paragraph("参见[[xref:fig-missing]]。")
    document.save(source)
    inspection = inspect_document(source)
    assert inspection.advanced.cross_reference_issues
    with pytest.raises(CrossReferenceInvalidError):
        generate_plan(inspection, deep_rule_pack())


def test_malformed_cross_reference_marker_blocks_plan(tmp_path: Path) -> None:
    source = tmp_path / "malformed-cross-reference.docx"
    document = Document()
    document.add_paragraph("参见[[XREF:fig-system]]。")
    document.save(source)
    inspection = inspect_document(source)
    assert inspection.advanced.cross_reference_issues == [
        "无法识别交叉引用标记：[[XREF:fig-system]]"
    ]
    with pytest.raises(CrossReferenceInvalidError):
        generate_plan(inspection, deep_rule_pack())


@pytest.mark.parametrize(
    ("caption", "expected_issue"),
    [
        (
            "题注：[[xref-target:fig-system]]系统架构",
            "目标 fig-system 必须位于题注段落开头",
        ),
        ("[[xref-target:fig-system]]", "目标 fig-system 缺少题注文字"),
    ],
)
def test_invalid_cross_reference_target_shape_blocks_plan(
    tmp_path: Path,
    caption: str,
    expected_issue: str,
) -> None:
    source = tmp_path / "invalid-cross-reference-target.docx"
    document = Document()
    document.add_paragraph("参见[[xref:fig-system]]。")
    document.add_paragraph(caption)
    document.save(source)
    inspection = inspect_document(source)
    assert expected_issue in inspection.advanced.cross_reference_issues
    with pytest.raises(CrossReferenceInvalidError):
        generate_plan(inspection, deep_rule_pack())


def test_cross_reference_targets_replace_existing_caption_numbers(
    tmp_path: Path,
) -> None:
    source = tmp_path / "numbered-cross-reference.docx"
    document = Document()
    document.add_paragraph("1 绪论", style="Heading 1")
    document.add_paragraph("参见[[xref:fig-system]]。")
    caption_style = document.styles.add_style(
        "PaperSetting.figure_caption", WD_STYLE_TYPE.PARAGRAPH
    )
    caption = document.add_paragraph(style=caption_style)
    caption.add_run("[[xref-target:fig-system]]图 2-3：系统架构")
    document.save(source)

    rule_pack = deep_rule_pack()
    inspection = inspect_document(source)
    assert inspection.advanced.cross_reference_target_counts == {"figure": 1}
    plan = generate_plan(inspection, rule_pack)
    operations = [
        item
        for item in plan.operations
        if item.operation_type in {"apply_automatic_numbering", "create_cross_references"}
    ]
    assert {item.operation_type for item in operations} == {
        "apply_automatic_numbering",
        "create_cross_references",
    }

    output = tmp_path / "numbered-cross-reference-result.docx"
    operation_ids = {item.operation_id for item in operations}
    executed = apply_plan(
        source,
        output,
        plan,
        operation_ids,
        confirmed_manual_operation_ids=operation_ids,
    )
    report = validate_result(source, output, executed)
    assert report.integrity_ok, report.differences

    result = Document(output)
    result_caption = result.paragraphs[2]
    assert result_caption.text == "图 1 系统架构"
    assert result_caption._p.pPr is not None
    assert result_caption._p.pPr.find(qn("w:numPr")) is None
    numbering_result = next(
        item.result
        for item in executed
        if item.operation_type == "apply_automatic_numbering"
    )
    assert numbering_result["excluded_cross_reference_targets"] == {
        "figure_caption": 1
    }

    result_inspection = inspect_document(output)
    assert result_inspection.advanced.sequence_numbered_role_counts == {
        "figure_caption": 1
    }
    second_plan = generate_plan(result_inspection, rule_pack)
    assert not any(
        item.operation_type in {"apply_automatic_numbering", "create_cross_references"}
        for item in second_plan.operations
    )


def test_cross_references_complete_through_job_worker(tmp_path: Path) -> None:
    source = tmp_path / "cross-reference-job.docx"
    _create_cross_reference_document(source)
    settings = Settings(
        env="test",
        database_url=f"sqlite:///{tmp_path / 'cross-reference-job.db'}",
        data_dir=tmp_path / "data",
        allowed_origins=["http://testserver"],
        allowed_hosts=["testserver"],
        web_dist_dir=tmp_path / "missing-web",
    )
    app = create_app(settings)
    with TestClient(app, base_url="http://testserver") as client:
        with source.open("rb") as handle:
            created = client.post(
                "/api/v1/jobs",
                files={
                    "document": (
                        source.name,
                        handle,
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    )
                },
                data={"rule_pack_id": "zh-thesis-deep", "render_preview": "false"},
            )
        job_id = created.json()["id"]
        assert app.state.job_service.process_next("cross-reference-worker")
        plan = client.get(f"/api/v1/jobs/{job_id}/plan").json()
        operation = next(
            item
            for item in plan["operations"]
            if item["operation_type"] == "create_cross_references"
        )
        rejected = [
            item["operation_id"]
            for item in plan["operations"]
            if item["operation_id"] != operation["operation_id"]
        ]
        approved = client.post(
            f"/api/v1/jobs/{job_id}/approve",
            json={
                "plan_version": plan["plan_version"],
                "approved_operation_ids": [operation["operation_id"]],
                "rejected_operation_ids": rejected,
                "confirmed_manual_operation_ids": [operation["operation_id"]],
            },
        )
        assert approved.status_code == 200, approved.text
        assert app.state.job_service.process_next("cross-reference-worker")
        completed = client.get(f"/api/v1/jobs/{job_id}").json()
        assert completed["status"] == "completed", completed
        report = client.get(f"/api/v1/jobs/{job_id}/report").json()
        assert report["integrity_ok"] is True
        assert report["operation_counts"]["applied"] == 1
        output = client.get(f"/api/v1/jobs/{job_id}/artifacts/formatted_docx")
        generated = tmp_path / "cross-reference-job-result.docx"
        generated.write_bytes(output.content)
        result_inspection = inspect_document(generated)
        assert result_inspection.advanced.cross_reference_existing_field_count == 5


def test_deep_word_operations_are_audited_and_idempotent(tmp_path: Path) -> None:
    source = tmp_path / "deep-source.docx"
    _create_deep_document(source)

    rule_pack = deep_rule_pack()
    inspection = inspect_document(source)
    plan = generate_plan(inspection, rule_pack)
    advanced_types = {
        operation.operation_type
        for operation in plan.operations
        if operation.execution_scope == "controlled_ooxml_rewrite"
    }
    assert advanced_types == {
        "reformat_formulas",
        "normalize_citations",
        "rebuild_toc",
        "apply_automatic_numbering",
    }

    approved = {operation.operation_id for operation in plan.operations}
    confirmed = {
        operation.operation_id
        for operation in plan.operations
        if operation.status == "manual_review"
    }
    output = tmp_path / "deep-formatted.docx"
    executed = apply_plan(
        source,
        output,
        plan,
        approved,
        confirmed_manual_operation_ids=confirmed,
    )
    report = validate_result(source, output, executed)
    assert report.integrity_ok, report.differences
    assert report.operation_counts["applied"] == len(plan.operations)

    result_document = Document(output)
    assert "[1–3]" in "".join(paragraph.text for paragraph in result_document.paragraphs)
    heading = next(
        paragraph for paragraph in result_document.paragraphs if "绪论" in paragraph.text
    )
    assert heading.text == "绪论"
    result_inspection = inspect_document(output)
    assert result_inspection.advanced.formula_math_font == "Cambria Math"
    assert result_inspection.advanced.citation_candidate_count == 0
    assert result_inspection.advanced.update_fields_on_open is True
    assert any("TOC" in value for value in result_inspection.advanced.toc_instructions)
    assert result_inspection.advanced.numbered_role_counts["heading_1"] == 1
    assert result_inspection.advanced.numbered_role_counts["heading_2"] == 1
    assert result_inspection.advanced.numbered_role_counts["reference_entry"] == 1
    assert result_inspection.advanced.numbered_role_counts["figure_caption"] == 1
    assert result_inspection.advanced.numbered_role_counts["table_caption"] == 1

    settings_math_font = result_document.settings.element.xpath("./m:mathPr/m:mathFont")
    assert settings_math_font[0].get(qn("m:val")) == "Cambria Math"
    second_plan = generate_plan(result_inspection, rule_pack)
    assert not any(
        operation.execution_scope == "controlled_ooxml_rewrite"
        for operation in second_plan.operations
    )


def test_deep_word_rule_pack_completes_through_api_and_worker(tmp_path: Path) -> None:
    source = tmp_path / "deep-api.docx"
    _create_deep_document(source)
    settings = Settings(
        env="test",
        database_url=f"sqlite:///{tmp_path / 'deep-api.db'}",
        data_dir=tmp_path / "data",
        allowed_origins=["http://testserver"],
        allowed_hosts=["testserver"],
        web_dist_dir=tmp_path / "missing-web",
    )
    app = create_app(settings)
    with TestClient(app, base_url="http://testserver") as client:
        packs = client.get("/api/v1/rule-packs").json()
        assert "zh-thesis-deep" in {pack["id"] for pack in packs}
        with source.open("rb") as handle:
            created = client.post(
                "/api/v1/jobs",
                files={
                    "document": (
                        source.name,
                        handle,
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    )
                },
                data={"rule_pack_id": "zh-thesis-deep", "render_preview": "false"},
            )
        job_id = created.json()["id"]
        assert app.state.job_service.process_next("deep-test-worker")
        plan = client.get(f"/api/v1/jobs/{job_id}/plan").json()
        approved = [operation["operation_id"] for operation in plan["operations"]]
        manual = [
            operation["operation_id"]
            for operation in plan["operations"]
            if operation["status"] == "manual_review"
        ]
        approval = client.post(
            f"/api/v1/jobs/{job_id}/approve",
            json={
                "plan_version": plan["plan_version"],
                "approved_operation_ids": approved,
                "rejected_operation_ids": [],
                "confirmed_manual_operation_ids": manual,
            },
        )
        assert approval.status_code == 200, approval.text
        assert app.state.job_service.process_next("deep-test-worker")
        completed = client.get(f"/api/v1/jobs/{job_id}").json()
        assert completed["status"] == "completed", completed
        report = client.get(f"/api/v1/jobs/{job_id}/report").json()
        assert report["integrity_ok"] is True
        assert report["compliance_rate"] == 1.0
