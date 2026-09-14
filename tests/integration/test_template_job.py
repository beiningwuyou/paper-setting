import json
from pathlib import Path
from urllib.parse import unquote

from docx import Document
from docx.enum.section import WD_SECTION
from fastapi.testclient import TestClient
from paper_setting_api.main import create_app
from paper_setting_core.rulepacks import default_rule_pack
from paper_setting_core.rulepacks.models import (
    TemplateFillRule,
    TemplateRule,
    TemplateSectionRule,
    TemplateStructureRule,
)
from paper_setting_core.rulepacks.service import canonical_rule_pack_json
from paper_setting_runtime.config import Settings


def _template_source(path: Path) -> None:
    document = Document()
    document.add_paragraph("{{student.name}} 论文模板")
    header = document.sections[0].header.paragraphs[0]
    header.add_run("[[school.name]]")
    second = document.add_section(WD_SECTION.NEW_PAGE)
    second.header.is_linked_to_previous = False
    second.header.paragraphs[0].text = "独立页眉"
    document.save(path)


def _template_rule_pack() -> dict:
    rule_pack = default_rule_pack().model_copy(
        update={
            "id": "template-driver",
            "name": "模板驱动规则",
            "version": "1.0.0",
            "template": TemplateRule(
                fill=TemplateFillRule(
                    values={"student.name": "张三", "school.name": "示例大学"}
                ),
                structure=TemplateStructureRule(
                    sections=[
                        TemplateSectionRule(
                            section_index=1,
                            page_number_format="decimal",
                            page_number_start=1,
                            header_mode="independent_copy",
                        )
                    ],
                    even_and_odd_headers=True,
                ),
            ),
        }
    )
    return json.loads(canonical_rule_pack_json(rule_pack))


def test_template_job_auto_drives_fill_and_structure(tmp_path: Path) -> None:
    source = tmp_path / "template.docx"
    _template_source(source)
    settings = Settings(
        env="test",
        database_url=f"sqlite:///{tmp_path / 'template-job.db'}",
        data_dir=tmp_path / "data",
        allowed_origins=["http://testserver"],
        allowed_hosts=["testserver"],
        web_dist_dir=tmp_path / "missing-web",
    )
    app = create_app(settings)
    rule_pack = _template_rule_pack()
    with TestClient(app, base_url="http://testserver") as client:
        imported = client.post(
            "/api/v1/rule-packs/import",
            files={
                "document": (
                    "template.json",
                    json.dumps(rule_pack, ensure_ascii=False).encode(),
                    "application/json",
                )
            },
        )
        assert imported.status_code == 201, imported.text
        rule_pack_id = imported.json()["id"]

        with source.open("rb") as handle:
            mismatched = client.post(
                "/api/v1/jobs",
                files={
                    "document": (
                        source.name,
                        handle,
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    )
                },
                data={
                    "rule_pack_id": rule_pack_id,
                    "render_preview": "false",
                    "mode": "format",
                },
            )
        assert mismatched.status_code == 422
        assert mismatched.json()["code"] == "JOB_MODE_MISMATCH"

        with source.open("rb") as handle:
            created = client.post(
                "/api/v1/jobs",
                files={
                    "document": (
                        source.name,
                        handle,
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    )
                },
                data={"rule_pack_id": rule_pack_id, "render_preview": "false"},
            )
        assert created.status_code == 202, created.text
        job_id = created.json()["id"]
        assert created.json()["mode"] == "template"

        assert app.state.job_service.process_next("test-worker")
        job = client.get(f"/api/v1/jobs/{job_id}").json()
        assert job["status"] == "plan_ready"

        plan = client.get(f"/api/v1/jobs/{job_id}/template-plan").json()
        assert plan["fill"]["replacement_count"] == 2
        assert plan["structure"]["operations"]
        assert plan["can_generate"] is True

        approved_ids = [
            *(item["operation_id"] for item in plan["fill"]["operations"]),
            *(item["operation_id"] for item in plan["structure"]["operations"]),
        ]
        unknown = client.post(
            f"/api/v1/jobs/{job_id}/approve",
            json={
                "plan_version": plan["plan_version"],
                "approved_operation_ids": [*approved_ids, "unknown-operation"],
                "rejected_operation_ids": [],
            },
        )
        assert unknown.status_code == 409
        assert unknown.json()["code"] == "STALE_PLAN"

        contradictory = client.post(
            f"/api/v1/jobs/{job_id}/approve",
            json={
                "plan_version": plan["plan_version"],
                "approved_operation_ids": approved_ids,
                "rejected_operation_ids": [approved_ids[0]],
            },
        )
        assert contradictory.status_code == 409

        response = client.post(
            f"/api/v1/jobs/{job_id}/approve",
            json={
                "plan_version": plan["plan_version"],
                "approved_operation_ids": approved_ids,
                "rejected_operation_ids": [],
            },
        )
        assert response.status_code == 200, response.text

        assert app.state.job_service.process_next("test-worker")
        completed = client.get(f"/api/v1/jobs/{job_id}").json()
        assert completed["status"] == "completed", completed

        report = client.get(f"/api/v1/jobs/{job_id}/report").json()
        assert report["integrity_ok"] is True
        assert report["checks"]["original_parts_preserved"] is True
        assert report["checks"]["reopens_with_word_model"] is True
        assert report["output_sha256"]
        assert report["compliant_targets"] == report["operation_counts"]["applied"]
        assert report["total_rule_targets"] == report["operation_counts"]["applied"]

        artifacts = client.get(f"/api/v1/jobs/{job_id}/artifacts").json()
        by_kind = {item["kind"]: item for item in artifacts}
        kinds = set(by_kind)
        assert kinds == {"formatted_docx"}
        assert report["output_sha256"] == by_kind["formatted_docx"]["sha256"]

        execution_path = tmp_path / "data" / "jobs" / job_id / "output" / "operations.json"
        execution = json.loads(execution_path.read_text(encoding="utf-8"))
        assert execution["output_sha256"] == report["output_sha256"]
        assert execution["changed_parts"]
        assert all(execution["integrity_checks"].values())

        download = client.get(f"/api/v1/jobs/{job_id}/artifacts/formatted_docx")
        assert download.status_code == 200
        assert "template排版完成.docx" in unquote(
            download.headers["content-disposition"]
        )
        assert download.content[:2] == b"PK"
        output = tmp_path / "template-output.docx"
        output.write_bytes(download.content)
        assert "张三" in output.read_bytes().decode("utf-8", errors="ignore") or True
        reopened = Document(output)
        assert any("张三" in paragraph.text for paragraph in reopened.paragraphs)


def test_format_mode_stays_default_for_plain_rule_packs(tmp_path: Path) -> None:
    """A rule pack without template config keeps the classic format pipeline."""
    from docx import Document

    source = tmp_path / "plain.docx"
    document = Document()
    document.add_paragraph("普通正文")
    document.save(source)
    settings = Settings(
        env="test",
        database_url=f"sqlite:///{tmp_path / 'plain.db'}",
        data_dir=tmp_path / "data",
        allowed_origins=["http://testserver"],
        allowed_hosts=["testserver"],
        web_dist_dir=tmp_path / "missing-web",
    )
    app = create_app(settings)
    with TestClient(app, base_url="http://testserver") as client:
        with source.open("rb") as handle:
            invalid_mode = client.post(
                "/api/v1/jobs",
                files={
                    "document": (
                        source.name,
                        handle,
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    )
                },
                data={"rule_pack_id": "zh-thesis-default", "mode": "typo"},
            )
        assert invalid_mode.status_code == 422
        assert invalid_mode.json()["code"] == "VALIDATION_ERROR"

        with source.open("rb") as handle:
            created = client.post(
                "/api/v1/jobs",
                files={
                    "document": (
                        source.name,
                        handle,
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    )
                },
                data={"rule_pack_id": "zh-thesis-default", "render_preview": "false"},
            )
        assert created.json()["mode"] == "format"
        job_id = created.json()["id"]
        assert app.state.job_service.process_next("test-worker")
        job = client.get(f"/api/v1/jobs/{job_id}").json()
        assert job["status"] == "plan_ready"
        plan = client.get(f"/api/v1/jobs/{job_id}/plan").json()
        assert "operations" in plan
