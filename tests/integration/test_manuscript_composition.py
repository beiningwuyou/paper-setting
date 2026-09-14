import base64
import zipfile
from pathlib import Path

from docx import Document
from paper_setting_core.templates import (
    TemplateSectionStructureConfig,
    apply_manuscript_composition,
    build_manuscript_composition_preview,
)

PNG_1X1 = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
)


def test_manuscript_is_mapped_and_injected_into_template(tmp_path: Path) -> None:
    image = tmp_path / "figure.png"
    image.write_bytes(PNG_1X1)

    manuscript = tmp_path / "manuscript.docx"
    source = Document()
    source.add_paragraph("智能论文排版研究", style="Title")
    source.add_paragraph("摘要")
    source.add_paragraph("本文提出一种可验证的论文排版方法。")
    source.add_paragraph("关键词：论文；排版；完整性")
    source.add_paragraph("学生姓名：张三")
    source.add_paragraph("第一章 绪论", style="Heading 1")
    source.add_paragraph("这是需要完整注入模板的正文内容。")
    source.add_picture(str(image))
    table = source.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "项目"
    table.cell(0, 1).text = "内容"
    source.save(manuscript)

    template = tmp_path / "university-template.docx"
    target = Document()
    target.add_paragraph("{{论文题目}}")
    target.add_paragraph("摘要：{{摘要}}")
    target.add_paragraph("关键词：{{关键词}}")
    target.add_paragraph("作者：{{学生姓名}}")
    target.add_paragraph("{{正文}}")
    target.save(template)

    configuration = TemplateSectionStructureConfig()
    preview = build_manuscript_composition_preview(
        manuscript,
        template,
        {},
        configuration,
    )
    assert preview.can_generate is True
    assert preview.body_injection.source_blocks == 4
    assert preview.body_injection.image_relationships == 1
    assert {item.key for item in preview.mappings if item.source != "unmapped"} == {
        "论文题目",
        "摘要",
        "关键词",
        "学生姓名",
    }

    output = tmp_path / "composed.docx"
    result = apply_manuscript_composition(
        manuscript,
        template,
        output,
        {},
        configuration,
        expected_source_sha256=preview.source_sha256,
        expected_template_sha256=preview.template_sha256,
        expected_plan_version=preview.plan_version,
    )
    assert all(result.integrity_checks.values())
    reopened = Document(output)
    text = "\n".join(paragraph.text for paragraph in reopened.paragraphs)
    assert "智能论文排版研究" in text
    assert "本文提出一种可验证的论文排版方法" in text
    assert "作者：张三" in text
    assert "第一章 绪论" in text
    assert "这是需要完整注入模板的正文内容" in text
    assert reopened.tables[0].cell(0, 1).text == "内容"
    with zipfile.ZipFile(output) as archive:
        assert any(name.startswith("word/media/injected-") for name in archive.namelist())


def test_composition_requires_a_single_body_anchor(tmp_path: Path) -> None:
    manuscript = tmp_path / "source.docx"
    source = Document()
    source.add_paragraph("正文")
    source.save(manuscript)

    template = tmp_path / "template.docx"
    target = Document()
    target.add_paragraph("模板没有正文占位符")
    target.save(template)

    preview = build_manuscript_composition_preview(
        manuscript,
        template,
        {},
        TemplateSectionStructureConfig(),
    )
    assert preview.can_generate is False
    assert any("{{document.body}}" in blocker for blocker in preview.blockers)
