import json
import uuid
from pathlib import Path
from urllib.parse import unquote

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from fastapi.testclient import TestClient
from paper_setting_api.main import create_app
from paper_setting_core.documents import conversion
from paper_setting_runtime.config import Settings

GUIDE_TEXT = """
页面采用 A4 纸，纵向，上下边距 2.5cm，左边距 3cm，右边距 2.5cm。
正文：宋体，小四，Times New Roman，1.5 倍行距，首行缩进 2 字符，两端对齐。
一级标题：黑体，三号，加粗，居中，段前 12 磅，段后 6 磅。
参考文献条目：宋体，五号，固定值 20 磅。
注释一律采用脚注，文末不列参考文献。每页脚注重新编号，编号格式为带圈数字。
"""


def test_legacy_doc_upload_is_normalized_and_original_is_preserved(
    sample_docx: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    settings = Settings(
        env="test",
        database_url=f"sqlite:///{tmp_path / 'legacy.db'}",
        data_dir=tmp_path / "data",
        allowed_origins=["http://testserver"],
        allowed_hosts=["testserver"],
        web_dist_dir=tmp_path / "missing-web",
    )

    def fake_convert(_source: Path, target: Path, *, timeout_seconds: int) -> None:
        assert timeout_seconds == settings.doc_conversion_timeout_seconds
        target.write_bytes(sample_docx.read_bytes())

    monkeypatch.setattr(conversion, "convert_legacy_doc", fake_convert)
    app = create_app(settings)
    legacy_payload = conversion.LEGACY_DOC_SIGNATURE + b"legacy word payload"
    with TestClient(app, base_url="http://testserver") as client:
        created = client.post(
            "/api/v1/jobs",
            files={"document": ("旧版论文.doc", legacy_payload, "application/msword")},
            data={"rule_pack_id": "zh-thesis-default", "render_preview": "false"},
        )

    assert created.status_code == 202, created.text
    job = created.json()
    assert job["source_filename"] == "旧版论文.doc"
    job_input = settings.jobs_dir / job["id"] / "input"
    assert (job_input / "source.docx").read_bytes() == sample_docx.read_bytes()
    assert (job_input / "source.original.doc").read_bytes() == legacy_payload


def test_full_api_workflow(sample_docx: Path, tmp_path: Path) -> None:
    settings = Settings(
        env="test",
        database_url=f"sqlite:///{tmp_path / 'test.db'}",
        data_dir=tmp_path / "data",
        allowed_origins=["http://testserver"],
        allowed_hosts=["testserver"],
        web_dist_dir=tmp_path / "missing-web",
    )
    app = create_app(settings)
    with TestClient(app, base_url="http://testserver") as client:
        response = client.get("/health")
        assert response.status_code == 200
        assert response.headers["X-Request-Id"]
        assert client.get("/ready").headers["X-Paper-Setting-Instance"] == "development"
        assert client.get("/health", headers={"Origin": "http://testserver"}).status_code == 200
        remote_response = client.get(
            "/health", headers={"Origin": "https://remote.example"}
        )
        assert remote_response.status_code == 403
        assert remote_response.headers["content-type"].startswith("application/problem+json")
        assert remote_response.headers["X-Request-Id"]
        assert remote_response.headers["Permissions-Policy"]
        assert remote_response.json()["code"] == "ORIGIN_NOT_ALLOWED"

        with sample_docx.open("rb") as handle:
            created = client.post(
                "/api/v1/jobs",
                files={
                    "document": (
                        sample_docx.name,
                        handle,
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    )
                },
                data={"rule_pack_id": "zh-thesis-default", "render_preview": "false"},
            )
        assert created.status_code == 202, created.text
        job_id = created.json()["id"]
        pending_inspection = client.get(f"/api/v1/jobs/{job_id}/inspection")
        assert pending_inspection.status_code == 409
        assert pending_inspection.json()["code"] == "PLAN_NOT_READY"

        assert app.state.job_service.process_next("test-worker")
        job = client.get(f"/api/v1/jobs/{job_id}").json()
        assert job["status"] == "plan_ready"
        inspection = client.get(f"/api/v1/jobs/{job_id}/inspection")
        assert inspection.status_code == 200
        assert inspection.json()["summary"]["paragraphs"] > 0
        plan = client.get(f"/api/v1/jobs/{job_id}/plan").json()
        assert "compliant" in plan["summary"]
        assert all("changed_fields" in item for item in plan["operations"])
        approved = [
            item["operation_id"] for item in plan["operations"] if item["status"] == "proposed"
        ]
        rejected = [
            item["operation_id"] for item in plan["operations"] if item["status"] != "proposed"
        ]
        response = client.post(
            f"/api/v1/jobs/{job_id}/approve",
            json={
                "plan_version": plan["plan_version"],
                "approved_operation_ids": approved,
                "rejected_operation_ids": rejected,
            },
        )
        assert response.status_code == 200, response.text
        assert app.state.job_service.process_next("test-worker")

        completed = client.get(f"/api/v1/jobs/{job_id}").json()
        assert completed["status"] == "completed", completed
        report = client.get(f"/api/v1/jobs/{job_id}/report").json()
        assert report["integrity_ok"] is True
        assert report["total_rule_targets"] >= report["already_compliant"]
        assert report["compliant_targets"] == report["total_rule_targets"]
        assert report["compliance_rate"] == 1.0
        artifacts = client.get(f"/api/v1/jobs/{job_id}/artifacts").json()
        assert {item["kind"] for item in artifacts} == {"formatted_docx"}
        download = client.get(f"/api/v1/jobs/{job_id}/artifacts/formatted_docx")
        assert download.status_code == 200
        assert "论文样稿排版完成.docx" in unquote(
            download.headers["content-disposition"]
        )
        assert download.content[:2] == b"PK"
        output_dir = tmp_path / "data" / "jobs" / job_id / "output"
        assert (output_dir / "report.json").exists()
        assert (output_dir / "operations.json").exists()
        assert not (output_dir / "report.html").exists()
        assert not (output_dir / "preview.pdf").exists()
        event_stream = client.get(f"/api/v1/jobs/{job_id}/events")
        assert event_stream.status_code == 200
        assert "event: job.completed" in event_stream.text


def test_request_id_and_rate_limit_are_consistent(tmp_path: Path) -> None:
    settings = Settings(
        env="test",
        database_url=f"sqlite:///{tmp_path / 'middleware.db'}",
        data_dir=tmp_path / "data",
        allowed_origins=["http://testserver"],
        allowed_hosts=["testserver"],
        api_rate_limit_per_minute=2,
        web_dist_dir=tmp_path / "missing-web",
    )
    app = create_app(settings)
    with TestClient(app, base_url="http://testserver") as client:
        supplied = "local-trace_123"
        accepted = client.get("/health", headers={"X-Request-Id": supplied})
        assert accepted.headers["X-Request-Id"] == supplied

        regenerated = client.get("/health", headers={"X-Request-Id": "invalid id with spaces"})
        uuid.UUID(regenerated.headers["X-Request-Id"])

        limited = client.get("/health")
        assert limited.status_code == 429
        assert limited.headers["content-type"].startswith("application/problem+json")
        assert limited.headers["X-RateLimit-Remaining"] == "0"
        assert limited.headers["X-Content-Type-Options"] == "nosniff"
        assert limited.json()["code"] == "RATE_LIMITED"
        assert limited.json()["request_id"] == limited.headers["X-Request-Id"]


def test_uploaded_job_can_be_cancelled(sample_docx: Path, tmp_path: Path) -> None:
    settings = Settings(
        env="test",
        database_url=f"sqlite:///{tmp_path / 'cancel.db'}",
        data_dir=tmp_path / "data",
        allowed_origins=["http://testserver"],
        allowed_hosts=["testserver"],
        web_dist_dir=tmp_path / "missing-web",
    )
    app = create_app(settings)
    with TestClient(app, base_url="http://testserver") as client:
        with sample_docx.open("rb") as handle:
            created = client.post(
                "/api/v1/jobs",
                files={
                    "document": (
                        sample_docx.name,
                        handle,
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    )
                },
                data={"rule_pack_id": "zh-thesis-default", "render_preview": "false"},
            )
        job_id = created.json()["id"]
        cancelled = client.post(f"/api/v1/jobs/{job_id}/cancel")
        assert cancelled.status_code == 200
        assert cancelled.json()["status"] == "cancelled"
        assert app.state.job_service.process_next("test-worker") is False


def test_rule_pack_draft_can_be_reviewed_and_imported(tmp_path: Path) -> None:
    settings = Settings(
        env="test",
        database_url=f"sqlite:///{tmp_path / 'rules.db'}",
        data_dir=tmp_path / "data",
        allowed_origins=["http://testserver"],
        allowed_hosts=["testserver"],
        web_dist_dir=tmp_path / "missing-web",
    )
    app = create_app(settings)
    with TestClient(app, base_url="http://testserver") as client:
        response = client.post(
            "/api/v1/rule-packs/drafts",
            data={
                "text": GUIDE_TEXT,
                "name": "从文本生成的规则",
                "rule_pack_id": "generated-thesis-rules",
            },
        )
        assert response.status_code == 200, response.text
        draft = response.json()
        assert draft["rule_pack"]["id"] == "generated-thesis-rules"
        assert draft["recognized_properties"] >= 20
        assert draft["evidence"]
        assert draft["rule_pack"]["advanced"]["notes"] == {
            "enabled": True,
            "convert_inline_citations": True,
            "convert_endnotes": True,
            "delete_bibliography": True,
            "numbering_restart": "each_page",
            "number_format": "decimal_enclosed_circle",
        }

        imported = client.post(
            "/api/v1/rule-packs/import",
            files={
                "document": (
                    "generated.json",
                    json.dumps(draft["rule_pack"], ensure_ascii=False).encode(),
                    "application/json",
                )
            },
        )
        assert imported.status_code == 201, imported.text
        assert imported.json()["id"] == "generated-thesis-rules"
        assert client.get("/api/v1/rule-packs/generated-thesis-rules").status_code == 200

        missing = client.post("/api/v1/rule-packs/drafts")
        assert missing.status_code == 422
        assert missing.json()["code"] == "RULE_SOURCE_INVALID"


def test_manual_review_requires_confirmation_and_preserves_fields(tmp_path: Path) -> None:
    source = tmp_path / "field-review.docx"
    document = Document()
    paragraph = document.add_paragraph("这是一段包含页码字段的正文内容，当前页码是 ")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    paragraph.add_run()._r.append(begin)
    paragraph.add_run()._r.append(instruction)
    paragraph.add_run()._r.append(end)
    document.save(source)

    settings = Settings(
        env="test",
        database_url=f"sqlite:///{tmp_path / 'manual-review.db'}",
        data_dir=tmp_path / "data",
        allowed_origins=["http://testserver"],
        allowed_hosts=["testserver"],
        web_dist_dir=tmp_path / "missing-web",
    )
    app = create_app(settings)
    with TestClient(app, base_url="http://testserver") as client:
        with source.open("rb") as handle:
            created = client.post(
                "/api/v1/jobs",
                files={
                    "document": (
                        source.name,
                        handle,
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    )
                },
                data={"rule_pack_id": "zh-thesis-default", "render_preview": "false"},
            )
        job_id = created.json()["id"]
        assert app.state.job_service.process_next("test-worker")
        plan = client.get(f"/api/v1/jobs/{job_id}/plan").json()
        approved = [item["operation_id"] for item in plan["operations"]]
        manual = [
            item["operation_id"]
            for item in plan["operations"]
            if item["status"] == "manual_review"
        ]
        assert manual
        assert any(
            item["execution_scope"] == "preserve_protected_content"
            for item in plan["operations"]
            if item["operation_id"] in manual
        )

        unconfirmed = client.post(
            f"/api/v1/jobs/{job_id}/approve",
            json={
                "plan_version": plan["plan_version"],
                "approved_operation_ids": approved,
                "rejected_operation_ids": [],
            },
        )
        assert unconfirmed.status_code == 409
        assert unconfirmed.json()["code"] == "APPROVAL_REQUIRED"

        confirmed = client.post(
            f"/api/v1/jobs/{job_id}/approve",
            json={
                "plan_version": plan["plan_version"],
                "approved_operation_ids": approved,
                "rejected_operation_ids": [],
                "confirmed_manual_operation_ids": manual,
            },
        )
        assert confirmed.status_code == 200, confirmed.text
        assert app.state.job_service.process_next("test-worker")
        completed = client.get(f"/api/v1/jobs/{job_id}").json()
        assert completed["status"] == "completed", completed
        report = client.get(f"/api/v1/jobs/{job_id}/report").json()
        assert report["integrity_ok"] is True
        assert report["checks"]["field_instructions_unchanged"] is True
        assert report["operation_counts"]["applied"] == len(approved)
