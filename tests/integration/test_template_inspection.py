import json
import zipfile
from pathlib import Path

import pytest
from docx import Document
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from fastapi.testclient import TestClient
from paper_setting_api.main import create_app
from paper_setting_core.errors import TemplateFillBlockedError
from paper_setting_core.templates import (
    TemplateSectionStructureConfig,
    apply_section_structure,
    apply_template_combined,
    build_section_structure_preview,
    build_template_combined_preview,
    build_template_fill_preview,
    fill_template,
    inspect_template,
    parse_template_values,
)
from paper_setting_core.templates.models import TemplateSectionUpdate
from paper_setting_runtime.config import Settings


def _set_page_numbering(section: object, number_format: str, start: int) -> None:
    properties = section._sectPr  # type: ignore[attr-defined]
    for existing in properties.xpath("./w:pgNumType"):
        properties.remove(existing)
    node = OxmlElement("w:pgNumType")
    node.set(qn("w:fmt"), number_format)
    node.set(qn("w:start"), str(start))
    properties.append(node)


def _append_content_control(
    document: Document,
    *,
    tag_value: str = "metadata.author",
    alias_value: str = "作者姓名",
    data_bound: bool = False,
) -> None:
    control = OxmlElement("w:sdt")
    properties = OxmlElement("w:sdtPr")
    tag = OxmlElement("w:tag")
    tag.set(qn("w:val"), tag_value)
    alias = OxmlElement("w:alias")
    alias.set(qn("w:val"), alias_value)
    properties.extend([tag, alias])
    if data_bound:
        binding = OxmlElement("w:dataBinding")
        binding.set(qn("w:storeItemID"), "{00000000-0000-0000-0000-000000000000}")
        binding.set(qn("w:xpath"), "/root/value")
        properties.append(binding)
    content = OxmlElement("w:sdtContent")
    paragraph = OxmlElement("w:p")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "作者姓名"
    run.append(text)
    paragraph.append(run)
    content.append(paragraph)
    control.extend([properties, content])
    document.element.body.insert(-1, control)


def _append_ref_field(paragraph: object) -> None:
    run = OxmlElement("w:r")
    instruction = OxmlElement("w:instrText")
    instruction.text = " REF figure_one \\h "
    run.append(instruction)
    paragraph._p.append(run)  # type: ignore[attr-defined]


def _template(path: Path) -> None:
    document = Document()
    document.add_paragraph("学校论文模板", style="Title")
    placeholder = document.add_paragraph()
    placeholder.add_run("{{student.")
    placeholder.add_run("name}}")
    _append_content_control(document)
    _append_ref_field(document.add_paragraph("图示引用："))
    first = document.sections[0]
    first.different_first_page_header_footer = True
    header = first.header.paragraphs[0]
    header.add_run("[[school.")
    header.add_run("name]]")
    first.footer.paragraphs[0].text = "页码"
    _set_page_numbering(first, "lowerRoman", 1)

    second = document.add_section(WD_SECTION.NEW_PAGE)
    _set_page_numbering(second, "decimal", 1)
    document.add_paragraph("第一章 绪论", style="Heading 1")
    document.save(path)


def test_template_inspection_finds_structure_and_capabilities(tmp_path: Path) -> None:
    source = tmp_path / "学校模板.docx"
    _template(source)

    inspection = inspect_template(source)

    assert inspection.summary.sections == 2
    assert inspection.summary.placeholders == 2
    assert inspection.summary.content_controls == 1
    assert inspection.summary.fields == 1
    assert {item.key for item in inspection.placeholders} == {"student.name", "school.name"}
    assert inspection.content_controls[0].tag == "metadata.author"
    assert [item.page_number_format for item in inspection.sections] == [
        "lowerRoman",
        "decimal",
    ]
    assert inspection.sections[0].different_first_page is True
    assert inspection.sections[1].inherits_headers is True
    capability_status = {item.id: item.status for item in inspection.capabilities}
    assert capability_status["template.inspect"] == "supported"
    assert capability_status["template.fill"] == "supported"
    assert capability_status["template.placeholders"] == "supported"
    assert capability_status["sections.page_numbering"] == "supported"
    assert capability_status["sections.headers_footers"] == "supported"
    assert capability_status["word.cross_references"] == "supported"
    assert capability_status["word.content_controls"] == "supported"
    assert any(item.role == "paper_title" for item in inspection.semantic_candidates)


def test_template_inspection_api(tmp_path: Path) -> None:
    source = tmp_path / "学校模板.docx"
    _template(source)
    settings = Settings(
        env="test",
        database_url=f"sqlite:///{tmp_path / 'templates.db'}",
        data_dir=tmp_path / "data",
        allowed_hosts=["testserver"],
        web_dist_dir=tmp_path / "missing-web",
    )
    app = create_app(settings)
    with TestClient(app, base_url="http://testserver") as client, source.open("rb") as handle:
        response = client.post(
            "/api/v1/templates/inspect",
            files={
                "document": (
                    source.name,
                    handle,
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )
        assert response.status_code == 200, response.text
        assert response.json()["summary"]["sections"] == 2
        assert response.json()["source_filename"] == source.name


def _part_text(path: Path, part_name: str) -> str:
    with zipfile.ZipFile(path) as archive:
        root = parse_xml(archive.read(part_name))
        return "".join(root.xpath(".//w:t/text()"))


def _field_instructions(path: Path) -> list[str]:
    with zipfile.ZipFile(path) as archive:
        root = parse_xml(archive.read("word/document.xml"))
        return [str(value) for value in root.xpath(".//w:instrText/text()")]


def test_template_fill_replaces_cross_run_header_and_content_control(tmp_path: Path) -> None:
    source = tmp_path / "学校模板.docx"
    output = tmp_path / "已填写.docx"
    _template(source)
    source_bytes = source.read_bytes()
    values = parse_template_values(
        json.dumps(
            {
                "student.name": "张三",
                "school.name": "测试大学",
                "metadata.author": "李四",
            },
            ensure_ascii=False,
        )
    )
    preview = build_template_fill_preview(source, values, source_filename=source.name)

    assert preview.can_generate is True
    assert preview.replacement_count == 3
    assert not preview.blocked_targets

    result = fill_template(
        source,
        output,
        values,
        source_filename=source.name,
        expected_source_sha256=preview.source_sha256,
        expected_plan_version=preview.plan_version,
    )

    assert source.read_bytes() == source_bytes
    assert all(result.integrity_checks.values())
    assert "张三" in _part_text(output, "word/document.xml")
    assert "李四" in _part_text(output, "word/document.xml")
    assert "{{student.name}}" not in _part_text(output, "word/document.xml")
    header_part = next(name for name in result.changed_parts if name.startswith("word/header"))
    assert "测试大学" in _part_text(output, header_part)
    assert _field_instructions(output) == _field_instructions(source)
    Document(output)


def test_template_fill_blocks_data_bound_content_control(tmp_path: Path) -> None:
    source = tmp_path / "bound.docx"
    output = tmp_path / "bound-filled.docx"
    document = Document()
    _append_content_control(document, tag_value="bound.value", data_bound=True)
    document.save(source)
    values = {"bound.value": "不应写入"}
    preview = build_template_fill_preview(source, values)

    assert preview.can_generate is False
    assert preview.blocked_targets[0].key == "bound.value"
    with pytest.raises(TemplateFillBlockedError):
        fill_template(
            source,
            output,
            values,
            expected_source_sha256=preview.source_sha256,
            expected_plan_version=preview.plan_version,
        )
    assert not output.exists()


def test_template_fill_api_preview_generate_and_stale_plan(tmp_path: Path) -> None:
    source = tmp_path / "学校模板.docx"
    _template(source)
    settings = Settings(
        env="test",
        database_url=f"sqlite:///{tmp_path / 'fill.db'}",
        data_dir=tmp_path / "data-fill",
        allowed_hosts=["testserver"],
        web_dist_dir=tmp_path / "missing-web",
    )
    app = create_app(settings)
    values = json.dumps({"student.name": "王五", "metadata.author": "赵六"}, ensure_ascii=False)
    with TestClient(app, base_url="http://testserver") as client:
        with source.open("rb") as handle:
            preview_response = client.post(
                "/api/v1/templates/fill/preview",
                data={"values": values},
                files={
                    "document": (
                        source.name,
                        handle,
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    )
                },
            )
        assert preview_response.status_code == 200, preview_response.text
        preview = preview_response.json()
        with source.open("rb") as handle:
            fill_response = client.post(
                "/api/v1/templates/fill",
                data={
                    "values": values,
                    "source_sha256": preview["source_sha256"],
                    "plan_version": preview["plan_version"],
                },
                files={
                    "document": (
                        source.name,
                        handle,
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    )
                },
            )
        assert fill_response.status_code == 200, fill_response.text
        assert fill_response.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )
        generated = tmp_path / "api-filled.docx"
        generated.write_bytes(fill_response.content)
        assert "王五" in _part_text(generated, "word/document.xml")

        with source.open("rb") as handle:
            stale_response = client.post(
                "/api/v1/templates/fill",
                data={
                    "values": values,
                    "source_sha256": preview["source_sha256"],
                    "plan_version": "stale-plan",
                },
                files={
                    "document": (
                        source.name,
                        handle,
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    )
                },
            )
        assert stale_response.status_code == 409
        assert stale_response.json()["code"] == "STALE_TEMPLATE_FILL_PLAN"


def test_section_structure_updates_page_numbers_and_clones_inherited_parts(
    tmp_path: Path,
) -> None:
    source = tmp_path / "sections.docx"
    output = tmp_path / "sections-updated.docx"
    _template(source)
    source_bytes = source.read_bytes()
    configuration = TemplateSectionStructureConfig(
        sections=[
            TemplateSectionUpdate(
                section_index=0,
                page_number_format="upperRoman",
                page_number_start=3,
                different_first_page=False,
            ),
            TemplateSectionUpdate(
                section_index=1,
                page_number_format="decimal",
                page_number_start=5,
                different_first_page=True,
                header_mode="independent_copy",
                footer_mode="independent_copy",
            ),
        ],
        even_and_odd_headers=True,
    )
    preview = build_section_structure_preview(source, configuration)

    assert preview.can_generate is True
    assert {item.operation_type for item in preview.operations} == {
        "page_numbering",
        "different_first_page",
        "even_and_odd_headers",
        "header_link",
        "footer_link",
    }
    result = apply_section_structure(
        source,
        output,
        configuration,
        expected_source_sha256=preview.source_sha256,
        expected_plan_version=preview.plan_version,
    )

    assert source.read_bytes() == source_bytes
    assert all(result.integrity_checks.values())
    assert any(name.startswith("word/header") for name in result.added_parts)
    assert any(name.startswith("word/footer") for name in result.added_parts)
    inspection = inspect_template(output)
    assert inspection.sections[0].page_number_format == "upperRoman"
    assert inspection.sections[0].page_number_start == 3
    assert inspection.sections[0].different_first_page is False
    assert inspection.sections[1].page_number_start == 5
    assert inspection.sections[1].different_first_page is True
    assert inspection.sections[1].inherits_headers is False
    assert inspection.sections[1].inherits_footers is False
    assert inspection.even_and_odd_headers is True
    cloned_header = next(
        item.part_name
        for item in inspection.sections[1].references
        if item.story == "header" and item.part_name
    )
    assert "[[school.name]]" in _part_text(output, cloned_header)


def test_section_structure_can_restore_inheritance(tmp_path: Path) -> None:
    source = tmp_path / "independent.docx"
    output = tmp_path / "inherited.docx"
    document = Document()
    document.sections[0].header.paragraphs[0].text = "前置页眉"
    second = document.add_section(WD_SECTION.NEW_PAGE)
    second.header.is_linked_to_previous = False
    second.footer.is_linked_to_previous = False
    second.header.paragraphs[0].text = "独立页眉"
    second.footer.paragraphs[0].text = "独立页脚"
    document.save(source)
    configuration = TemplateSectionStructureConfig(
        sections=[
            TemplateSectionUpdate(
                section_index=1,
                header_mode="inherit",
                footer_mode="inherit",
                clear_page_numbering=True,
            )
        ]
    )
    preview = build_section_structure_preview(source, configuration)
    result = apply_section_structure(
        source,
        output,
        configuration,
        expected_source_sha256=preview.source_sha256,
        expected_plan_version=preview.plan_version,
    )

    assert all(result.integrity_checks.values())
    inspection = inspect_template(output)
    assert inspection.sections[1].inherits_headers is True
    assert inspection.sections[1].inherits_footers is True
    assert inspection.sections[1].references == []


def test_section_structure_api_preview_and_generate(tmp_path: Path) -> None:
    source = tmp_path / "api-sections.docx"
    _template(source)
    settings = Settings(
        env="test",
        database_url=f"sqlite:///{tmp_path / 'structure.db'}",
        data_dir=tmp_path / "data-structure",
        allowed_hosts=["testserver"],
        web_dist_dir=tmp_path / "missing-web",
    )
    app = create_app(settings)
    configuration = json.dumps(
        {
            "sections": [
                {
                    "section_index": 1,
                    "page_number_format": "decimal",
                    "page_number_start": 7,
                    "header_mode": "independent_copy",
                }
            ],
            "even_and_odd_headers": True,
        }
    )
    with TestClient(app, base_url="http://testserver") as client:
        with source.open("rb") as handle:
            preview_response = client.post(
                "/api/v1/templates/structure/preview",
                data={"configuration": configuration},
                files={
                    "document": (
                        source.name,
                        handle,
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    )
                },
            )
        assert preview_response.status_code == 200, preview_response.text
        preview = preview_response.json()
        with source.open("rb") as handle:
            response = client.post(
                "/api/v1/templates/structure",
                data={
                    "configuration": configuration,
                    "source_sha256": preview["source_sha256"],
                    "plan_version": preview["plan_version"],
                },
                files={
                    "document": (
                        source.name,
                        handle,
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    )
                },
            )
        assert response.status_code == 200, response.text
        generated = tmp_path / "structured-api.docx"
        generated.write_bytes(response.content)
        inspection = inspect_template(generated)
        assert inspection.sections[1].page_number_start == 7
        assert inspection.sections[1].inherits_headers is False
        assert inspection.even_and_odd_headers is True


def test_template_combined_fills_fields_and_restructures_in_one_copy(tmp_path: Path) -> None:
    source = tmp_path / "学校模板.docx"
    output = tmp_path / "组合副本.docx"
    _template(source)
    source_bytes = source.read_bytes()
    values = parse_template_values(
        json.dumps(
            {
                "student.name": "张三",
                "school.name": "测试大学",
                "metadata.author": "李四",
            },
            ensure_ascii=False,
        )
    )
    configuration = TemplateSectionStructureConfig(
        sections=[
            TemplateSectionUpdate(
                section_index=0,
                page_number_format="upperRoman",
                page_number_start=1,
                different_first_page=False,
            ),
            TemplateSectionUpdate(
                section_index=1,
                page_number_format="decimal",
                page_number_start=1,
                header_mode="independent_copy",
                footer_mode="independent_copy",
            ),
        ],
        even_and_odd_headers=True,
    )
    preview = build_template_combined_preview(
        source, values, configuration, source_filename=source.name
    )
    assert preview.can_generate is True
    assert preview.fill.replacement_count == 3
    assert preview.structure.operations
    assert not preview.blockers

    result = apply_template_combined(
        source,
        output,
        values,
        configuration,
        source_filename=source.name,
        expected_source_sha256=preview.source_sha256,
        expected_plan_version=preview.plan_version,
    )

    assert source.read_bytes() == source_bytes
    assert all(result.integrity_checks.values())
    inspection = inspect_template(output)
    assert "张三" in _part_text(output, "word/document.xml")
    assert "李四" in _part_text(output, "word/document.xml")
    assert "{{student.name}}" not in _part_text(output, "word/document.xml")
    assert inspection.even_and_odd_headers is True
    assert inspection.sections[0].page_number_format == "upperRoman"
    assert inspection.sections[1].page_number_start == 1
    assert inspection.sections[1].inherits_headers is False
    assert inspection.sections[1].inherits_footers is False
    cloned_header = next(
        item.part_name
        for item in inspection.sections[1].references
        if item.story == "header" and item.part_name
    )
    assert "测试大学" in _part_text(output, cloned_header)
    assert _field_instructions(output) == _field_instructions(source)
    Document(output)


def test_template_combined_preview_blocks_protected_targets(tmp_path: Path) -> None:
    source = tmp_path / "bound.docx"
    document = Document()
    _append_content_control(document, tag_value="bound.value", data_bound=True)
    document.save(source)
    preview = build_template_combined_preview(
        source,
        {"bound.value": "不应写入"},
        TemplateSectionStructureConfig(),
    )

    assert preview.can_generate is False
    assert preview.blockers
    assert any("bound.value" in blocker for blocker in preview.blockers)


def test_template_combined_api_preview_and_generate(tmp_path: Path) -> None:
    source = tmp_path / "api-combined.docx"
    _template(source)
    settings = Settings(
        env="test",
        database_url=f"sqlite:///{tmp_path / 'combined.db'}",
        data_dir=tmp_path / "data-combined",
        allowed_hosts=["testserver"],
        web_dist_dir=tmp_path / "missing-web",
    )
    app = create_app(settings)
    values = json.dumps({"student.name": "王五", "school.name": "某大学"}, ensure_ascii=False)
    configuration = json.dumps(
        {
            "sections": [
                {
                    "section_index": 1,
                    "page_number_format": "decimal",
                    "page_number_start": 9,
                    "header_mode": "independent_copy",
                }
            ],
            "even_and_odd_headers": True,
        }
    )
    with TestClient(app, base_url="http://testserver") as client:
        with source.open("rb") as handle:
            preview_response = client.post(
                "/api/v1/templates/combined/preview",
                data={"values": values, "configuration": configuration},
                files={
                    "document": (
                        source.name,
                        handle,
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    )
                },
            )
        assert preview_response.status_code == 200, preview_response.text
        preview = preview_response.json()
        assert preview["fill"]["replacement_count"] == 2
        assert preview["structure"]["operations"]
        with source.open("rb") as handle:
            response = client.post(
                "/api/v1/templates/combined",
                data={
                    "values": values,
                    "configuration": configuration,
                    "source_sha256": preview["source_sha256"],
                    "plan_version": preview["plan_version"],
                },
                files={
                    "document": (
                        source.name,
                        handle,
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    )
                },
            )
        assert response.status_code == 200, response.text
        generated = tmp_path / "combined-api.docx"
        generated.write_bytes(response.content)
        inspection = inspect_template(generated)
        assert "王五" in _part_text(generated, "word/document.xml")
        assert inspection.sections[1].page_number_start == 9
        assert inspection.sections[1].inherits_headers is False
        assert inspection.even_and_odd_headers is True
        assert response.headers["X-Template-Replacements"] == "2"

        with source.open("rb") as handle:
            stale_response = client.post(
                "/api/v1/templates/combined",
                data={
                    "values": values,
                    "configuration": configuration,
                    "source_sha256": preview["source_sha256"],
                    "plan_version": "stale-plan",
                },
                files={
                    "document": (
                        source.name,
                        handle,
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    )
                },
            )
        assert stale_response.status_code == 409
        assert stale_response.json()["code"] == "STALE_TEMPLATE_COMBINED_PLAN"
