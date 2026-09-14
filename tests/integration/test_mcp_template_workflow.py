from pathlib import Path

from docx import Document
from paper_setting_api.main import create_app
from paper_setting_runtime.config import Settings


def test_mcp_agent_can_create_template_rule_pack_and_job(
    tmp_path: Path,
    monkeypatch,
) -> None:
    from paper_setting_mcp import main as mcp_main

    source = tmp_path / "agent-template.docx"
    document = Document()
    document.add_paragraph("{{student.name}} 的论文")
    document.save(source)

    settings = Settings(
        env="test",
        database_url=f"sqlite:///{tmp_path / 'mcp-template.db'}",
        data_dir=tmp_path / "data",
        allowed_origins=["http://testserver"],
        allowed_hosts=["testserver"],
        allowed_input_roots=[tmp_path],
        web_dist_dir=tmp_path / "missing-web",
    )
    app = create_app(settings)
    monkeypatch.setattr(mcp_main, "settings", settings)
    monkeypatch.setattr(mcp_main, "rule_pack_service", app.state.rule_pack_service)
    monkeypatch.setattr(mcp_main, "job_service", app.state.job_service)

    inspection = mcp_main.inspect_template(str(source))
    assert inspection["source_filename"] == source.name
    assert inspection["placeholders"][0]["key"] == "student.name"

    base = mcp_main.get_rule_pack("zh-thesis-default")
    assert base["id"] == "zh-thesis-default"

    created_pack = mcp_main.create_template_rule_pack(
        rule_pack_id="agent-template-pack",
        name="Agent 模板规则",
        values={"student.name": "张三"},
        structure={"sections": [{"section_index": 0, "page_number_start": 1}]},
    )
    assert created_pack["id"] == "agent-template-pack"
    assert created_pack["capability_report"]["executable"] is True

    created_job = mcp_main.create_job(
        str(source),
        rule_pack_id="agent-template-pack",
        render_preview=False,
    )
    assert created_job["mode"] == "template"
    assert app.state.job_service.process_next("mcp-test-worker")

    current = mcp_main.get_job(created_job["id"])
    assert current["status"] == "plan_ready"
    plan = mcp_main.get_template_plan(created_job["id"])
    assert plan["can_generate"] is True
    assert plan["fill"]["replacement_count"] == 1
    assert plan["structure"]["operations"]

    app.state.engine.dispose()
