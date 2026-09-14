from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Mm, Pt, RGBColor
from docx.oxml.ns import qn
from paper_setting_core.documents import inspect_document
from paper_setting_core.formatting import apply_plan
from paper_setting_core.planning import generate_plan
from paper_setting_core.planning.service import _changed_fields
from paper_setting_core.rulepacks.models import RulePack
from paper_setting_core.rulepacks.policy import (
    FormattingPolicy,
    default_formatting_policy,
    resolve_formatting,
)
from paper_setting_core.rulepacks.service import canonical_rule_pack_json
from paper_setting_core.verification import validate_result


def sparse_rules(**updates) -> RulePack:
    return RulePack.model_validate(
        {
            "id": "sparse",
            "name": "仅规定正文字号",
            "version": "1",
            "roles": [
                {
                    "id": "body",
                    "role": "body",
                    "description": "正文",
                    "source": "test",
                    "character": {"size_pt": 13},
                }
            ],
            **updates,
        }
    )


def test_sparse_page_fields_survive_serialization():
    rules = sparse_rules(page={"margin_left_mm": 30})
    restored = RulePack.model_validate_json(canonical_rule_pack_json(rules))
    _, page, sources = resolve_formatting(restored, FormattingPolicy(mode="preserve"))
    assert page == {"margin_left_mm": 30}
    assert sources["page"] == {"margin_left_mm": "rule_pack"}


def test_baseline_merges_by_property_and_never_mutates_rules():
    rules = sparse_rules()
    before = rules.model_dump()
    effective, _, sources = resolve_formatting(rules, default_formatting_policy())
    body = effective.rule_for("body")
    assert body.character.size_pt == 13
    assert body.character.color == "000000"
    assert body.character.bold is False
    assert body.character.italic is False
    assert body.character.underline is False
    assert body.paragraph.space_after_pt == 0
    assert sources["body"]["character.size_pt"] == "rule_pack"
    assert sources["body"]["character.color"] == "default_template"
    assert effective.rule_for("heading_1") is not None
    assert effective.rule_for("table_body") is None
    assert rules.model_dump() == before


def test_false_and_zero_are_explicit_overrides():
    rules = sparse_rules(
        roles=[
            {
                "id": "heading",
                "role": "heading_1",
                "description": "标题",
                "source": "test",
                "paragraph": {"space_before_pt": 0},
                "character": {"bold": False},
            }
        ]
    )
    effective, _, _ = resolve_formatting(rules, default_formatting_policy())
    assert effective.rule_for("heading_1").character.bold is False
    assert effective.rule_for("heading_1").paragraph.space_before_pt == 0


def test_inherited_zero_paragraph_spacing_does_not_create_direct_formatting():
    assert _changed_fields(
        {"space_before_pt": None, "space_after_pt": None},
        {"space_before_pt": 0, "space_after_pt": 0},
        "paragraph",
    ) == []
    assert _changed_fields(
        {"space_before_pt": 6, "space_after_pt": 10},
        {"space_before_pt": 0, "space_after_pt": 0},
        "paragraph",
    ) == ["paragraph.space_before_pt", "paragraph.space_after_pt"]


def test_body_spacing_can_be_written_and_inspected_in_lines(tmp_path: Path):
    source, output = tmp_path / "source.docx", tmp_path / "out.docx"
    doc = Document()
    paragraph = doc.add_paragraph("这是一段需要按行设置段前段后的正文。")
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(10)
    doc.save(source)
    rules = sparse_rules(
        roles=[
            {
                "id": "body",
                "role": "body",
                "description": "正文",
                "source": "test",
                "paragraph": {"space_before_lines": 0, "space_after_lines": 0},
            }
        ]
    )
    plan = generate_plan(inspect_document(source), rules, FormattingPolicy(mode="preserve"))
    apply_plan(source, output, plan, {op.operation_id for op in plan.operations})
    result = Document(output).paragraphs[0]
    spacing = result._p.pPr.find(qn("w:spacing"))
    assert spacing.get(qn("w:beforeLines")) == "0"
    assert spacing.get(qn("w:afterLines")) == "0"
    assert spacing.get(qn("w:before")) is None
    assert spacing.get(qn("w:after")) is None
    inspected = inspect_document(output).items[0].effective_format.paragraph
    assert inspected["space_before_lines"] == 0
    assert inspected["space_after_lines"] == 0


def test_standardize_removes_accidental_direct_emphasis_and_is_idempotent(tmp_path: Path):
    path, output = tmp_path / "source.docx", tmp_path / "out.docx"
    doc = Document()
    paragraph = doc.add_paragraph("这是一段包含不同字号和颜色的正文，用于验证默认基线补全。")
    paragraph.runs[0].font.color.rgb = RGBColor.from_string("FF0000")
    paragraph.runs[0].font.size = Pt(21)
    emphasis = paragraph.add_run("重点文字")
    emphasis.bold = True
    emphasis.italic = True
    emphasis.underline = True
    emphasis.font.superscript = True
    paragraph.paragraph_format.space_after = Pt(33)
    doc.save(path)
    policy = default_formatting_policy()
    plan = generate_plan(inspect_document(path), sparse_rules(), policy)
    executed = apply_plan(path, output, plan, {op.operation_id for op in plan.operations})
    result = Document(output).paragraphs[0]
    assert result.runs[0].font.size.pt == 13
    assert str(result.runs[0].font.color.rgb) == "000000"
    assert result.paragraph_format.space_after.pt == 0
    assert result.runs[1].bold is False
    assert result.runs[1].italic is False
    assert result.runs[1].underline is False
    assert result.runs[1].font.superscript
    assert not generate_plan(inspect_document(output), sparse_rules(), policy).operations
    report = validate_result(path, output, executed)
    assert report.integrity_ok
    assert report.format_source_counts["default_template"] > 0
    assert report.format_source_counts["rule_pack"] > 0


def test_preserve_keeps_inherited_format_and_rejected_paragraph(tmp_path: Path):
    path, output = tmp_path / "source.docx", tmp_path / "out.docx"
    doc = Document()
    style = doc.styles.add_style("Original Body", WD_STYLE_TYPE.PARAGRAPH)
    style.font.color.rgb = RGBColor.from_string("112233")
    style.font.italic = True
    style.paragraph_format.space_after = Pt(27)
    for _ in range(2):
        doc.add_paragraph("这是一段正文，保守排版应当保留未规定的继承格式。", style)
    doc.sections[0].left_margin = Mm(38)
    doc.save(path)
    plan = generate_plan(inspect_document(path), sparse_rules(), FormattingPolicy(mode="preserve"))
    assert all(op.operation_type != "apply_page_format" for op in plan.operations)
    op = plan.operations[0]
    apply_plan(
        path, output, plan, {op.operation_id}, confirmed_manual_operation_ids={op.operation_id}
    )
    after = inspect_document(output)
    assert after.items[0].effective_format.character["color"] == "112233"
    assert after.items[0].effective_format.paragraph["space_after_pt"] == 27
    assert after.items[0].effective_format.character["size_pt"] == 13
    assert after.items[1].effective_format.character["size_pt"] != 13
    assert Document(output).paragraphs[0].style.base_style.font.italic
    assert round(Document(output).sections[0].left_margin.mm) == 38


def test_partial_page_does_not_replace_other_section_settings(tmp_path: Path):
    path, output = tmp_path / "source.docx", tmp_path / "out.docx"
    doc = Document()
    doc.sections[0].right_margin = Mm(41)
    doc.sections[0].page_width = Mm(220)
    doc.save(path)
    rules = sparse_rules(page={"margin_left_mm": 30})
    plan = generate_plan(inspect_document(path), rules, FormattingPolicy(mode="preserve"))
    apply_plan(path, output, plan, {op.operation_id for op in plan.operations})
    section = Document(output).sections[0]
    assert round(section.left_margin.mm) == 30
    assert round(section.right_margin.mm) == 41
    assert round(section.page_width.mm) == 220


def test_policy_changes_plan_version_even_without_operations(tmp_path: Path):
    path = tmp_path / "empty.docx"
    Document().save(path)
    inspection = inspect_document(path)
    policy = FormattingPolicy(mode="preserve")
    first = generate_plan(inspection, sparse_rules(), policy)
    policy.baseline_status = "configured"
    second = generate_plan(inspection, sparse_rules(), policy)
    assert first.plan_version != second.plan_version
