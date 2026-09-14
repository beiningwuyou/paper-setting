from __future__ import annotations

import zipfile
from pathlib import Path

from docx import Document
from paper_setting_core.verification import snapshot_package, validate_result


def _create_document(path: Path) -> None:
    document = Document()
    document.add_paragraph("目录占位项不是 DOCX 真实部件。")
    document.save(str(path))


def test_directory_placeholders_do_not_change_part_inventory(tmp_path: Path) -> None:
    source = tmp_path / "source-with-directories.docx"
    output = tmp_path / "roundtrip.docx"
    _create_document(source)
    with zipfile.ZipFile(source, "a") as archive:
        for directory in ("_rels/", "docProps/", "word/", "word/_rels/", "word/theme/"):
            archive.writestr(directory, b"")

    Document(str(source)).save(str(output))
    before = snapshot_package(source)
    after = snapshot_package(output)
    report = validate_result(source, output, [])

    assert before.part_names == after.part_names
    assert all(not name.endswith("/") for name in before.part_names)
    assert report.integrity_ok
    assert report.checks["part_inventory_unchanged"]


def test_missing_real_file_still_fails_part_inventory_check(tmp_path: Path) -> None:
    source = tmp_path / "source.docx"
    output = tmp_path / "missing-part.docx"
    _create_document(source)
    with zipfile.ZipFile(source) as original, zipfile.ZipFile(output, "w") as changed:
        for entry in original.infolist():
            if entry.filename != "docProps/app.xml":
                changed.writestr(entry, original.read(entry.filename))

    report = validate_result(source, output, [])

    assert not report.integrity_ok
    assert not report.checks["part_inventory_unchanged"]
    assert "part_inventory_unchanged" in report.differences
