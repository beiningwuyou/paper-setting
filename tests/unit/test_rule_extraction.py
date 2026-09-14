from io import BytesIO

import pytest
from docx import Document
from paper_setting_core.errors import RuleSourceInvalidError
from paper_setting_core.rulepacks import extract_rule_pack_draft, read_rule_source_file

GUIDE_TEXT = """
页面采用 A4 纸，纵向，上下边距 2.5cm，左边距 3cm，右边距 2.5cm。
正文：宋体，小四，Times New Roman，1.5 倍行距，首行缩进 2 字符，两端对齐。
一级标题：黑体，三号，加粗，居中，段前 12 磅，段后 6 磅，与下段同页。
参考文献条目：宋体，五号，固定值 20 磅。
"""


def test_extracts_page_role_and_line_spacing_rules() -> None:
    draft = extract_rule_pack_draft(GUIDE_TEXT, name="测试学位论文规则")

    assert draft.rule_pack.name == "测试学位论文规则"
    assert draft.rule_pack.schema_version == "2.0"
    assert draft.capability_report.executable is True
    assert draft.rule_pack.page.margin_left_mm == 30
    assert draft.rule_pack.page.margin_top_mm == 25
    body = draft.rule_pack.rule_for("body")
    assert body is not None
    assert body.character.east_asia_font == "宋体"
    assert body.character.latin_font == "Times New Roman"
    assert body.character.size_pt == 12
    assert body.paragraph.first_line_indent_pt == 24
    assert body.paragraph.line_spacing == 1.5
    assert body.paragraph.line_spacing_mode == "multiple"
    reference = draft.rule_pack.rule_for("reference_entry")
    assert reference is not None
    assert reference.paragraph.line_spacing == 20
    assert reference.paragraph.line_spacing_mode == "exact"
    assert draft.recognized_properties == len(draft.evidence)
    assert len(
        [item for item in draft.evidence if item.property_path == "page.margin_bottom_mm"]
    ) == 1


def test_reads_rules_from_docx_paragraphs_and_tables() -> None:
    document = Document()
    document.add_paragraph("页面：A4 纸，纵向。")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "正文"
    table.cell(0, 1).text = "宋体，小四，单倍行距"
    buffer = BytesIO()
    document.save(buffer)

    text, source_format = read_rule_source_file("guide.docx", buffer.getvalue())
    draft = extract_rule_pack_draft(text, source_filename="guide.docx", source_format=source_format)

    assert source_format == "docx"
    assert "正文" in text
    assert draft.rule_pack.rule_for("body") is not None


def test_rejects_unrecognized_or_unsupported_sources() -> None:
    with pytest.raises(RuleSourceInvalidError):
        extract_rule_pack_draft("这是一段与排版无关的普通文本。")
    with pytest.raises(RuleSourceInvalidError):
        read_rule_source_file("guide.pdf", b"not a pdf")


def test_extracts_notes_and_bibliography_removal_rules_without_role_context() -> None:
    draft = extract_rule_pack_draft(
        "5.注释一律采用脚注，文末不列参考文献。每页脚注重新编号，编号格式为带圈数字。",
        name="哲学分析",
    )

    notes = draft.rule_pack.advanced.notes
    assert notes.enabled is True
    assert notes.convert_inline_citations is True
    assert notes.convert_endnotes is True
    assert notes.delete_bibliography is True
    assert notes.numbering_restart == "each_page"
    assert notes.number_format == "decimal_enclosed_circle"
    assert {item.property_path for item in draft.evidence} >= {
        "advanced.notes.convert_inline_citations",
        "advanced.notes.convert_endnotes",
        "advanced.notes.delete_bibliography",
        "advanced.notes.numbering_restart",
        "advanced.notes.number_format",
    }
