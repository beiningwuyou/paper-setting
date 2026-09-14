from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from paper_setting_core.classification import classify_paragraph
from paper_setting_core.documents import inspect_document
from paper_setting_core.formatting import apply_plan
from paper_setting_core.planning import generate_plan
from paper_setting_core.rulepacks.defaults import default_rule_pack
from paper_setting_core.rulepacks.policy import FormattingPolicy


def test_generated_lists_and_extended_roles():
    examples = [
        ("第1章 绪论 1", "TOC 1", "toc_entry_1"),
        ("1.2 研究方法 3", "TOC 2", "toc_entry_2"),
        ("1.2.3 方法 5", "TOC 3", "toc_entry_3"),
        ("图1-1 系统 7", "Table of Figures", "figure_table_list_entry"),
        ("摘  要", "Normal", "abstract_heading"),
        ("Key Words: example", "Normal", "keywords"),
        ("1.2.3.4 四级", "Normal", "heading_4"),
        ("四级", "Heading 4", "heading_4"),
        ("Figure 3-1 Example", "Caption", "figure_caption_en"),
        ("Table 3-1 Example", "Caption", "table_caption_en"),
        ("致  谢", "Heading 1", "acknowledgments_heading"),
        ("博士/硕士学位论文", "Normal", "degree_label"),
    ]
    for text, style, role in examples:
        assert classify_paragraph(text, style, order=10)[0] == role


def test_abstract_and_bibliography_boundaries(tmp_path):
    doc = Document()
    rows = [
        ("博士学位论文", "degree_label"),
        ("原创性声明内容及签名日期不得应用正文行距。", "cover_frontmatter"),
        ("摘  要", "abstract_heading"),
        ("中文摘要的一段用于验证分区的文字。", "abstract_body"),
        ("关键词：测试", "keywords"),
        ("Abstract", "abstract_heading"),
        ("This is an abstract.", "abstract_body"),
        ("Key Words: Test", "keywords"),
        ("第1章 绪论", "heading_1"),
        ("这是一段正文，不得被当成英文摘要。", "body"),
        ("Figure 1-1 Test", "figure_caption_en"),
        ("注：测试图片的解释。", "caption_note"),
        ("参考文献", "bibliography_heading"),
        ("中文文献", "bibliography_group_heading"),
        ("张某. 书籍[M]. 出版社, 2024.", "reference_entry"),
        ("附录一 研究材料", "appendix_heading"),
        ("这里是附录的正文，不再属于参考文献。", "appendix_body"),
        ("致  谢", "acknowledgments_heading"),
        ("感谢帮助本研究的所有老师与同学。", "acknowledgments_body"),
        ("作者简历", "cv_heading"),
        ("2020—2024 学习经历。", "cv_body"),
    ]
    for text, _ in rows:
        doc.add_paragraph(text)
    path = tmp_path / "boundaries.docx"
    doc.save(path)
    assert [i.semantic_role for i in inspect_document(path).items if i.story == "body"] == [
        role for _, role in rows
    ]


def test_chapter_ends_abstract_without_keywords(tmp_path):
    doc = Document()
    doc.add_paragraph("Abstract")
    doc.add_paragraph("Text of abstract.")
    doc.add_paragraph("A chapter", "Heading 1")
    doc.add_paragraph("This is body text after a chapter.")
    path = tmp_path / "no-keywords.docx"
    doc.save(path)
    assert inspect_document(path).items[3].semantic_role == "body"


def test_table_note_skips_cells_and_reference_examples_are_not_entries(tmp_path):
    doc = Document()
    doc.add_paragraph("表1-1 示例")
    doc.add_paragraph("Table 1-1 Example")
    doc.add_table(rows=2, cols=2).cell(0, 0).text = "数据"
    doc.add_paragraph("注：表格注释。")
    doc.add_paragraph("参考文献")
    doc.add_paragraph("普通图书：")
    doc.add_paragraph("（顺序编码制）示例：")
    doc.add_paragraph("致谢")
    doc.add_paragraph("2024年10月")
    path = tmp_path / "table-note.docx"
    doc.save(path)
    items = {i.text_preview: i.semantic_role for i in inspect_document(path).items}
    assert items["注：表格注释。"] == "caption_note"
    assert items["普通图书："] == "bibliography_group_heading"
    assert items["（顺序编码制）示例："] == "template_instruction"
    assert items["2024年10月"] == "date_line"


def test_preserve_native_style_identity_and_numbering(tmp_path):
    doc = Document()
    p = doc.add_paragraph("绪论", "Heading 1")
    p.runs[0].font.size = Pt(28)
    num = OxmlElement("w:numPr")
    num_id = OxmlElement("w:numId")
    num_id.set(qn("w:val"), "1")
    num.append(num_id)
    p._p.get_or_add_pPr().append(num)
    native_styles = doc.styles.element.xml
    native_numbering = num.xml
    source, output = tmp_path / "source.docx", tmp_path / "out.docx"
    doc.save(source)
    pack = default_rule_pack()
    pack.rule_for("heading_1").preserve_style_identity = True
    policy = FormattingPolicy(mode="preserve")
    plan = generate_plan(inspect_document(source), pack, policy)
    approved = {op.operation_id for op in plan.operations}
    ops = apply_plan(source, output, plan, approved, confirmed_manual_operation_ids=approved)
    assert all(op.status == "applied" for op in ops)
    result = Document(output)
    assert result.paragraphs[0].style.name == "Heading 1"
    assert result.styles.element.xml == native_styles
    assert result.paragraphs[0]._p.pPr.numPr.xml == native_numbering
    assert not generate_plan(inspect_document(output), pack, policy).operations
