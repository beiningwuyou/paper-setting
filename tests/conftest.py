from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE


@pytest.fixture
def sample_docx(tmp_path: Path) -> Path:
    path = tmp_path / "论文样稿.docx"
    document = Document()
    if "标题 1" not in [style.name for style in document.styles]:
        document.styles.add_style("标题 1", WD_STYLE_TYPE.PARAGRAPH)
    document.add_paragraph("Agent 论文排版工具研究", style="Title")
    document.add_paragraph("摘要")
    document.add_paragraph("本文研究一种面向智能体的确定性论文排版方法。")
    document.add_paragraph("关键词：智能体；论文排版；Word")
    document.add_paragraph("第一章 绪论", style="Heading 1")
    document.add_paragraph("论文排版需要在保持内容完整的前提下统一页面与段落格式。")
    document.add_paragraph("1.1 研究背景", style="Heading 2")
    paragraph = document.add_paragraph("这是包含")
    run = paragraph.add_run("强调内容")
    run.bold = True
    paragraph.add_run("的正文段落，用于确认执行后文本不会发生变化。")
    document.add_paragraph("图 1-1 系统处理流程")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "项目"
    table.cell(0, 1).text = "说明"
    table.cell(1, 0).text = "格式"
    table.cell(1, 1).text = "保持内容不变"
    document.add_paragraph("参考文献")
    document.add_paragraph("[1] 张三. 论文排版方法研究[J]. 示例期刊, 2026.")
    document.save(path)
    return path

